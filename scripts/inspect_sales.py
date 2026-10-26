"""Read-only source inspection for the SOP dashboard."""
from pathlib import Path
import argparse
import json
import sys
import openpyxl


def main():
    sys.stdout.reconfigure(encoding="utf-8")
    parser = argparse.ArgumentParser()
    parser.add_argument("source", nargs="?", default="input_file")
    parser.add_argument("--file", default="")
    parser.add_argument("--sheet", default="")
    parser.add_argument("--rows", type=int, default=12)
    parser.add_argument("--cols", type=int, default=30)
    parser.add_argument("--start", type=int, default=1)
    args = parser.parse_args()
    for path in Path(args.source).glob("*.xlsx"):
        if path.name == "SOP计划维护.xlsx" and not args.file:
            continue
        if args.file not in path.name:
            continue
        workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
        for sheet in workbook.worksheets:
            if args.sheet not in sheet.title:
                continue
            print(json.dumps({"file": path.name, "sheet": sheet.title, "dimensions": [sheet.max_row, sheet.max_column]}, ensure_ascii=False))
            for row_num, row in enumerate(sheet.iter_rows(min_row=args.start, max_row=args.start + args.rows - 1, max_col=args.cols, values_only=True), args.start):
                print(json.dumps({"row": row_num, "values": row}, ensure_ascii=False, default=str))
        workbook.close()


if __name__ == "__main__":
    main()
