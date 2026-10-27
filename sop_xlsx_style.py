"""Shared presentation for offline SOP workbooks written with openpyxl."""
from __future__ import annotations

import os
import time
import xml.etree.ElementTree as ET
import zipfile
from copy import copy
from pathlib import Path

from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


INK = "243649"
BLUE = "0871B8"
TITLE_BLUE = "086BAE"
WHITE = "FFFFFF"
PAST = "F0F2F4"
BLOCKED = "DCE2E8"
RED = "B62025"
GRID = Border(*(Side(style="thin", color="8498AA") for _ in range(4)))


def plan_guidance(include_locks=False):
    return ("单位：辆。只填版本月计划量；可在主表新增、删除行或修改名称，删除表示停用并保留历史。合计自动计算。"
            + ("锁单行无需填写。" if include_locks else ""))


def _fill(color):
    return PatternFill(fill_type="solid", fgColor=color)


def setup_plan_sheet(sheet, *, last_row, last_column, visible_last_column,
                     columns, header_row=3, first_data_row=5,
                     identity_columns=3):
    """Mirror the shared blue, two-row working grid of plan_sheet_style.mjs."""
    sheet.sheet_view.showGridLines = False
    body_font = Font(name="Microsoft YaHei", size=11, color=INK)
    body_alignment = Alignment(vertical="center")
    for row in sheet.iter_rows(min_row=1, max_row=last_row, max_col=last_column):
        for cell in row:
            cell.font = body_font
            cell.alignment = body_alignment
    for row in range(1, last_row + 1):
        sheet.row_dimensions[row].height = 28
    if identity_columns not in (2, 3):
        raise ValueError("身份列须为品牌、车型，或品牌、车型、版本。")
    for column, width in enumerate((10, 20, 32)[:identity_columns], 1):
        sheet.column_dimensions[get_column_letter(column)].width = width
    for column in range(identity_columns + 1, visible_last_column + 1):
        sheet.column_dimensions[get_column_letter(column)].width = 10
    for item in columns:
        if item["kind"] == "year_total":
            sheet.column_dimensions[get_column_letter(item["column"])].width = 14
    for column in range(visible_last_column + 1, last_column + 1):
        dimension = sheet.column_dimensions[get_column_letter(column)]
        dimension.width = 0
        dimension.hidden = True

    sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=identity_columns)
    sheet.merge_cells(start_row=1, start_column=identity_columns + 1,
                      end_row=1, end_column=visible_last_column)
    sheet.merge_cells(start_row=2, start_column=1, end_row=2, end_column=visible_last_column)
    title = sheet.cell(1, 1)
    title.font = Font(name="Microsoft YaHei", size=18, color=TITLE_BLUE, bold=True)
    title.alignment = Alignment(horizontal="center", vertical="center")
    sheet.row_dimensions[1].height = 40
    period = sheet.cell(1, identity_columns + 1)
    period.font = Font(name="Microsoft YaHei", size=12, color="52677C", bold=True)
    period.alignment = Alignment(horizontal="left", vertical="center")
    subtitle = sheet.cell(2, 1)
    subtitle.font = Font(name="Microsoft YaHei", size=10, color="52677C")
    subtitle.alignment = Alignment(horizontal="left", vertical="center")
    sheet.row_dimensions[2].height = 29

    for row in range(header_row, first_data_row):
        sheet.row_dimensions[row].height = 30
        for column in range(1, last_column + 1):
            cell = sheet.cell(row, column)
            cell.fill = _fill(BLUE)
            cell.font = Font(name="Microsoft YaHei", size=11, bold=True, color=WHITE)
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = GRID
    for item in columns:
        if item["kind"] == "month" and item.get("historical"):
            sheet.cell(first_data_row - 1, item["column"]).fill = _fill("6F8498")


def style_plan_row(sheet, *, row, kind, visible_last_column, columns,
                   version="", month_states=None, identity_columns=3):
    if identity_columns not in (2, 3):
        raise ValueError("身份列须为品牌、车型，或品牌、车型、版本。")
    month_states = month_states or {}
    for column in range(1, visible_last_column + 1):
        cell = sheet.cell(row, column)
        cell.border = GRID
        if column > identity_columns:
            cell.number_format = "#,##0"
            cell.alignment = Alignment(horizontal="right", vertical="center")

    if kind in {"product", "manual_entry"}:
        for column in range(1, identity_columns + 1):
            cell = sheet.cell(row, column)
            if kind == "manual_entry":
                cell.fill = _fill(WHITE)
            cell.font = Font(name="Microsoft YaHei", size=11, bold=True, color=INK)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        if kind == "product":
            sheet.row_dimensions[row].height = 46 if len(version) > 20 else 30
        for item in columns:
            if item["kind"] != "month":
                continue
            cell = sheet.cell(row, item["column"])
            state = month_states.get(cell.coordinate, {})
            active = state.get("editable") is not False
            cell.fill = _fill(BLOCKED if not active else PAST if item.get("historical") else WHITE)
            if kind == "product" and state.get("applicable") is False:
                font = copy(cell.font)
                font.color = "68798A"
                cell.font = font
    else:
        fill = _fill("DCEAF7" if kind == "grand_total" else "EAF2F9")
        for column in range(1, visible_last_column + 1):
            cell = sheet.cell(row, column)
            cell.fill = fill
            cell.font = Font(name="Microsoft YaHei", size=11, bold=True,
                             color=RED if column > identity_columns else INK)

    for item in columns:
        if item["kind"] == "year_total":
            cell = sheet.cell(row, item["column"])
            cell.fill = _fill("FFF8E4")
            cell.font = Font(name="Microsoft YaHei", size=11, bold=True, color=RED)


def write_formula_caches(path, sheet_names, cached_values, hidden_columns=None):
    """Populate formula caches and conceal import-only columns in the saved XLSX."""
    path = Path(path)
    hidden_columns = hidden_columns or {}
    main_ns = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
    ET.register_namespace("", main_ns)
    ET.register_namespace("r", "http://schemas.openxmlformats.org/officeDocument/2006/relationships")
    tag = lambda name: f"{{{main_ns}}}{name}"
    replacements = {}
    with zipfile.ZipFile(path) as archive:
        entries = [(info, archive.read(info.filename)) for info in archive.infolist()]
        entry_map = {info.filename: content for info, content in entries}
        for index, name in enumerate(sheet_names, 1):
            values = cached_values.get(name, {})
            hidden_range = hidden_columns.get(name)
            if not values and not hidden_range:
                continue
            member = f"xl/worksheets/sheet{index}.xml"
            root = ET.fromstring(entry_map[member])
            cells = {cell.attrib.get("r"): cell for cell in root.findall(f"{tag('sheetData')}/{tag('row')}/{tag('c')}")}
            for address, value in values.items():
                cell = cells.get(address)
                if cell is None or cell.find(tag("f")) is None:
                    raise ValueError(f"缺少汇总公式：{name}!{address}")
                for old in cell.findall(tag("v")):
                    cell.remove(old)
                if value is None:
                    cell.set("t", "str")
                    ET.SubElement(cell, tag("v"))
                else:
                    if type(value) is not int or value < 0:
                        raise ValueError(f"汇总缓存必须是非负整数：{name}!{address}")
                    cell.attrib.pop("t", None)
                    ET.SubElement(cell, tag("v")).text = str(value)
            if hidden_range:
                first, last = hidden_range
                cols = root.find(tag("cols"))
                if cols is None:
                    cols = ET.Element(tag("cols"))
                    sheet_data = root.find(tag("sheetData"))
                    root.insert(list(root).index(sheet_data), cols)
                covered = set()
                for col in list(cols):
                    low, high = int(col.attrib["min"]), int(col.attrib["max"])
                    if high < first or low > last:
                        continue
                    if low < first:
                        col.set("max", str(first - 1))
                        concealed = copy(col)
                        concealed.set("min", str(first))
                        concealed.set("max", str(high))
                        cols.insert(list(cols).index(col) + 1, concealed)
                        col, low = concealed, first
                    col.attrib.update({"hidden": "1", "width": "0", "customWidth": "1"})
                    covered.update(range(low, min(high, last) + 1))
                for number in range(first, last + 1):
                    if number not in covered:
                        ET.SubElement(cols, tag("col"), {"min": str(number), "max": str(number),
                                                         "hidden": "1", "width": "0", "customWidth": "1"})
            replacements[member] = ET.tostring(root, encoding="utf-8", xml_declaration=True)
    temporary = path.with_name(path.name + ".cache.tmp")
    try:
        with zipfile.ZipFile(temporary, "w") as output:
            for info, content in entries:
                output.writestr(info, replacements.get(info.filename, content))
        # Windows indexing and preview tools can briefly hold newly written
        # .xlsx files. Retry the same atomic swap while the lock clears.
        for attempt in range(8):
            try:
                os.replace(temporary, path)
                break
            except PermissionError:
                if attempt == 7:
                    raise
                time.sleep(0.15 * (attempt + 1))
    finally:
        if temporary.exists():
            temporary.unlink()
