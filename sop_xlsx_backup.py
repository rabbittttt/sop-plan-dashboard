"""Write frozen snapshot backups when the Node spreadsheet package is unavailable.

The input is the immutable snapshot payload, never the live plan store. All
strings are written as literal text, including timestamps and formula-like
notes, so an Excel backup cannot turn user text into a formula.
"""
from __future__ import annotations

import json
import math
import re
import time
import zipfile
import xml.etree.ElementTree as ET
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


MONTH = re.compile(r"^\d{4}-(0[1-9]|1[0-2])$")
FONT_NAME = "Microsoft YaHei"
NAVY = "FF173B64"
BODY = "FF253A54"
MUTED = "FF60758E"
PALE = "FFF2F5F9"


def _scalar(value):
    if isinstance(value, (dict, list)):
        return json.dumps(value, ensure_ascii=False, separators=(",", ":"))
    return value


def _coalesce(value, fallback):
    return fallback if value is None else value


def _put(sheet, row, column, value):
    value = _scalar(value)
    if isinstance(value, str) and len(value) > 32767:
        raise ValueError("Snapshot text exceeds Excel cell limit")
    cell = sheet.cell(row, column, value)
    if isinstance(value, str):
        cell.data_type = "s"  # Preserve exact text; never evaluate user formulas.
    return cell


def _row(sheet, row, values):
    for column, value in enumerate(values, 1):
        _put(sheet, row, column, value)


def _base(sheet, last_row, last_column):
    sheet.sheet_view.showGridLines = False
    font = Font(name=FONT_NAME, size=11, color=BODY)
    alignment = Alignment(vertical="center")
    for row in sheet.iter_rows(min_row=1, max_row=last_row, max_col=last_column):
        for cell in row:
            cell.font = font
            cell.alignment = alignment
    for index in range(1, last_row + 1):
        sheet.row_dimensions[index].height = 29


def _heading(sheet, row, last_column):
    fill = PatternFill("solid", fgColor=NAVY)
    font = Font(name=FONT_NAME, size=11, bold=True, color="FFFFFFFF")
    border = Border(right=Side(style="thin", color="FFFFFFFF"))
    for cells in sheet.iter_rows(min_row=row, max_row=row, max_col=last_column):
        for cell in cells:
            cell.fill = fill
            cell.font = font
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = border
    sheet.row_dimensions[row].height = 32


def _title(sheet, size=15):
    sheet["A2"].font = Font(name=FONT_NAME, size=size, bold=True, color="FF183C65")
    sheet["A3"].font = Font(name=FONT_NAME, size=11, color=MUTED)


def _source_label(source):
    if isinstance(source, str):
        return source
    if not isinstance(source, dict):
        return ""
    name = source.get("name") or source.get("file") or source.get("path") or ""
    return " / ".join(str(part) for part in (name, source.get("sheet") or "") if part)


def _preserve_empty_strings(output, workbook):
    """Openpyxl omits empty text cells; the old writer retained literal ''."""
    empty = [[cell.coordinate for row in sheet for cell in row if cell.value == ""] for sheet in workbook]
    if not any(empty):
        return
    namespace = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
    ET.register_namespace("", namespace)
    ET.register_namespace("r", "http://schemas.openxmlformats.org/officeDocument/2006/relationships")
    with zipfile.ZipFile(output) as archive:
        parts = {item.filename: (item, archive.read(item.filename)) for item in archive.infolist()}
    for index, addresses in enumerate(empty, 1):
        if not addresses:
            continue
        name = f"xl/worksheets/sheet{index}.xml"
        root = ET.fromstring(parts[name][1])
        data = root.find(f"{{{namespace}}}sheetData")
        rows = {int(row.attrib["r"]): row for row in data}
        cells = {cell.attrib["r"]: cell for row in data for cell in row if cell.tag == f"{{{namespace}}}c"}
        for address in addresses:
            cell = cells.get(address)
            if cell is None:
                number = int("".join(char for char in address if char.isdigit()))
                row = rows.get(number)
                if row is None:
                    row = ET.SubElement(data, f"{{{namespace}}}row", {"r": str(number)})
                    rows[number] = row
                cell = ET.SubElement(row, f"{{{namespace}}}c", {"r": address})
            cell.set("t", "inlineStr")
            for child in list(cell):
                cell.remove(child)
            inline = ET.SubElement(cell, f"{{{namespace}}}is")
            ET.SubElement(inline, f"{{{namespace}}}t").text = ""
        parts[name] = (parts[name][0], ET.tostring(root, encoding="utf-8", xml_declaration=True))
    temporary = output.with_name(output.name + ".text.tmp")
    with zipfile.ZipFile(temporary, "w", zipfile.ZIP_DEFLATED) as archive:
        for info, contents in parts.values():
            archive.writestr(info, contents)
    for attempt in range(10):
        try:
            temporary.replace(output)
            break
        except PermissionError:
            if attempt == 9:
                raise
            time.sleep(0.1)


def _snapshot_rows(frozen):
    snapshot = frozen.get("snapshot") or {}
    meta = frozen.get("meta") or {}
    plans = frozen.get("plans") if isinstance(frozen.get("plans"), list) else []
    original_products = frozen.get("products") if isinstance(frozen.get("products"), list) else []
    products = [dict(product) for product in original_products]
    directory = {}
    for product in products:
        pid = product.get("id")
        if not isinstance(pid, str) or not pid or pid in directory:
            raise ValueError("Snapshot product IDs must be unique nonempty strings")
        directory[pid] = product
    months = set(meta.get("months") if isinstance(meta.get("months"), list) else [])
    lookup = {}
    orphan_count = 0
    for plan in plans:
        pid, month = plan.get("product_id"), plan.get("month")
        if not isinstance(pid, str) or not pid:
            raise ValueError("Snapshot plan product_id is missing")
        if not isinstance(month, str) or not MONTH.fullmatch(month):
            raise ValueError("Invalid snapshot plan month")
        key = (pid, month)
        if key in lookup:
            raise ValueError("Duplicate snapshot plan key")
        for field in ("sop", "forecast"):
            quantity = plan.get(field)
            if quantity is not None and (type(quantity) not in (int, float) or not math.isfinite(quantity)):
                raise ValueError(f"Invalid snapshot {field} quantity")
        lookup[key] = plan
        months.add(month)
        if pid not in directory:
            product = {"id": pid, "brand": "未记录品牌", "model": "快照目录缺失", "version": "原产品ID保留"}
            products.append(product)
            directory[pid] = product
            orphan_count += 1
    if any(not isinstance(month, str) or not MONTH.fullmatch(month) for month in months):
        raise ValueError("Invalid snapshot metadata month")
    return snapshot, meta, plans, original_products, products, directory, sorted(months), lookup, orphan_count


def write_snapshot_workbook(payload: dict) -> None:
    """Write the XLSX at ``payload['output_path']`` using only openpyxl."""
    if not isinstance(payload, dict) or not payload.get("output_path") or not isinstance(payload.get("snapshot_data"), dict):
        raise ValueError("output_path and snapshot_data are required")
    frozen = payload["snapshot_data"]
    snapshot, meta, plans, original_products, products, directory, months, lookup, orphan_count = _snapshot_rows(frozen)
    exported_at = str(_coalesce(payload.get("exported_at"), datetime.now(timezone.utc).isoformat()))

    workbook = Workbook()
    grid = workbook.active
    grid.title = "SOP计划"
    detail = workbook.create_sheet("计划明细")
    info = workbook.create_sheet("备份信息")

    grid_last_row = max(10, 9 + len(products))
    grid_last_col = max(4, 4 + len(months))
    _base(grid, grid_last_row, grid_last_col)
    _put(grid, 2, 1, "鸿蒙智行 SOP计划快照")
    _put(grid, 3, 1, "固定值备份，不参与维护表同步。空白表示待维护，0表示明确不安排。")
    for row, label, value in (
        (4, "业务日期", snapshot.get("business_date")),
        (5, "快照截点", _coalesce(snapshot.get("cutoff_at"), meta.get("snapshot_cutoff_at"))),
        (6, "快照生成", snapshot.get("created_at")),
        (7, "Excel生成", exported_at),
    ):
        _put(grid, row, 1, label)
        _put(grid, row, 2, value)
    _row(grid, 9, ["品牌", "车型", "版本", "产品ID", *months])
    if products:
        for row, product in enumerate(products, 10):
            values = [(lookup.get((product["id"], month)) or {}).get("sop") for month in months]
            _row(grid, row, [product.get("brand") or "", product.get("model") or "", product.get("version") or "",
                             product["id"], *values])
    else:
        _put(grid, 10, 1, "本快照未记录产品与SOP计划。")
    _title(grid, size=16)
    grid.row_dimensions[2].height = 34
    for row in range(4, 8):
        grid.cell(row, 1).font = Font(name=FONT_NAME, size=11, color="FF64778D")
    for column, width in enumerate((13, 24, 40, 28), 1):
        grid.column_dimensions[get_column_letter(column)].width = width
    for column in range(5, grid_last_col + 1):
        grid.column_dimensions[get_column_letter(column)].width = 14
    _heading(grid, 9, grid_last_col)
    if products:
        identity_fill = PatternFill("solid", fgColor=PALE)
        number_fill = PatternFill("solid", fgColor="FFFAFCFF")
        group_border = Border(top=Side(style="thin", color="FFCCD8E7"))
        for row, product in enumerate(products, 10):
            for column in range(1, 5):
                cell = grid.cell(row, column)
                cell.fill = identity_fill
                cell.alignment = Alignment(vertical="center", wrap_text=True)
                if column == 4:
                    cell.font = Font(name=FONT_NAME, size=10, color="FF61738A")
            for column in range(5, grid_last_col + 1):
                cell = grid.cell(row, column)
                cell.fill = number_fill
                cell.number_format = "#,##0"
                cell.alignment = Alignment(horizontal="right", vertical="center")
            if row == 10 or any(product.get(field) != products[row - 11].get(field) for field in ("brand", "model")):
                for column in range(1, grid_last_col + 1):
                    grid.cell(row, column).border = group_border
            grid.row_dimensions[row].height = 44 if len(str(product.get("version") or "")) > 24 else 34
    grid.freeze_panes = "E10"
    grid.sheet_properties.tabColor = NAVY

    plan_fields = ["product_id", "month", "forecast", "sop", "owner", "reason", "notes", "revision", "updated_at"]
    extra_fields = sorted({field for plan in plans for field in plan if field not in plan_fields})
    headers = ["product_id", "品牌", "车型", "版本", *plan_fields[1:], *extra_fields]
    detail_last_row = max(6, 5 + len(plans))
    _base(detail, detail_last_row, len(headers))
    _put(detail, 2, 1, "计划明细（冻结记录）")
    _put(detail, 3, 1, "保留原计划字段与产品ID。forecast为原快照留存字段，不是本备份的输入项。")
    _row(detail, 5, headers)
    if plans:
        for row, plan in enumerate(plans, 6):
            product = directory[plan["product_id"]]
            _row(detail, row, [plan["product_id"], product.get("brand") or "", product.get("model") or "",
                               product.get("version") or "", *(plan.get(field) for field in plan_fields[1:]),
                               *(plan.get(field) for field in extra_fields)])
    else:
        _put(detail, 6, 1, "本快照没有已保存的计划记录。")
    _title(detail)
    widths = (27, 10, 20, 38, 13, 15, 15, 15, 20, 48, 11, 42)
    for column, _ in enumerate(headers, 1):
        detail.column_dimensions[get_column_letter(column)].width = widths[column - 1] if column <= len(widths) else 30
    _heading(detail, 5, len(headers))
    if plans:
        alternate = PatternFill("solid", fgColor="FFF5F8FC")
        for row in range(6, detail_last_row + 1):
            detail.row_dimensions[row].height = 42
            for column in range(1, len(headers) + 1):
                cell = detail.cell(row, column)
                cell.alignment = Alignment(vertical="center", wrap_text=column != 12)
                if row % 2 == 0:
                    cell.fill = alternate
            for column in (6, 7):
                detail.cell(row, column).number_format = "#,##0"
                detail.cell(row, column).alignment = Alignment(horizontal="right", vertical="center", wrap_text=True)
            detail.cell(row, 11).number_format = "0"
    detail.freeze_panes = "F6"
    detail.sheet_properties.tabColor = "FF7890AE"

    sources = frozen.get("sources") if isinstance(frozen.get("sources"), list) and frozen["sources"] else meta.get("source_files") if isinstance(meta.get("source_files"), list) else []
    kind = {"daily": "每日快照", "sunday": "周日考核快照", "manual": "会议快照"}.get(snapshot.get("kind"), snapshot.get("kind") or "未记录")
    info_rows = [
        ["快照ID", snapshot.get("id")], ["快照类型", kind], ["快照名称", snapshot.get("label")],
        ["业务日期", snapshot.get("business_date")], ["规定截点", _coalesce(snapshot.get("cutoff_at"), meta.get("snapshot_cutoff_at"))],
        ["快照实际生成时间", snapshot.get("created_at")], ["Excel实际生成时间", exported_at],
        ["时间口径", "保留快照原始时间和时区偏移。业务截点按北京时间。"],
        ["备份性质", "冻结SOP计划备份，不参与维护表同步。修改本文件不会更新看板。"],
        ["数量单位", "辆"], ["空白与0", "空白表示未维护或快照未记录；0为明确保存的零计划。"],
        ["计划月份", "、".join(months) if months else "快照未记录计划月份"],
        ["快照产品数", len(original_products)], ["目录缺失产品数", orphan_count], ["原始计划记录数", len(plans)],
        ["已填写SOP的记录数", sum(plan.get("sop") is not None for plan in plans)],
        ["来源数据截至", _coalesce(meta.get("data_through_label"), meta.get("data_through"))],
        ["取数范围", "仅使用这一快照的产品目录、计划记录和已开放月份，不取当前计划。"],
    ]
    info_rows.extend([f"来源{index}", _source_label(source)] for index, source in enumerate(sources, 1))
    info_last_row = 5 + len(info_rows)
    _base(info, info_last_row, 2)
    _put(info, 2, 1, "备份信息")
    _put(info, 3, 1, "计划主表用于回看，计划明细保留原始字段用于核对。")
    _row(info, 5, ["项目", "冻结值或说明"])
    for row, values in enumerate(info_rows, 6):
        _row(info, row, values)
    _title(info)
    info.column_dimensions["A"].width = 27
    info.column_dimensions["B"].width = 103
    _heading(info, 5, 2)
    for row, values in enumerate(info_rows, 6):
        info.cell(row, 1).fill = PatternFill("solid", fgColor=PALE)
        info.cell(row, 2).alignment = Alignment(vertical="center", wrap_text=True)
        if len(str(values[1] or "")) > 65:
            info.row_dimensions[row].height = 44
    info.freeze_panes = "A6"
    info.sheet_properties.tabColor = "FF9AAABD"

    output = Path(payload["output_path"])
    output.parent.mkdir(parents=True, exist_ok=True)
    try:
        workbook.save(output)
    finally:
        workbook.close()
    _preserve_empty_strings(output, workbook)
