"""Safe two-way synchronization between the managed SOP workbook and SQLite.

openpyxl is used only to read. The bundled artifact-tool authoring script writes XLSX.
The hidden baseline is verified against a persistent export ledger, then merged per SOP cell.
"""
from __future__ import annotations

import hashlib
import io
import json
import math
import os
import subprocess
import threading
import uuid
from pathlib import Path

from sop_runtime import node_executable
from sop_store import ConflictError, TZ, ValidationError, catalog_name_key, horizon, month_add, month_range, stamp

SCHEMA_VERSION = 5
LAYOUT_VERSION = "continuous-inline-v1"
PREVIOUS_CONTINUOUS_LAYOUT = "continuous-blue-v2"
LEGACY_CONTINUOUS_LAYOUT = "continuous-blue-v1"
LEGACY_ANNUAL_LAYOUT = "annual-blue-v1"


def _column(number):
    result = ""
    while number:
        number, remainder = divmod(number - 1, 26)
        result = chr(65 + remainder) + result
    return result


def _annual_layout(products, months, editable_months=None):
    """Trusted coordinate map; display years never expand editable business months."""
    editable = set(months if editable_months is None else editable_months)
    grouped = {}
    for product in products:
        grouped.setdefault(product.get("brand", ""), {}).setdefault(product.get("model", ""), []).append(product)
    sheets, cells = [], []
    for year in sorted({month[:4] for month in months}):
        name = f"{year}年SOP"
        rows, merges, formulas, brand_totals = [], [], {}, []
        index = 4
        for brand, models in grouped.items():
            brand_start, model_totals = index, []
            for model, members in models.items():
                start = index
                for product in members:
                    rows.append({"row": index, "kind": "product", "product_id": product["id"],
                                 "brand": brand, "model": model, "version": product.get("version", ""),
                                 "label_a": brand if index == brand_start else None,
                                 "label_b": model if index == start else None,
                                 "label_c": product.get("version") or None})
                    for month in months:
                        if month.startswith(year + "-"):
                            cells.append({"product_id": product["id"], "month": month,
                                          "sheet": name, "cell": f"{_column(int(month[-2:]) + 3)}{index}",
                                          "editable": month in editable})
                    formulas[f"P{index}"] = f'=IF(COUNT(D{index}:O{index})=0,"",SUM(D{index}:O{index}))'
                    index += 1
                if index - start > 1:
                    merges.append(f"B{start}:B{index-1}")
                rows.append({"row": index, "kind": "model_total", "brand": brand, "model": model,
                             "label_a": None, "label_b": f"{model} 合计", "label_c": None})
                merges.append(f"B{index}:C{index}")
                for col in range(4, 17):
                    letter = _column(col)
                    ref = f"{letter}{start}:{letter}{index-1}"
                    formulas[f"{letter}{index}"] = f'=IF(COUNT({ref})=0,"",SUM({ref}))'
                model_totals.append(index)
                index += 1
            rows.append({"row": index, "kind": "brand_total", "brand": brand,
                         "label_a": None, "label_b": f"{brand} 合计", "label_c": None})
            merges += [f"A{brand_start}:A{index}", f"B{index}:C{index}"]
            for col in range(4, 17):
                letter = _column(col)
                refs = ",".join(f"{letter}{r}" for r in model_totals)
                formulas[f"{letter}{index}"] = f'=IF(COUNT({refs})=0,"",SUM({refs}))'
            brand_totals.append(index)
            index += 2
        rows.append({"row": index, "kind": "grand_total", "label_a": "总计", "label_b": None, "label_c": None})
        merges.append(f"A{index}:C{index}")
        for col in range(4, 17):
            letter = _column(col)
            refs = ",".join(f"{letter}{r}" for r in brand_totals)
            formulas[f"{letter}{index}"] = f'=IF(COUNT({refs})=0,"",SUM({refs}))'
        sheets.append({"name": name, "year": year, "rows": rows, "merges": merges,
                       "formulas": formulas, "last_row": index,
                       "active_months": [m for m in months if m.startswith(year + "-") and m in editable],
                       "shown_months": [m for m in months if m.startswith(year + "-")]})
    return {"sheets": sheets, "cells": cells}


def _signature(context):
    keys = ["months", "products", "baseline"]
    if context.get("schema_version") in (2, 3, 4, 5):
        keys += ["schema_version", "layout_version", "layout", "plan_end", "editable_months"]
    if context.get("schema_version") in (3, 4, 5):
        keys += ["current_month", "business_months"]
    return hashlib.sha256(json.dumps({key: context[key] for key in keys}, ensure_ascii=False, sort_keys=True).encode("utf-8")).hexdigest()


def _continuous_layout(products, months, editable_months, current_month, plans=None, include_manual=True):
    """Single chronological input grid. Year totals never form input mappings."""
    plans = plans or {}
    name, columns, years = "SOP计划", [], []
    for year in sorted({month[:4] for month in months}):
        first = len(columns) + 4
        year_months = [m for m in months if m.startswith(year + "-")]
        for month in year_months:
            columns.append({"column": len(columns) + 4, "kind": "month", "month": month,
                            "label": f"{int(month[-2:])}月", "historical": month < current_month})
        total = len(columns) + 4
        columns.append({"column": total, "kind": "year_total", "year": year,
                        "label": "年度合计" if len(year_months) == 12 else "期间合计"})
        years.append({"year": year, "first_column": first, "last_column": total,
                      "last_month_column": total - 1})
    product_column, kind_column = len(columns) + 4, len(columns) + 5
    headers = [[None] * kind_column for _ in range(2)]
    headers[0][:3] = ["品牌", "车型", "版本"]
    for year in years:
        headers[0][year["first_column"] - 1] = f"{year['year']}年"
    for col in columns:
        headers[1][col["column"] - 1] = col["label"]
    headers[1][-2:] = ["product_id", "row_kind"]
    grouped, rows, cells, formulas, brand_totals = {}, [], [], {}, []
    for product in products:
        grouped.setdefault(product.get("brand", ""), {}).setdefault(product.get("model", ""), []).append(product)
    merges = ["A3:A4", "B3:B4", "C3:C4"]
    merges += [f"{_column(y['first_column'])}3:{_column(y['last_column'])}3" for y in years]
    editable = set(editable_months)
    index = 5
    def totals(row, members):
        for col in columns:
            letter = _column(col["column"])
            refs = ",".join(f"{letter}{r}" for r in members)
            formulas[f"{letter}{row}"] = f'=IF(COUNT({refs})=0,"",SUM({refs}))' if refs else '=""'
    for brand, models in grouped.items():
        brand_start, model_totals = index, []
        for model, members in models.items():
            start = index
            for product in members:
                active = product.get("active", True)
                version = product.get("version", "")
                label = version + ("（停用）" if not active else "")
                rows.append({"row": index, "kind": "product", "product_id": product["id"],
                             "brand": brand, "model": model, "version": version, "active": active,
                             "label_a": brand if index == brand_start else None,
                             "label_b": model if index == start else None, "label_c": label or None})
                for col in columns:
                    if col["kind"] == "month":
                        month = col["month"]
                        record = plans.get((product["id"], month), {})
                        historical_value = month < current_month and (record.get("sop") is not None or record.get("forecast") is not None)
                        in_lifecycle = (not product.get("start_month") or month >= product["start_month"]) and (not product.get("end_month") or month <= product["end_month"])
                        applicable = in_lifecycle or historical_value
                        cells.append({"product_id": product["id"], "month": month, "sheet": name,
                                      "cell": f"{_column(col['column'])}{index}",
                                      "applicable": applicable,
                                      "editable": applicable and month in editable and (active or month < current_month)})
                for year in years:
                    ref = f"{_column(year['first_column'])}{index}:{_column(year['last_month_column'])}{index}"
                    formulas[f"{_column(year['last_column'])}{index}"] = f'=IF(COUNT({ref})=0,"",SUM({ref}))'
                index += 1
            if index - start > 1:
                merges.append(f"B{start}:B{index - 1}")
            rows.append({"row": index, "kind": "model_total", "label_a": None,
                         "label_b": f"{model} 合计", "label_c": None})
            merges.append(f"B{index}:C{index}")
            totals(index, range(start, index))
            model_totals.append(index)
            index += 1
        rows.append({"row": index, "kind": "brand_total", "label_a": None,
                     "label_b": f"{brand} 合计", "label_c": None})
        merges += [f"A{brand_start}:A{index}", f"B{index}:C{index}"]
        totals(index, model_totals)
        brand_totals.append(index)
        index += 2
    rows.append({"row": index, "kind": "grand_total", "label_a": "总计", "label_b": None, "label_c": None})
    merges.append(f"A{index}:C{index}")
    totals(index, brand_totals)
    manual_section = None
    if include_manual:
        title_row = index + 2
        first_row = title_row + 3
        manual_section = {"title_row": title_row, "header_row": title_row + 1,
                          "first_row": first_row, "last_row": first_row + 9,
                          "title": "新增车型/版本：填写品牌、车型、版本及计划量，保存后自动并入上方目录。空白行忽略。"}
        merges += [f"A{title_row}:{_column(product_column - 1)}{title_row}",
                   f"A{title_row+1}:A{title_row+2}", f"B{title_row+1}:B{title_row+2}",
                   f"C{title_row+1}:C{title_row+2}"]
        merges += [f"{_column(y['first_column'])}{title_row+1}:{_column(y['last_column'])}{title_row+1}" for y in years]
        for row in range(first_row, first_row + 10):
            rows.append({"row": row, "kind": "manual_entry"})
            for year in years:
                ref = f"{_column(year['first_column'])}{row}:{_column(year['last_month_column'])}{row}"
                formulas[f"{_column(year['last_column'])}{row}"] = f'=IF(COUNT({ref})=0,"",SUM({ref}))'
        index = first_row + 9
    return {"sheets": [{"name": name, "rows": rows, "merges": merges, "formulas": formulas,
                        "columns": columns, "years": years, "headers": headers, "header_row": 3,
                        "first_data_row": 5, "last_row": index, "last_column": kind_column,
                        "product_column": product_column, "kind_column": kind_column,
                        "visible_last_column": product_column - 1, "manual_section": manual_section}], "cells": cells}


def _inline_layout(products, months, editable_months, current_month, plans=None):
    """Editable identity rows share the normal grid; IDs travel with Excel rows."""
    import re
    plans = plans or {}
    layout = _continuous_layout(products, months, editable_months, current_month, plans, include_manual=False)
    spec = layout["sheets"][0]
    spec["key_column"] = spec["last_column"] + 1
    spec["last_column"] = spec["key_column"]
    for header in spec["headers"]:
        header.append(None)
    spec["headers"][1][-1] = "row_key"
    # Repeated labels keep identity explicit after inserting or removing a row.
    spec["merges"] = [merge for merge in spec["merges"]
                      if not (re.fullmatch(r"([ABC])(\d+):\1(\d+)", merge)
                              and int(re.fullmatch(r"([ABC])(\d+):\1(\d+)", merge).group(2)) >= 5)]
    catalog = {p["id"]: p for p in products}
    for row in spec["rows"]:
        row["key"] = "plan:" + row["product_id"] if row["kind"] == "product" else "row:" + str(row["row"])
        row["metric"] = "SOP计划量"
        if row["kind"] == "product":
            product = catalog[row["product_id"]]
            row.update(label_a=product["brand"], label_b=product["model"], label_c=product["version"],
                       hidden=not product.get("active", True))
            row["values"] = {month: plans.get((product["id"], month), {}).get("sop") for month in months}
        else:
            row["values"] = {}
        row["labels"] = [row.get("label_" + name) for name in "abc"]
    spec["products"] = products
    spec["editable_months_by_product"] = {p["id"]: [c["month"] for c in layout["cells"]
        if c["product_id"] == p["id"] and c["editable"]] for p in products}
    spec["manual_editable_months"] = editable_months
    spec["plan_revisions_by_product"] = {p["id"]: {month: record["revision"] for (pid, month), record in plans.items()
        if pid == p["id"]} for p in products}
    return layout


def _cell_quantity(value):
    if value is None or value == "":
        return None
    if isinstance(value, bool) or not isinstance(value, (int, float)) or not math.isfinite(value) or value < 0 or value > 10**12 or int(value) != value:
        raise ValidationError("SOP 数量必须为非负整数，空白表示未填写；不能填写公式。")
    return int(value)


def _key(pid, month):
    return f"{pid}\t{month}"


class WorkbookSync:
    def __init__(self, store, path, writer=None, root=None):
        self.store = store
        self.path = Path(path).resolve()
        self.root = Path(root or Path(__file__).resolve().parent).resolve()
        self.writer = writer or self._node_writer
        self.lock = threading.RLock()
        self.path.parent.mkdir(parents=True, exist_ok=True)
        self._last_file_stat = None
        self._last_digest = None
        self._export_signature = None
        self._input_checked = False
        self._state = {"path": str(self.path), "status": "idle", "message": "等待同步计划表。", "last_sync": None, "conflicts": []}
        with store.lock:
            store.db.executescript("""
                CREATE TABLE IF NOT EXISTS workbook_exports (
                    token TEXT PRIMARY KEY, created_at TEXT NOT NULL, payload TEXT NOT NULL);
                CREATE TRIGGER IF NOT EXISTS workbook_export_no_update BEFORE UPDATE ON workbook_exports
                    BEGIN SELECT RAISE(ABORT, 'Workbook export baselines are immutable'); END;
                CREATE TRIGGER IF NOT EXISTS workbook_export_no_delete BEFORE DELETE ON workbook_exports
                    BEGIN SELECT RAISE(ABORT, 'Workbook export baselines are immutable'); END;
                CREATE TABLE IF NOT EXISTS workbook_imports (
                    token TEXT PRIMARY KEY, accepted_values TEXT NOT NULL, last_digest TEXT NOT NULL, imported_at TEXT NOT NULL);
            """)

    def status(self):
        # Returning a copied small object avoids acquiring the sync lock during a long XLSX write.
        return {**self._state, "conflicts": list(self._state["conflicts"])}

    def _set_status(self, status, message, conflicts=None, synced=False):
        self._state = {**self._state, "status": status, "message": message, "conflicts": conflicts or []}
        if synced:
            self._state["last_sync"] = stamp(self.store.clock())
        return self.status()

    def _node_writer(self, payload):
        node = node_executable()
        script = self.root / "scripts/build_plan_workbook.mjs"
        if not node or not script.is_file():
            raise ValidationError("计划表生成组件尚未就绪，请检查本机 Node 运行环境和生成脚本。")
        completed = subprocess.run([node, str(script)], input=json.dumps(payload, ensure_ascii=False),
                                   text=True, encoding="utf-8", capture_output=True, timeout=120, cwd=self.root,
                                   creationflags=subprocess.CREATE_NO_WINDOW if os.name == "nt" else 0)
        if completed.returncode:
            details = (completed.stderr or completed.stdout or "未知错误")[-1200:]
            raise ValidationError("计划表生成失败：" + details)

    def _context(self):
        sales = self.store.read_sales()
        current = self.store.clock().astimezone(TZ).strftime("%Y-%m")
        products = sales["products"]
        plan_end = horizon(current, self.store.settings()["bridge_m6"])
        plans = {(r["product_id"], r["month"]): dict(r) for r in self.store.db.execute("SELECT * FROM plans")}
        business_months = sorted(set(month_range(current, plan_end))
                        | {month_add(current, offset) for offset in (-3, -2, -1)}
                        | {r["month"] for r in sales["sales"]}
                        | {month for _, month in plans})
        months = month_range(business_months[0][:4] + "-01", max(plan_end, business_months[-1]))
        baseline = [{"product_id": p["id"], "month": month,
                     "baseline_sop": plans.get((p["id"], month), {}).get("sop"),
                     "revision": plans.get((p["id"], month), {}).get("revision", 0)}
                    for p in products for month in months]
        editable_months = [month for month in months if month <= plan_end]
        signature_data = {"months": months, "products": products, "baseline": baseline,
                          "plan_end": plan_end, "editable_months": editable_months,
                          "current_month": current, "business_months": business_months,
                          "schema_version": SCHEMA_VERSION, "layout_version": LAYOUT_VERSION,
                          "layout": _inline_layout(products, months, editable_months, current, plans)}
        signature = _signature(signature_data)
        return {**signature_data, "plans": list(plans.values())}, signature

    def _workbook_product_ids(self):
        current = {p["id"] for p in self.store.read_sales()["products"]}
        aliases = self.store.source_aliases() if hasattr(self.store, "source_aliases") else {}
        # A confirmed alias only allows an old row to be inspected. It never
        # redirects quantities; the complete emptiness check below is required.
        return current | {source for source, target in aliases.items() if target in current}

    def _exclude_empty_bound_sources(self, expected, baseline, accepted, values):
        current = {p["id"] for p in self.store.read_sales()["products"]}
        aliases = self.store.source_aliases() if hasattr(self.store, "source_aliases") else {}
        removed = {p["id"] for p in expected["products"]} - current
        for pid in removed:
            if aliases.get(pid) not in current:
                raise ValidationError("计划表包含已不在产品目录中的版本，已停止导入。")
            keys = [key for key, base in baseline.items() if base["product_id"] == pid]
            if any(baseline[key]["baseline_sop"] is not None or accepted.get(key) is not None or values.get(key) is not None for key in keys):
                raise ValidationError("已关联来源的旧版本行中存在计划值或未同步修改（含0），为避免丢失已停止更新，请先核对处理。")
        return {key: value for key, value in values.items() if baseline[key]["product_id"] not in removed}

    def _parse(self, content):
        import openpyxl
        try:
            workbook = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=False)
        except Exception as error:
            raise ValidationError("计划表无法读取，请保存并关闭 Excel 后重试。") from error
        try:
            if not {"_sync", "_meta"} <= set(workbook.sheetnames):
                raise ValidationError("计划表缺少同步基准，不能自动导入。请使用本看板导出的原始维护表。")
            if workbook["_sync"].sheet_state == "visible" or workbook["_meta"].sheet_state == "visible":
                raise ValidationError("计划表同步基准页已被改动，请恢复隐藏页后重试。")
            metadata = {}
            for row in list(workbook["_meta"].iter_rows(values_only=True))[1:]:
                if row and row[0] is not None:
                    metadata[str(row[0])] = row[1] if len(row) > 1 else None
            token = metadata.get("token")
            if not isinstance(token, str) or not token:
                raise ValidationError("计划表缺少有效同步令牌。")
            schema = str(metadata.get("schema_version"))
            if schema not in {"1", "2", "3", "4", "5"}:
                raise ValidationError("计划表同步版本不受支持。")
            entry = self.store.db.execute("SELECT payload FROM workbook_exports WHERE token=?", (token,)).fetchone()
            if not entry:
                raise ValidationError("计划表不属于此看板的数据版本，不能覆盖当前计划。")
            expected = json.loads(entry["payload"])
            if schema != str(expected.get("schema_version", 1)):
                raise ValidationError("计划表布局版本与原始导出记录不一致。")
            expected_baseline = {_key(b["product_id"], b["month"]): b for b in expected["baseline"]}
            baseline_rows = list(workbook["_sync"].iter_rows(values_only=True))
            if not baseline_rows or tuple(baseline_rows[0][:4]) != ("product_id", "month", "baseline_sop", "revision"):
                raise ValidationError("计划表同步基准列已被改动。")
            baseline = {}
            for row in baseline_rows[1:]:
                if all(value is None for value in row):
                    continue
                if len(row) < 4:
                    raise ValidationError("计划表同步基准不完整。")
                pid, month, value, revision = row[:4]
                key = _key(pid, month)
                if key in baseline or key not in expected_baseline:
                    raise ValidationError("计划表同步基准存在重复或未知产品月份。")
                if isinstance(revision, bool) or not isinstance(revision, (int, float)) or int(revision) != revision:
                    raise ValidationError("同步基准版本号无效。")
                baseline[key] = {"product_id": pid, "month": month, "baseline_sop": _cell_quantity(value), "revision": int(revision)}
                if baseline[key] != expected_baseline[key]:
                    raise ValidationError("计划表隐藏的同步基准被改动，已停止导入。")
            if baseline.keys() != expected_baseline.keys():
                raise ValidationError("计划表同步基准不完整。")
            if schema in {"2", "3", "4", "5"}:
                if schema == "5":
                    values = self._parse_inline(workbook, metadata, expected)
                else:
                    values = self._parse_continuous(workbook, metadata, expected) if schema in {"3", "4"} else self._parse_annual(workbook, metadata, expected)
                accepted_row = self.store.db.execute("SELECT accepted_values FROM workbook_imports WHERE token=?", (token,)).fetchone()
                accepted = json.loads(accepted_row["accepted_values"]) if accepted_row else {key: row["baseline_sop"] for key, row in baseline.items()}
                return token, expected, baseline, accepted, self._exclude_empty_bound_sources(expected, baseline, accepted, values)
            if "SOP计划" not in workbook.sheetnames:
                raise ValidationError("旧版计划表缺少SOP计划页。")
            visible = list(workbook["SOP计划"].iter_rows(values_only=False))
            if not visible:
                raise ValidationError("计划表为空。")
            headers = [cell.value for cell in visible[0]]
            while headers and headers[-1] is None:
                headers.pop()
            if headers != ["product_id", "品牌", "车型", "版本", *expected["months"]]:
                raise ValidationError("计划表列标题被改动，请保留产品列及原始月份列。")
            expected_ids = {p["id"] for p in expected["products"]}
            known_ids = self._workbook_product_ids()
            seen, values = set(), {}
            for cells in visible[1:]:
                if all(cell.value is None for cell in cells):
                    continue
                pid = cells[0].value
                if not isinstance(pid, str) or pid not in expected_ids or pid not in known_ids:
                    raise ValidationError("计划表包含未知产品或产品标识被改动。")
                if pid in seen:
                    raise ValidationError("计划表出现重复产品行，已停止整表导入。")
                seen.add(pid)
                if any(cell.value is not None for cell in cells[len(headers):]):
                    raise ValidationError("计划表存在未识别的附加数据列。")
                for index, month in enumerate(expected["months"], 4):
                    cell = cells[index] if index < len(cells) else None
                    if cell is not None and cell.data_type == "f":
                        raise ValidationError(f"{pid} / {month} 包含公式，请填写非负整数。")
                    values[_key(pid, month)] = _cell_quantity(cell.value if cell is not None else None)
            if seen != expected_ids:
                raise ValidationError("计划表产品行被删除，请保留所有版本行。")
            accepted_row = self.store.db.execute("SELECT accepted_values FROM workbook_imports WHERE token=?", (token,)).fetchone()
            accepted = json.loads(accepted_row["accepted_values"]) if accepted_row else {key: row["baseline_sop"] for key, row in baseline.items()}
            return token, expected, baseline, accepted, self._exclude_empty_bound_sources(expected, baseline, accepted, values)
        finally:
            workbook.close()

    def _parse_annual(self, workbook, metadata, expected):
        import openpyxl
        if metadata.get("layout_version") != expected.get("layout_version") or expected.get("layout_version") != LEGACY_ANNUAL_LAYOUT:
            raise ValidationError("年度计划表布局标识被改动或不受支持。")
        if "_layout" not in workbook.sheetnames or workbook["_layout"].sheet_state == "visible":
            raise ValidationError("年度计划表缺少隐藏的产品月份映射。")
        layout = expected["layout"]
        mapping = list(workbook["_layout"].iter_rows(values_only=True))
        wanted = [("product_id", "month", "sheet", "cell", "editable")]
        wanted += [(c["product_id"], c["month"], c["sheet"], c["cell"], int(c["editable"])) for c in layout["cells"]]
        actual = [tuple(r[:5]) for r in mapping if any(v is not None for v in r)]
        if actual != wanted or any(any(v is not None for v in r[5:]) for r in mapping):
            raise ValidationError("年度计划表隐藏的单元格映射已被改动。")
        known_ids = self._workbook_product_ids()
        expected_headers = ["品牌", "车型", "版本", *[f"{m}月" for m in range(1, 13)], "年度合计", "product_id", "row_kind"]
        grids, values = {}, {}
        for sheet_spec in layout["sheets"]:
            name = sheet_spec["name"]
            if name not in workbook.sheetnames:
                raise ValidationError(f"年度计划表缺少{name}页。")
            rows = list(workbook[name].iter_rows(values_only=False))
            def cell_value(row, col):
                return rows[row-1][col-1].value if row <= len(rows) and col <= len(rows[row-1]) else None
            if [cell_value(3, col) for col in range(1, 19)] != expected_headers:
                raise ValidationError(f"{name}的月份或标识标题已被改动。")
            for row in rows[sheet_spec["last_row"]:]:
                if any(c.value is not None for c in row):
                    raise ValidationError(f"{name}存在模板外的附加行，不能自动导入。")
            if any(any(c.value is not None for c in row[18:]) for row in rows):
                raise ValidationError(f"{name}存在模板外的附加列，不能自动导入。")
            allowed_rows = {r["row"] for r in sheet_spec["rows"]}
            for row_number in range(4, sheet_spec["last_row"] + 1):
                if row_number not in allowed_rows and any(cell_value(row_number, c) is not None for c in range(1, 19)):
                    raise ValidationError(f"{name}的分隔行中存在未识别数据。")
            for spec in sheet_spec["rows"]:
                row_num = spec["row"]
                if [cell_value(row_num, col) for col in (1, 2, 3)] != [spec.get(f"label_{letter}") or None for letter in ("a", "b", "c")]:
                    raise ValidationError(f"{name}的品牌、车型或版本标签已被改动。")
                if cell_value(row_num, 18) != spec["kind"]:
                    raise ValidationError(f"{name}的行标识已被改动。")
                if spec["kind"] == "product":
                    pid = spec["product_id"]
                    if pid not in known_ids or cell_value(row_num, 17) != pid:
                        raise ValidationError(f"{name}的产品标识或版本行已被改动。")
                    shown = set(sheet_spec["shown_months"])
                    for month_number in range(1, 13):
                        month = f"{sheet_spec['year']}-{month_number:02d}"
                        if month not in shown and cell_value(row_num, month_number + 3) is not None:
                            raise ValidationError(f"{name} / {month}为灰色未开放月份，不能填写计划。")
                elif cell_value(row_num, 17) is not None:
                    raise ValidationError(f"{name}的汇总行中存在产品标识。")
            # Total cells stay formulas, never inputs. Tampering is visible rather
            # than silently treating a changed total as a maintained SOP value.
            for address, formula in sheet_spec["formulas"].items():
                formula_row, formula_col = openpyxl.utils.cell.coordinate_to_tuple(address)
                if cell_value(formula_row, formula_col) != formula:
                    raise ValidationError(f"{name} / {address}汇总公式已被改动，请只填写版本月份的白色输入格。")
            grids[name] = rows
        baseline = {_key(b["product_id"], b["month"]): b["baseline_sop"] for b in expected["baseline"]}
        for mapping in layout["cells"]:
            row_number, column = openpyxl.utils.cell.coordinate_to_tuple(mapping["cell"])
            rows = grids[mapping["sheet"]]
            cell = rows[row_number-1][column-1]
            if cell.data_type == "f":
                raise ValidationError(f"{mapping['product_id']} / {mapping['month']} 包含公式，请填写非负整数。")
            key = _key(mapping["product_id"], mapping["month"])
            value = _cell_quantity(cell.value)
            if not mapping["editable"]:
                if value != baseline[key]:
                    raise ValidationError(f"{mapping['month']}为灰色未开放月份，已保留原计划，不能修改。")
                continue
            values[key] = value
        return values

    def _parse_inline(self, workbook, metadata, expected):
        from sop_table import parse_inline_sheet
        if metadata.get("layout_version") != LAYOUT_VERSION or expected.get("layout_version") != LAYOUT_VERSION:
            raise ValidationError("主表布局标识已被改动，请恢复原表后重试。")
        if "_layout" not in workbook.sheetnames or workbook["_layout"].sheet_state == "visible":
            raise ValidationError("计划表缺少隐藏的版本月份映射。")
        mappings = list(workbook["_layout"].iter_rows(values_only=True))
        wanted = [("product_id", "month", "sheet", "cell", "editable", "applicable")]
        wanted += [(c["product_id"], c["month"], c["sheet"], c["cell"], int(c["editable"]), int(c["applicable"])) for c in expected["layout"]["cells"]]
        if [tuple(row[:6]) for row in mappings if any(v is not None for v in row)] != wanted or any(any(v is not None for v in row[6:]) for row in mappings):
            raise ValidationError("隐藏的版本月份基准已被改动。")
        spec = expected["layout"]["sheets"][0]
        if spec["name"] not in workbook.sheetnames:
            raise ValidationError("计划表缺少 SOP计划 工作表。")
        parsed = parse_inline_sheet(workbook[spec["name"]], spec)
        expected["_manual_entries"] = parsed["entries"]
        expected["_catalog_changes"] = parsed["catalog_changes"]
        expected["_catalog_values"] = parsed.get("catalog_values", {})
        expected["_removed_ids"] = parsed.get("removed_ids", [])
        # The parser tolerates formulas Excel adjusts during row insertion. Rebuild
        # moved rows and totals from the committed directory after accepting edits.
        from openpyxl.utils.cell import coordinate_to_tuple
        rows = list(workbook[spec["name"]].iter_rows(values_only=True))
        def cell_at(row, column):
            return rows[row-1][column-1] if row <= len(rows) and column <= len(rows[row-1]) else None
        expected["_structure_changed"] = any(cell_at(row["row"], spec["key_column"]) != row["key"] for row in spec["rows"])
        expected["_structure_changed"] |= any(cell_at(*coordinate_to_tuple(address)) != formula for address, formula in spec["formulas"].items())
        return {_key(item["product_id"], item["month"]): item["sop"] for item in parsed["values"]}

    def _parse_continuous(self, workbook, metadata, expected):
        import openpyxl
        supported_layout = PREVIOUS_CONTINUOUS_LAYOUT if expected.get("schema_version") == 4 else LEGACY_CONTINUOUS_LAYOUT
        if metadata.get("layout_version") != expected.get("layout_version") or expected.get("layout_version") != supported_layout:
            raise ValidationError("跨年计划表布局标识被改动或不受支持。")
        if "_layout" not in workbook.sheetnames or workbook["_layout"].sheet_state == "visible":
            raise ValidationError("跨年计划表缺少隐藏的产品月份映射。")
        layout = expected["layout"]
        mapping_rows = list(workbook["_layout"].iter_rows(values_only=True))
        wanted = [("product_id", "month", "sheet", "cell", "editable", "applicable")]
        wanted += [(c["product_id"], c["month"], c["sheet"], c["cell"], int(c["editable"]), int(c["applicable"])) for c in layout["cells"]]
        if [tuple(r[:6]) for r in mapping_rows if any(v is not None for v in r)] != wanted or any(any(v is not None for v in r[6:]) for r in mapping_rows):
            raise ValidationError("跨年计划表隐藏的单元格映射已被改动。")
        known_ids = self._workbook_product_ids()
        baseline = {_key(b["product_id"], b["month"]): b["baseline_sop"] for b in expected["baseline"]}
        grids, values, manual_entries = {}, {}, []
        for spec in layout["sheets"]:
            name = spec["name"]
            if name not in workbook.sheetnames:
                raise ValidationError(f"跨年计划表缺少{name}页。")
            rows = list(workbook[name].iter_rows(values_only=False))
            def cell_value(row, col):
                return rows[row-1][col-1].value if row <= len(rows) and col <= len(rows[row-1]) else None
            width = spec["last_column"]
            for offset, header in enumerate(spec["headers"]):
                if [cell_value(spec["header_row"] + offset, col) for col in range(1, width + 1)] != header:
                    raise ValidationError("跨年计划表的年份、月份或标识标题已被改动。")
            if any(any(c.value is not None for c in row) for row in rows[spec["last_row"]:]):
                raise ValidationError("跨年计划表存在模板外的附加行，不能自动导入。")
            if any(any(c.value is not None for c in row[width:]) for row in rows):
                raise ValidationError("跨年计划表存在模板外的附加列，不能自动导入。")
            allowed = {r["row"] for r in spec["rows"]}
            manual_section = spec.get("manual_section") if expected.get("schema_version") == 4 else None
            if manual_section:
                title_row = manual_section["title_row"]
                if cell_value(title_row, 1) != manual_section["title"] or any(cell_value(title_row, c) is not None for c in range(2, width + 1)):
                    raise ValidationError("新增车型/版本区域标题已被改动。")
                for offset, header in enumerate(spec["headers"]):
                    if [cell_value(manual_section["header_row"] + offset, col) for col in range(1, width + 1)] != header:
                        raise ValidationError("新增车型/版本区域的年份或月份标题已被改动。")
                allowed.update(range(title_row, manual_section["first_row"]))
            for row_num in range(spec["first_data_row"], spec["last_row"] + 1):
                if row_num not in allowed and any(cell_value(row_num, c) is not None for c in range(1, width + 1)):
                    raise ValidationError("跨年计划表的分隔行中存在未识别数据。")
            for item in spec["rows"]:
                row_num = item["row"]
                is_manual = item["kind"] == "manual_entry" and manual_section is not None
                if not is_manual and [cell_value(row_num, col) for col in (1, 2, 3)] != [item.get(f"label_{letter}") or None for letter in ("a", "b", "c")]:
                    raise ValidationError("跨年计划表的品牌、车型或版本标签已被改动，请在看板中维护目录。")
                if cell_value(row_num, spec["kind_column"]) != item["kind"]:
                    raise ValidationError("跨年计划表的行标识已被改动。")
                pid = cell_value(row_num, spec["product_column"])
                if item["kind"] == "product":
                    if pid != item["product_id"] or pid not in known_ids:
                        raise ValidationError("跨年计划表的产品标识已被改动。")
                elif pid is not None:
                    raise ValidationError("跨年计划表的汇总行中存在产品标识。")
                if is_manual:
                    labels = [cell_value(row_num, col) for col in (1, 2, 3)]
                    quantities = []
                    for column in spec["columns"]:
                        if column["kind"] != "month":
                            continue
                        cell = rows[row_num-1][column["column"]-1]
                        if cell.data_type == "f":
                            raise ValidationError(f"新增区域第{row_num}行的计划量不能填写公式。")
                        value = _cell_quantity(cell.value)
                        if value is not None:
                            quantities.append({"month": column["month"], "sop": value})
                    if not quantities and all(value is None or isinstance(value, str) and not value.strip() for value in labels):
                        continue
                    if any(not isinstance(value, str) or not value.strip() or rows[row_num-1][col].data_type == "f" for col, value in enumerate(labels)):
                        raise ValidationError(f"新增区域第{row_num}行请完整填写品牌、车型和版本名称。")
                    manual_entries.append(dict(zip(("brand", "model", "version"), labels), plans=quantities))
            for address, formula in spec["formulas"].items():
                r, c = openpyxl.utils.cell.coordinate_to_tuple(address)
                if cell_value(r, c) != formula:
                    raise ValidationError(f"{address}汇总公式已被改动，请只填写版本月份输入格。")
            grids[name] = rows
        for mapping in layout["cells"]:
            r, c = openpyxl.utils.cell.coordinate_to_tuple(mapping["cell"])
            cell = grids[mapping["sheet"]][r-1][c-1]
            if cell.data_type == "f":
                raise ValidationError(f"{mapping['product_id']} / {mapping['month']} 包含公式，请填写非负整数。")
            key = _key(mapping["product_id"], mapping["month"])
            not_applicable = not mapping["applicable"] and baseline[key] is None
            if not_applicable:
                if cell.value != "未适用":
                    raise ValidationError(f"{mapping['month']}不在该版本适用期内，请保留“未适用”标记。")
                value = None
            else:
                value = _cell_quantity(cell.value)
            if not mapping["editable"]:
                if value != baseline[key]:
                    raise ValidationError(f"{mapping['month']}为未开放月份或停用版本的当前及未来月份，原计划只读保留，不能修改。")
                continue
            values[key] = value
        expected["_manual_entries"] = manual_entries
        return values

    def _import(self, content, digest, choice=None):
        token, expected, baseline, accepted, values = self._parse(content)
        if self._export_signature is None:
            self._export_signature = _signature(expected)
        current = {(r["product_id"], r["month"]): dict(r) for r in self.store.db.execute("SELECT * FROM plans")}
        changes, conflicts, next_accepted = [], [], dict(accepted)
        current_month = self.store.clock().astimezone(TZ).strftime("%Y-%m")
        plan_end = horizon(current_month, self.store.settings()["bridge_m6"])
        for key, value in values.items():
            if value == accepted[key]:
                continue
            base = baseline[key]
            pid, month = base["product_id"], base["month"]
            if month > plan_end:
                raise ValidationError(f"{month}已超出当前开放计划范围，不能导入修改；请恢复该格后重试，原计划会保留。")
            row = current.get((pid, month), {})
            database_value = row.get("sop")
            conflicting = database_value != accepted[key] and value != database_value
            if conflicting:
                conflicts.append({"product_id": pid, "month": month, "baseline": accepted[key], "workbook": value, "database": database_value})
                if choice is None:
                    continue
                if choice == "db":
                    # This user-selected discard is applied by the subsequent regenerated workbook.
                    next_accepted[key] = value
                    continue
            if value != database_value:
                changes.append({"product_id": pid, "month": month, "sop": value, "forecast": row.get("forecast"),
                                "revision": row.get("revision", 0), "reason": "表格同步", "notes": row.get("notes", "")})
            next_accepted[key] = value
        if conflicts and choice is None:
            return self._set_status("conflict", "表格和网页同时修改了相同的 SOP 单元格，请选择保留哪一方。", conflicts)
        entries = expected.get("_manual_entries", [])
        catalog_changes = expected.get("_catalog_changes", [])
        # The managed file can stay open through several saves. Compare its names
        # with the last accepted save, just as monthly quantities already do.
        prior_catalog = accepted.get("_catalog", {})
        observed_catalog = expected.get("_catalog_values", {})
        if observed_catalog:
            catalog_changes = []
            original_catalog = {p["id"]: p for p in expected["products"]}
            revision_map = expected["layout"]["sheets"][0].get("plan_revisions_by_product", {})
            accepted_revisions = accepted.get("_catalog_plan_revisions", {})
            for pid, observed in observed_catalog.items():
                previous = prior_catalog.get(pid, original_catalog[pid])
                if pid in expected.get("_removed_ids", []):
                    # A missing row carries no replacement label. Preserve the
                    # name accepted from any earlier save of this open workbook.
                    observed.update({field: previous[field] for field in ("brand", "model", "version")})
                changed = {field: observed[field] for field in ("brand", "model", "version", "active")
                           if field in observed and observed[field] != previous.get(field, True if field == "active" else None)}
                if changed:
                    item = {"id": pid, "baseline": previous, **changed}
                    if changed.get("active") is False:
                        item["baseline_plan_revisions"] = accepted_revisions.get(pid, revision_map.get(pid, {}))
                    catalog_changes.append(item)
        previous_manual = accepted.get("_manual", {})
        previous_manual_records = accepted.get("_manual_records", {})
        if expected.get("schema_version") == 5 and previous_manual_records:
            observed_names = {json.dumps(catalog_name_key(entry), ensure_ascii=False) for entry in entries}
            unresolved = set(previous_manual_records) - observed_names
            if unresolved:
                names_to_restore = "、".join(" / ".join(json.loads(name)) for name in sorted(unresolved))
                raise ValidationError("新增版本的编号尚未写回表格。请先保留这些已登记名称并保存、关闭 Excel，完成同步后再改名或删行：" + names_to_restore)
        next_manual = dict(previous_manual)
        names = {}
        for product in expected["products"]:
            names.setdefault(catalog_name_key(product), []).append(product["id"])
        for entry in entries:
            name_key = catalog_name_key(entry)
            named = json.dumps(name_key, ensure_ascii=False)
            prior = previous_manual.get(named, {})
            supplied = {p["month"]: p["sop"] for p in entry["plans"]}
            # A saved addition may remain on disk while Excel holds the file.
            # Keep its own accepted baseline, including subsequent clears.
            all_months = set(supplied) | set(prior)
            matching = names.get(name_key, [])
            entry["plans"] = [{"month": month, "sop": supplied.get(month),
                "baseline_sop": prior.get(month, accepted.get(_key(matching[0], month)) if len(matching) == 1 else None)}
                for month in sorted(all_months)]
            next_manual[named] = {month: supplied.get(month) for month in all_months}
        imported = len(changes)
        created = 0
        catalog_updated = deactivated = 0
        if expected.get("schema_version") == 5 and (entries or catalog_changes) and not self._structural_file_ready():
            return self._set_status("pending", "本次包含新增、改名或删行。请保存并关闭 Excel，整次修改随后自动同步；历史数据会保留。")
        if changes or entries or catalog_changes:
            try:
                result = self.store.import_workbook_plans({"changes": changes, "entries": entries,
                    "catalog_changes": catalog_changes,
                    "choice": choice, "owner": "Excel计划表", "reason": "表格同步"})
            except ConflictError as error:
                return self._set_status("conflict", "表格和网页同时修改了相同的 SOP 单元格，请选择保留哪一方。", error.conflicts)
            imported, created = result["saved_count"], result["created_count"]
            catalog_updated, deactivated = result.get("catalog_updated_count", 0), result.get("deactivated_count", 0)
            if entries and expected.get("schema_version") == 5:
                original_ids = {p["id"] for p in expected["products"]}
                records = dict(previous_manual_records)
                for entry, product in zip(entries, result.get("products", [])):
                    if product["id"] not in original_ids:
                        records[json.dumps(catalog_name_key(entry), ensure_ascii=False)] = product["id"]
                if records:
                    next_accepted["_manual_records"] = records
        if observed_catalog:
            next_accepted["_catalog"] = observed_catalog
            # Only acknowledge revisions written by this Excel save. Unseen web
            # revisions must still conflict if a later saved row deletes a product.
            acknowledged = {pid: dict(accepted_revisions.get(pid, revision_map.get(pid, {}))) for pid in observed_catalog}
            for change in changes:
                pid, month = change["product_id"], change["month"]
                row = self.store.db.execute("SELECT revision FROM plans WHERE product_id=? AND month=?", (pid, month)).fetchone()
                if row and pid in acknowledged:
                    acknowledged[pid][month] = row["revision"]
            next_accepted["_catalog_plan_revisions"] = acknowledged
        if next_manual:
            next_accepted["_manual"] = next_manual
        # Track accepted values while Excel holds the file open and prevents baseline writeback.
        self.store.db.execute("INSERT OR REPLACE INTO workbook_imports(token,accepted_values,last_digest,imported_at) VALUES (?,?,?,?)", (
            token, json.dumps(next_accepted, ensure_ascii=False), digest, stamp(self.store.clock())))
        self._last_digest = digest
        if imported or created or catalog_updated:
            self._state["last_sync"] = stamp(self.store.clock())
        if expected.get("_structure_changed"):
            self._export_signature = None
        return {"imported": imported, "created": created, "catalog_updated": catalog_updated, "deactivated": deactivated, "token": token}

    def _structural_file_ready(self):
        """Probe delete-sharing only; never delete or alter the user's workbook.

        New rows need their stable IDs written back immediately. While Excel holds
        the file, defer the whole structural save so subsequent edits stay local.
        """
        if os.name != "nt":
            return True
        import ctypes
        from ctypes import wintypes
        kernel = ctypes.WinDLL("kernel32", use_last_error=True)
        create = kernel.CreateFileW
        create.argtypes = (wintypes.LPCWSTR, wintypes.DWORD, wintypes.DWORD, wintypes.LPVOID,
                           wintypes.DWORD, wintypes.DWORD, wintypes.HANDLE)
        create.restype = wintypes.HANDLE
        handle = create(str(self.path), 0x00010000, 0x00000007, None, 3, 0x80, None)
        if handle == ctypes.c_void_p(-1).value:
            error = ctypes.get_last_error()
            if error in (32, 33):
                return False
            raise ValidationError("无法更新计划表，请确认文件可写且未被占用：" + ctypes.FormatError(error))
        close = kernel.CloseHandle
        close.argtypes, close.restype = (wintypes.HANDLE,), wintypes.BOOL
        close(handle)
        return True

    def _export(self, context, signature, observed_digest):
        token = str(uuid.uuid4())
        generated = stamp(self.store.clock())
        pending_path = self.path.with_name(f".{self.path.stem}.{token}.pending.xlsx")
        payload = {**context, "token": token, "generated_at": generated, "output_path": str(pending_path)}
        try:
            self.writer(payload)
            if not pending_path.is_file():
                raise ValidationError("计划表生成未返回有效文件。")
            # A user may have saved Excel while the writer was working; never replace that newer edit.
            latest_digest = hashlib.sha256(self.path.read_bytes()).hexdigest() if self.path.exists() else None
            if latest_digest != observed_digest:
                self._input_checked = False
                return self._set_status("pending", "生成期间检测到新的表格保存，下一次同步将先读取这些修改。")
            self.store.db.execute("INSERT INTO workbook_exports(token,created_at,payload) VALUES (?,?,?)", (
                token, generated, json.dumps(context, ensure_ascii=False)))
            os.replace(pending_path, self.path)
            self._last_digest = hashlib.sha256(self.path.read_bytes()).hexdigest()
            stat = self.path.stat()
            self._last_file_stat = (stat.st_mtime_ns, stat.st_size)
            self._export_signature = signature
            return self._set_status("synced", "网页与计划表已同步。", synced=True)
        except PermissionError:
            return self._set_status("pending", "计划已保存在看板中；请保存并关闭 Excel，计划表将自动更新。")
        finally:
            if pending_path.exists():
                try:
                    pending_path.unlink()
                except OSError:
                    pass

    def sync(self, force_export=False, choice=None):
        if choice not in {None, "workbook", "db"}:
            raise ValidationError("请选择保留 workbook 或 db。")
        with self.lock, self.store.lock:
            self._input_checked = False
            try:
                context, signature = self._context()
                observed_digest = None
                if self.path.exists():
                    stat = self.path.stat()
                    observed_stat = (stat.st_mtime_ns, stat.st_size)
                    if observed_stat != self._last_file_stat or self._last_digest is None or choice is not None:
                        content = self.path.read_bytes()
                        observed_digest = hashlib.sha256(content).hexdigest()
                        if observed_digest != self._last_digest or choice is not None:
                            result = self._import(content, observed_digest, choice)
                            if result.get("status") in {"conflict", "pending"}:
                                return result
                            context, signature = self._context()
                        self._last_file_stat = observed_stat
                    else:
                        observed_digest = self._last_digest
                elif self._last_digest is not None:
                    return self._set_status("error", "维护表已被删除或移动，请放回原路径后重试。")
                self._input_checked = True
                if force_export or signature != self._export_signature or not self.path.exists() or self._state["status"] == "pending":
                    return self._export(context, signature, observed_digest)
                return self._set_status("synced", "网页与计划表已同步。")
            except PermissionError:
                return self._set_status("pending", "计划表暂被 Excel 占用，请保存并关闭后重试。")
            except Exception as error:
                return self._set_status("error", str(error))

    def save_web_plans(self, body):
        with self.lock:
            before = self.sync()
            if before["status"] == "conflict":
                raise ConflictError(before["conflicts"])
            if before["status"] in {"error", "idle"}:
                raise ValidationError(before["message"])
            if not self._input_checked:
                raise ValidationError("尚未读取计划表的已保存修改，请保存并关闭 Excel 后再保存网页计划。")
            with self.store.lock:
                result = self.store.save_plans(body)
            after = self.sync()
            return {**result, "workbook_sync": after}
