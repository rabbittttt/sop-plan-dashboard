"""Trusted, cell-scoped round trips for downloaded SOP workbooks.

Export baselines are immutable. An unchanged old copy never undoes a newer web
edit, and a changed cell is merged only against its own original export value.
Downloads may be authored by the bundled Node renderer or the offline Python writer.
"""
from __future__ import annotations

import io
import json
import re
import uuid
import zipfile
import xml.etree.ElementTree as ET
from contextlib import contextmanager, nullcontext
from pathlib import Path

from sop_export import validate_export
from sop_plan_scope import is_model_plan, plan_total
from sop_store import ConflictError, TZ, ValidationError, catalog_name_key, horizon, stamp, text_field
from sop_workbook import _cell_quantity
from sop_table import parse_inline_sheet

MAX_UPLOAD_BYTES = 12 * 1024 * 1024
MAX_EXPANDED_BYTES = 96 * 1024 * 1024
MAX_SHEET_CELLS = 1_500_000
IMPORT_SHEET = "_sop_import"
PLAN_SHEET = "SOP计划"
ADDRESS = re.compile(r"^([A-Z]{1,3})([1-9][0-9]{0,5})$")


def _key(pid, month):
    return json.dumps([pid, month], ensure_ascii=False, separators=(",", ":"))


def _sum_known(values):
    known = [v for v in values if v is not None]
    return sum(known) if known else None


def _value(cell):
    value = cell.value
    if value is None or value == "":
        return None
    if not isinstance(value, (str, int, float, bool)):
        # Dates and other non-plan cell types are never accepted as quantities.
        return [cell.data_type, str(value)]
    return [cell.data_type, value]


def _raw(cells, address):
    stored = cells.get(address)
    return stored[1] if stored is not None else None


def _xml_dimensions(raw):
    """Read actual cell coordinates, rather than trusting optional dimension tags."""
    last_row, max_row, max_column, count, coordinates = 0, 0, 0, 0, set()
    for event, element in ET.iterparse(io.BytesIO(raw), events=("start", "end")):
        tag = element.tag.rsplit("}", 1)[-1]
        if event == "start" and tag == "row":
            row = element.attrib.get("r", "")
            if not row.isdigit() or not last_row < int(row) <= 400000:
                raise ValidationError("Excel 行编号无效、重复或超出允许范围。")
            last_row = int(row)
            max_row = max(max_row, last_row)
        if event == "end" and tag == "c":
            address = element.attrib.get("r", "")
            matched = ADDRESS.fullmatch(address)
            if not matched or address in coordinates or int(matched.group(2)) != last_row:
                raise ValidationError("Excel 单元格编号无效或重复。")
            coordinates.add(address)
            column = 0
            for char in matched.group(1):
                column = column * 26 + ord(char) - ord("A") + 1
            max_column = max(max_column, column)
            count += 1
            if max_column > 200 or count > 2_000_000:
                raise ValidationError("Excel 工作表范围过大，请缩小导出范围。")
        if event == "end":
            element.clear()
    return max(1, max_row), max(1, max_column)


@contextmanager
def _read_workbook(content):
    """Bound the ZIP and worksheet dimensions before walking untrusted cells."""
    import openpyxl

    if not isinstance(content, bytes) or not 0 < len(content) <= MAX_UPLOAD_BYTES:
        raise ValidationError("请选择不超过 12 MB 的 .xlsx 计划表。")
    try:
        dimensions = {}
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            items = archive.infolist()
            names = [item.filename for item in items]
            if len(items) > 2000 or len(names) != len(set(names)):
                raise ValidationError("Excel 文件包含重复或过多的内部文件。")
            if sum(item.file_size for item in items) > MAX_EXPANDED_BYTES:
                raise ValidationError("Excel 文件解压后过大，请缩小导出范围。")
            if not {"[Content_Types].xml", "xl/workbook.xml"}.issubset(names):
                raise ValidationError("文件不是有效的 .xlsx 工作簿。")
            for item in items:
                if item.file_size > 32 * 1024 * 1024 or item.flag_bits & 1:
                    raise ValidationError("Excel 文件过大或受到加密保护，无法导入。")
                name = item.filename.lower()
                if "vbaproject" in name or name.startswith("xl/externallinks/"):
                    raise ValidationError("请使用本看板导出的 .xlsx 文件，不支持宏或外部工作簿链接。")
                if name.endswith(".xml"):
                    raw = archive.read(item)
                    if b"<!DOCTYPE" in raw.upper() or b"<!ENTITY" in raw.upper():
                        raise ValidationError("Excel 文件包含不支持的 XML 内容。")
                    if name.startswith("xl/worksheets/"):
                        dimensions[item.filename.lstrip("/")] = _xml_dimensions(raw)
        workbook = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=False, keep_links=False)
    except ValidationError:
        raise
    except Exception as error:
        raise ValidationError("Excel 无法读取，请保存为 .xlsx 后重试。") from error
    try:
        if len(workbook.sheetnames) > 10:
            raise ValidationError("Excel 的工作表数量与导出模板不一致。")
        for sheet in workbook:
            # The artifact writer omits dimension tags; Excel can also leave stale
            # ones. The ZIP preflight already verified the actual coordinates.
            actual_rows, actual_columns = dimensions.get(sheet._worksheet_path.lstrip("/"), (0, 0))
            if not actual_rows or not actual_columns:
                raise ValidationError("Excel 工作表内容不完整。")
            sheet._max_row, sheet._max_column = actual_rows, actual_columns
            row_limit, column_limit, cell_limit = (400000, 5, 2_000_000) if sheet.title == IMPORT_SHEET else (15000, 200, MAX_SHEET_CELLS)
            if (sheet.max_row or 0) > row_limit or (sheet.max_column or 0) > column_limit or (sheet.max_row or 0) * (sheet.max_column or 0) > cell_limit:
                raise ValidationError("Excel 工作表范围过大或已改变，请重新导出。")
        yield workbook
    finally:
        workbook.close()


def _snapshot(workbook):
    result = {}
    for sheet in workbook:
        cells = {}
        for row in sheet.iter_rows():
            for cell in row:
                value = _value(cell)
                if value is not None:
                    cells[cell.coordinate] = value
        result[sheet.title] = {"state": sheet.sheet_state, "cells": cells,
                               "max_row": sheet.max_row, "max_column": sheet.max_column}
    return result


def _metadata(workbook):
    if IMPORT_SHEET not in workbook.sheetnames:
        raise ValidationError("这个文件没有可回填的导出标识。请在看板重新导出并勾选“包含版本明细”，修改后再导入。")
    sheet = workbook[IMPORT_SHEET]
    if sheet.sheet_state == "visible":
        raise ValidationError("导出标识页已被改动，请恢复原文件或重新导出。")
    schema = sheet.cell(1, 2).value
    if sheet.cell(1, 1).value != "SOP_IMPORT" or schema not in (1, 2):
        raise ValidationError("不支持该导出格式，请在本看板重新导出。")
    if (sheet.cell(2, 1).value, sheet.cell(3, 1).value, sheet.cell(4, 1).value) != ("token", "include_versions", "sheet"):
        raise ValidationError("导出标识不完整，请重新导出。")
    if sheet.cell(4, 2).value != PLAN_SHEET or sheet.max_column > 5:
        raise ValidationError("导出表格结构不正确，请重新导出。")
    token, versions = sheet.cell(2, 2).value, sheet.cell(3, 2).value
    if not isinstance(token, str) or not re.fullmatch(r"[0-9a-f]{32}", token) or type(versions) is not bool:
        raise ValidationError("导出标识无效，请重新导出。")
    if schema == 2:
        if sheet.cell(5, 1).value != "manifest_json":
            raise ValidationError("主表行映射缺失，请重新导出。")
        chunks = []
        for row in sheet.iter_rows(min_row=6, max_col=5, values_only=True):
            if all(value is None for value in row):
                continue
            if row[0] != len(chunks) + 1 or not isinstance(row[1], str) or len(row[1]) > 30000 or any(value is not None for value in row[2:]):
                raise ValidationError("主表行映射分块不完整，请重新导出。")
            chunks.append(row[1])
        try:
            manifest = json.loads("".join(chunks))
        except (TypeError, ValueError) as error:
            raise ValidationError("主表行映射无法读取，请重新导出。") from error
        if not isinstance(manifest, dict):
            raise ValidationError("主表行映射格式无效。")
        return token, versions, manifest
    if [sheet.cell(5, col).value for col in range(1, 6)] != ["kind", "product_id", "month", "cell", "baseline_sop"]:
        raise ValidationError("导出单元格映射不完整，请重新导出。")
    mappings = []
    for row in sheet.iter_rows(min_row=6, max_col=5, values_only=True):
        if not any(value is not None for value in row):
            continue
        kind, pid, month, cell, baseline = row
        if kind not in {"plan", "manual"} or not isinstance(cell, str) or not ADDRESS.fullmatch(cell):
            raise ValidationError("导出单元格映射无效，请重新导出。")
        mappings.append({"kind": kind, "product_id": pid, "month": month, "cell": cell, "baseline_sop": baseline})
    return token, versions, mappings


class PlanRoundtrip:
    def __init__(self, store, workbook_sync=None):
        self.store = store
        self.workbook_sync = workbook_sync
        with store.lock:
            store.db.executescript("""
                CREATE TABLE IF NOT EXISTS plan_download_exports (
                    token TEXT PRIMARY KEY, created_at TEXT NOT NULL, payload TEXT NOT NULL);
                CREATE TABLE IF NOT EXISTS plan_download_entry_imports (
                    token TEXT PRIMARY KEY, payload TEXT NOT NULL);
                CREATE TRIGGER IF NOT EXISTS plan_download_no_update BEFORE UPDATE ON plan_download_exports
                    BEGIN SELECT RAISE(ABORT, 'Download export baselines are immutable'); END;
                CREATE TRIGGER IF NOT EXISTS plan_download_no_delete BEFORE DELETE ON plan_download_exports
                    BEGIN SELECT RAISE(ABORT, 'Download export baselines are immutable'); END;
            """)

    def _lock(self):
        return self.workbook_sync.lock if self.workbook_sync is not None else nullcontext()

    def _sync_before(self):
        if self.workbook_sync is None:
            return
        result = self.workbook_sync.sync()
        # A missing primary workbook has no saved edits to read. Its automatic
        # rebuild may fail or a user may have removed it after the last sync;
        # downloaded exports and their later imports can still use the database.
        if not self.workbook_sync.path.exists():
            return
        if result.get("status") == "conflict":
            error = ConflictError(result.get("conflicts", []))
            error.source = "maintenance"
            raise error
        if result.get("status") in {"error", "idle"} or not self.workbook_sync._input_checked:
            raise ValidationError(result.get("message") or "维护表的已保存修改尚未读取，请先完成同步。")

    def prepare(self, body):
        if not isinstance(body, dict):
            raise ValidationError("导出请求格式不正确。")
        source = dict(body)
        requested_rules = source.pop("generation_rule_signature", None)
        requested_mode = source.pop("generation_mode", None)
        ids = source.pop("row_product_ids", None)
        selected_ids = source.pop("selected_product_ids", None)
        clean = validate_export(source)
        if not isinstance(ids, list) or len(ids) != len(clean["rows"]):
            raise ValidationError("导出行缺少稳定版本编码，请刷新看板后重试。")
        if (not isinstance(selected_ids, list) or not selected_ids or len(selected_ids) > 5000
                or any(not isinstance(pid, str) for pid in selected_ids) or len(set(selected_ids)) != len(selected_ids)):
            raise ValidationError("导出范围缺少版本编码或包含重复版本，请刷新看板后重试。")
        with self._lock():
            self._sync_before()
            with self.store.lock:
                sales = self.store.read_sales()
                view = sales.get("generation_view", {})
                if ((requested_rules is not None and requested_rules != view.get("rule_signature"))
                        or (requested_mode is not None and requested_mode != view.get("mode"))):
                    raise ValidationError("代际规则或展示模式已变化，请刷新看板后重新导出。")
                products = {p["id"]: p for p in sales["products"]}
                if any(pid not in products for pid in selected_ids):
                    raise ValidationError("车型版本目录已变化，请刷新看板后重新导出。")
                selected = {pid: products[pid] for pid in selected_ids}
                groups = {}
                for pid, product in selected.items():
                    groups.setdefault((product["brand"], product["model"]), []).append(pid)
                plans = {(r["product_id"], r["month"]): dict(r) for r in self.store.view_plans(sales)}
                current = self.store.clock().astimezone(TZ).strftime("%Y-%m")
                end = horizon(current, self.store.settings()["bridge_m6"])
                clean["current_month"] = current
                metric_col = 4 if clean["include_versions"] else 3
                metrics = ["SOP计划量"] + (["锁单数"] if clean["include_locks"] else [])
                seen_summary, seen_details = set(), set()
                counts = []
                for index, (row, pid) in enumerate(zip(clean["rows"], ids)):
                    group = (row[1], row[2])
                    metric = row[metric_col]
                    if group not in groups:
                        raise ValidationError("导出车型与已选版本范围不一致，请刷新后重试。")
                    if row[0] == "车型汇总":
                        if pid is not None or (clean["include_versions"] and row[3] != "") or (group, metric) in seen_summary:
                            raise ValidationError("导出车型汇总重复或版本编码不正确。")
                        seen_summary.add((group, metric))
                        members = groups[group]
                    else:
                        if (not isinstance(pid, str) or pid not in selected or (pid, metric) in seen_details
                                or tuple(row[1:4]) != tuple(selected[pid][k] for k in ("brand", "model", "version"))):
                            raise ValidationError("导出版本编码与名称不一致或出现重复，请刷新后重试。")
                        seen_details.add((pid, metric))
                        members = [pid]
                    row_counts = {}
                    if metric == "SOP计划量":
                        for offset, month in enumerate(clean["months"]):
                            row[metric_col + 1 + offset] = plan_total([selected[p] for p in members], month, plans)
                        for year in clean["annual_plans"][index]:
                            values = [plan_total([selected[p] for p in members], f"{year}-{m:02d}", plans) for m in range(1, 13)]
                            clean["annual_plans"][index][year] = _sum_known(values)
                            row_counts[year] = sum(v is not None for v in values)
                    else:
                        row_counts = {year: 0 for year in clean["annual_plans"][index]}
                    counts.append(row_counts)
                if seen_summary != {(group, metric) for group in groups for metric in metrics}:
                    raise ValidationError("导出车型汇总不完整，请刷新后重试。")
                if clean["include_versions"] and seen_details != {
                        (pid, metric) for pid in selected for metric in metrics
                        if metric != "锁单数" or not is_model_plan(selected[pid])}:
                    raise ValidationError("导出版本明细不完整，请刷新后重试。")
                editable, baseline = {}, {}
                for pid, product in selected.items():
                    allowed = []
                    for month in clean["months"]:
                        row = plans.get((pid, month), {})
                        applicable = (not product.get("start_month") or month >= product["start_month"]) and (not product.get("end_month") or month <= product["end_month"])
                        correction = month < current and (row.get("sop") is not None or row.get("forecast") is not None)
                        if (not product.get("read_only") and month <= end and (product.get("active", True) or month < current) and (applicable or correction)):
                            allowed.append(month)
                            baseline[_key(pid, month)] = {"product_id": pid, "month": month,
                                                         "baseline_sop": row.get("sop"), "revision": row.get("revision", 0)}
                    editable[pid] = allowed
                context = {"token": uuid.uuid4().hex, "row_product_ids": ids, "inline_schema": 2,
                           "include_versions": clean["include_versions"], "months": clean["months"],
                           "selected_product_ids": selected_ids, "products": list(selected.values()),
                           "editable_months_by_product": editable,
                           "manual_editable_months": [m for m in clean["months"] if m <= end],
                           "annual_known_counts": counts, "baseline": baseline,
                           "plan_values_by_product": {pid: {month: plans.get((pid, month), {}).get("sop") for month in clean["months"]} for pid in selected},
                           "plan_revisions_by_product": {pid: {month: row["revision"] for (p, month), row in plans.items() if p == pid} for pid in selected}}
                return clean, context

    def register(self, path, roundtrip):
        content = Path(path).read_bytes()
        with _read_workbook(content) as workbook:
            token, versions, mappings = _metadata(workbook)
            if token != roundtrip.get("token") or versions != roundtrip.get("include_versions"):
                raise ValidationError("导出同步标识与生成请求不一致。")
            if set(workbook.sheetnames) != {PLAN_SHEET, IMPORT_SHEET}:
                raise ValidationError("导出表格的工作表与模板不一致。")
            snapshot = _snapshot(workbook)
            if isinstance(mappings, dict):
                context = self._register_inline(workbook, mappings, roundtrip, snapshot)
                with self.store.lock:
                    self.store.db.execute("INSERT INTO plan_download_exports VALUES (?,?,?)",
                                          (token, stamp(self.store.clock()), json.dumps(context, ensure_ascii=False, allow_nan=False)))
                return {"ok": True, "token": token}
            sheet_cells = snapshot[PLAN_SHEET]["cells"]
            expected = roundtrip["baseline"] if versions else {}
            seen, addresses, manual_rows = set(), set(), {}
            for item in mappings:
                address = item["cell"]
                if address in addresses:
                    raise ValidationError("导出映射存在重复单元格。")
                addresses.add(address)
                if item["kind"] == "plan":
                    key = _key(item["product_id"], item["month"])
                    if key not in expected or key in seen or item["baseline_sop"] != expected[key]["baseline_sop"]:
                        raise ValidationError("导出计划基准与服务器记录不一致。")
                    if _cell_quantity(_raw(sheet_cells, address)) != expected[key]["baseline_sop"]:
                        raise ValidationError("导出计划量与服务器记录不一致。")
                    seen.add(key)
                else:
                    if not versions or item["product_id"] not in (None, "") or item["baseline_sop"] is not None or _raw(sheet_cells, address) not in (None, ""):
                        raise ValidationError("导出新增区域不正确。")
                    row = int(ADDRESS.fullmatch(address).group(2))
                    if item["month"] in manual_rows.setdefault(row, {}):
                        raise ValidationError("导出新增区域存在重复字段。")
                    manual_rows[row][item["month"]] = address
            if seen != set(expected):
                raise ValidationError("导出版本计划的单元格映射不完整。")
            expected_manual = {"brand", "model", "version", *roundtrip["manual_editable_months"]}
            if versions and (len(manual_rows) != 10 or any(set(fields) != expected_manual for fields in manual_rows.values())):
                raise ValidationError("导出新增车型区域与模板不一致。")
            context = {**roundtrip, "mappings": mappings, "manual_rows": manual_rows, "snapshot": snapshot}
        with self.store.lock:
            self.store.db.execute("INSERT INTO plan_download_exports VALUES (?,?,?)",
                                  (token, stamp(self.store.clock()), json.dumps(context, ensure_ascii=False, allow_nan=False)))
        return {"ok": True, "token": token}

    def _register_inline(self, workbook, manifest, roundtrip, snapshot):
        """Only a freshly server-rendered manifest can become a trusted baseline."""
        for field in ("products", "editable_months_by_product", "manual_editable_months", "plan_revisions_by_product"):
            if manifest.get(field) != roundtrip.get(field):
                raise ValidationError("主表行映射与服务器生成范围不一致。")
        required = {"columns", "products", "rows", "product_column", "kind_column", "key_column", "last_column", "headers",
                    "editable_months_by_product", "manual_editable_months", "plan_revisions_by_product"}
        if not required.issubset(manifest):
            raise ValidationError("主表行映射字段不完整。")
        months = [c.get("month") for c in manifest["columns"] if c.get("kind") == "month"]
        if months != roundtrip["months"]:
            raise ValidationError("主表月份与导出范围不一致。")
        visible = snapshot[PLAN_SHEET]["cells"]
        import openpyxl
        def cell_at(row, column):
            return _raw(visible, f"{openpyxl.utils.get_column_letter(column)}{row}")
        seen = set()
        for row in manifest["rows"]:
            for field, value in (("product_column", row.get("product_id")), ("kind_column", row["kind"]), ("key_column", row["key"])):
                if cell_at(row["row"], manifest[field]) != value:
                    raise ValidationError("主表隐藏行标识与生成记录不一致。")
            if row["kind"] == "product" and row.get("metric", "SOP计划量") == "SOP计划量":
                pid = row["product_id"]
                if pid in seen or pid not in roundtrip["selected_product_ids"] or row.get("values") != roundtrip["plan_values_by_product"][pid]:
                    raise ValidationError("主表版本计划基准与服务器数据不一致。")
                seen.add(pid)
        if roundtrip["include_versions"]:
            if seen != set(roundtrip["selected_product_ids"]):
                raise ValidationError("主表版本明细映射不完整。")
            parsed = parse_inline_sheet(workbook[PLAN_SHEET], manifest)
            if parsed["catalog_changes"] or parsed["entries"] or parsed["removed_ids"]:
                raise ValidationError("导出主表中的名称或版本行与服务器记录不一致。")
            for item in parsed["values"]:
                if item["sop"] != roundtrip["baseline"][_key(item["product_id"], item["month"])]["baseline_sop"]:
                    raise ValidationError("主表导出计划量与服务器基准不一致。")
        elif seen:
            raise ValidationError("车型汇总导出不应包含可回填的版本明细。")
        mappings = []
        for row in manifest["rows"]:
            if row["kind"] != "product" or row.get("metric", "SOP计划量") != "SOP计划量":
                continue
            for col in manifest["columns"]:
                if col.get("kind") == "month" and col["month"] in roundtrip["editable_months_by_product"][row["product_id"]]:
                    mappings.append({"kind": "plan", "product_id": row["product_id"], "month": col["month"],
                                     "cell": f"{openpyxl.utils.get_column_letter(col['column'])}{row['row']}",
                                     "baseline_sop": row["values"].get(col["month"])})
        return {**roundtrip, "manifest": manifest, "mappings": mappings, "snapshot": snapshot, "inline_schema": 2}

    def _parse(self, content):
        with _read_workbook(content) as workbook:
            token, versions, _ = _metadata(workbook)
            with self.store.lock:
                row = self.store.db.execute("SELECT payload FROM plan_download_exports WHERE token=?", (token,)).fetchone()
            if row is None:
                raise ValidationError("这个文件不属于当前看板，或导出标识已被修改，请在本看板重新导出。")
            expected = json.loads(row["payload"])
            if not versions or not expected["include_versions"]:
                raise ValidationError("该文件只有车型汇总，不能确定各版本计划量。请勾选“包含版本明细”重新导出后回填。")
            actual = _snapshot(workbook)
            if set(actual) != set(expected["snapshot"]):
                raise ValidationError("工作表已被增加、删除或重命名，请保留原导出格式。")
            if expected.get("manifest") is not None:
                if actual[IMPORT_SHEET]["cells"] != expected["snapshot"][IMPORT_SHEET]["cells"]:
                    raise ValidationError("隐藏的原始行映射已被改动，请保留该标识页。")
                parsed = parse_inline_sheet(workbook[PLAN_SHEET], expected["manifest"])
                expected["_catalog_changes"] = parsed["catalog_changes"]
                values = {_key(r["product_id"], r["month"]): r["sop"] for r in parsed["values"]}
                return token, expected, values, parsed["entries"]
            editable = {m["cell"] for m in expected["mappings"]}
            for name, original in expected["snapshot"].items():
                current = actual[name]
                if name == IMPORT_SHEET and current["state"] == "visible":
                    raise ValidationError("导出标识页已被改动，请重新导出。")
                for address in set(original["cells"]) | set(current["cells"]):
                    if name == PLAN_SHEET and address in editable:
                        continue
                    if original["cells"].get(address) != current["cells"].get(address):
                        raise ValidationError(f"{name}!{address} 是锁单、汇总、表头或同步标识，请只修改版本 SOP 月计划和底部新增区域。")
            sheet_cells = actual[PLAN_SHEET]["cells"]
            values = {}
            for item in expected["mappings"]:
                if item["kind"] == "plan":
                    values[_key(item["product_id"], item["month"])] = _cell_quantity(_raw(sheet_cells, item["cell"]))
            entries = []
            for row, fields in expected["manual_rows"].items():
                raw_labels = {field: sheet_cells.get(fields[field]) for field in ("brand", "model", "version")}
                if any(cell and cell[0] == "f" for cell in raw_labels.values()):
                    raise ValidationError("新增区域的品牌、车型和版本不能填写公式。")
                labels = {field: text_field(cell[1] if cell else None, field, 200) for field, cell in raw_labels.items()}
                plans = [{"month": month, "sop": _cell_quantity(_raw(sheet_cells, fields[month])), "baseline_sop": None}
                         for month in expected["manual_editable_months"]]
                if not any(labels.values()) and all(p["sop"] is None for p in plans):
                    continue
                if not all(labels.values()):
                    raise ValidationError(f"第 {row} 行新增版本请同时填写品牌、车型和版本。")
                entries.append({**labels, "plans": plans})
            return token, expected, values, entries

    def import_content(self, content, choice=None):
        if choice not in {None, "db", "workbook"}:
            raise ValidationError("请选择保留网页值或使用表格值。")
        token, exported, values, entries = self._parse(content)
        with self._lock():
            self._sync_before()
            with self.store.lock:
                # Old downloads may still contain now-hidden summary source
                # rows. Preserve their identities and saved edits on import.
                catalog = {p["id"]: p for p in self.store.read_sales(include_excluded=True)["products"]}
                if exported.get("manifest") is not None:
                    previous = self.store.db.execute("SELECT payload FROM plan_download_entry_imports WHERE token=?", (token,)).fetchone()
                    accepted_entries = json.loads(previous["payload"]) if previous else {}
                    supplied_names = {json.dumps(catalog_name_key(entry), ensure_ascii=False, separators=(",", ":")) for entry in entries}
                    if set(accepted_entries) - supplied_names and supplied_names - set(accepted_entries):
                        raise ValidationError("此文件的新增行已导入过，旧副本中无法安全识别这些行的再次更名。请重新导出最新明细，再更名或删除。")
                    for name in supplied_names & set(accepted_entries):
                        product = catalog.get(accepted_entries[name])
                        if product is None or json.dumps(catalog_name_key(product), ensure_ascii=False, separators=(",", ":")) != name:
                            raise ValidationError("此文件以前新增的版本已在网页改名或合并。请重新导出最新明细，避免重复建立版本。")
                old_products = {p["id"]: p for p in exported["products"]}
                cells = {_key(m["product_id"], m["month"]): m["cell"] for m in exported["mappings"] if m["kind"] == "plan"}
                changes, conflicts = [], []
                for key, value in values.items():
                    base = exported["baseline"][key]
                    pid, month, original = base["product_id"], base["month"], base["baseline_sop"]
                    if value == original:
                        continue
                    if pid not in catalog:
                        raise ValidationError("文件中修改的版本已被合并或移除，请重新导出当前版本后维护。")
                    current = self.store.db.execute("SELECT * FROM plans WHERE product_id=? AND month=?", (pid, month)).fetchone()
                    current = dict(current) if current else {}
                    database = current.get("sop")
                    if value == database:
                        continue
                    if database != original:
                        product = catalog.get(pid, old_products[pid])
                        conflicts.append({"product_id": pid, "month": month, "baseline": original, "workbook": value,
                                          "database": database, "current": current or None, "cell": cells[key],
                                          **{field: product.get(field, "") for field in ("brand", "model", "version")}})
                        if choice != "workbook":
                            continue
                    changes.append({"product_id": pid, "month": month, "sop": value,
                                    "forecast": current.get("forecast"), "revision": current.get("revision", 0),
                                    "notes": current.get("notes", "")})
                if conflicts and choice is None:
                    return {"ok": False, "status": "conflict", "conflicts": conflicts, "imported": 0, "created": 0}
                try:
                    result = self.store.import_workbook_plans({"changes": changes, "entries": entries, "choice": choice,
                                                              "catalog_changes": exported.get("_catalog_changes", []),
                                                              "entry_tracking_token": token if exported.get("manifest") is not None else None,
                                                              "owner": "Excel导入", "reason": "导出计划表回填", "notes": ""})
                except ConflictError as error:
                    for conflict in error.conflicts:
                        index = conflict.get("manual_entry")
                        if isinstance(index, int) and 0 <= index < len(entries):
                            conflict.update({field: entries[index][field] for field in ("brand", "model", "version")})
                    return {"ok": False, "status": "conflict", "conflicts": conflicts + error.conflicts, "imported": 0, "created": 0}
            response = {"ok": True, "status": "imported", "token": token, "imported": result["saved_count"],
                        "created": result["created_count"], "conflicts": conflicts,
                        "catalog_updated": result.get("catalog_updated_count", 0),
                        "deactivated": result.get("deactivated_count", 0)}
            if exported.get("manifest") is not None and result["created_count"]:
                response["warning"] = "新增版本已保存；后续更名或删除这些新增行时，请先重新导出最新明细。"
            if self.workbook_sync is not None:
                # The database commit has succeeded. A blocked Excel replacement is
                # a pending writeback, never a failed import that invites a retry.
                try:
                    after = self.workbook_sync.sync()
                except Exception as error:
                    after = {"status": "error", "message": str(error)}
                response["workbook_sync"] = after
                if after.get("status") != "synced":
                    response["warning"] = (response.get("warning", "") + " 计划已导入；维护表待同步：" + after.get("message", "请保存并关闭 Excel 后同步。")).strip()
            return response
