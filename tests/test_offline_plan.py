"""Acceptance checks for the openpyxl maintenance workbook writer."""
import json
import os
import tempfile
import unittest
from datetime import datetime
from pathlib import Path
from unittest.mock import patch

from openpyxl import load_workbook

from sop_store import SOPStore, TZ
from sop_workbook import WorkbookSync
from tests.test_inline_workbook import change_rows
from tests.test_workbook_sync import edit_cell


class OfflinePlanWorkbookTest(unittest.TestCase):
    def setUp(self):
        temp = tempfile.TemporaryDirectory(prefix="sop-offline-plan-")
        self.addCleanup(temp.cleanup)
        self.root = Path(temp.name)
        source = self.root / "sales.json"
        self.source = source
        source.write_text(json.dumps({"meta": {}, "products": [
            {"id": "a", "brand": "问界", "model": "M5", "version": "Max"},
            {"id": "b", "brand": "问界", "model": "M5", "version": "Ultra"}],
            "sales": []}, ensure_ascii=False), encoding="utf-8")
        self.store = SOPStore(self.root / "test.sqlite3", source,
                              lambda: datetime(2026, 10, 27, 12, tzinfo=TZ))
        self.addCleanup(self.store.close)
        self.path = self.root / "SOP计划维护.xlsx"
        self.sync = WorkbookSync(self.store, self.path)
        environment = patch.dict(os.environ, {"SOP_EXCEL_ENGINE": "python"})
        environment.start()
        self.addCleanup(environment.stop)

    def _spec(self):
        return self.sync._parse(self.path.read_bytes())[1]["layout"]["sheets"][0]

    def test_default_sync_writes_readable_styled_workbook_and_imports_edit(self):
        self.store.save_plans({"owner": "验收", "changes": [
            {"product_id": "a", "month": "2026-10", "sop": 0, "forecast": 10, "revision": 0},
            {"product_id": "b", "month": "2026-10", "sop": 125, "forecast": None, "revision": 0}]})
        status = self.sync.sync()
        self.assertEqual(status["status"], "synced", status)
        token, expected, _, _, values = self.sync._parse(self.path.read_bytes())
        self.assertTrue(token)
        self.assertEqual((values["a\t2026-10"], values["b\t2026-10"]), (0, 125))
        self.assertEqual(expected["months"][-1], "2027-12")

        workbook = load_workbook(self.path, data_only=True)
        self.addCleanup(workbook.close)
        sheet = workbook["SOP计划"]
        self.assertEqual(sheet.freeze_panes, "D5")
        self.assertIn("A5:A6", {str(area) for area in sheet.merged_cells.ranges})
        self.assertIn("B5:B6", {str(area) for area in sheet.merged_cells.ranges})
        self.assertEqual(sheet["P5"].value, 0)
        self.assertEqual(sheet["P6"].value, 125)
        self.assertEqual(sheet["M7"].value, 125)
        self.assertEqual(sheet["P10"].value, 125)
        self.assertEqual(sheet["D5"].fill.fgColor.rgb[-6:], "F0F2F4")
        self.assertEqual(sheet["M5"].font.sz, 11)
        self.assertEqual(sheet["P6"].font.color.rgb[-6:], "B62025")
        self.assertFalse(sheet["M5"].protection.locked)
        self.assertTrue(sheet["P5"].protection.locked)
        self.assertEqual(workbook["_layout"].sheet_state, "hidden")
        from openpyxl.utils import get_column_letter
        self.assertTrue(sheet.column_dimensions[get_column_letter(expected["layout"]["sheets"][0]["product_column"])].hidden)

        address = next(cell["cell"] for cell in expected["layout"]["cells"]
                       if cell["product_id"] == "a" and cell["month"] == "2026-10")
        edit_cell(self.path, address, 15)
        status = self.sync.sync()
        self.assertEqual(status["status"], "synced", status)
        plan = self.store.db.execute("SELECT sop, forecast FROM plans WHERE product_id='a' AND month='2026-10'").fetchone()
        self.assertEqual((plan["sop"], plan["forecast"]), (15, 10))

    def test_inline_add_and_delete_still_import_with_python_writer(self):
        self.assertEqual(self.sync.sync()["status"], "synced")
        spec = self._spec()
        second = next(row["row"] for row in spec["rows"] if row.get("product_id") == "b")
        month = next(column["column"] for column in spec["columns"] if column.get("month") == "2026-10")
        from openpyxl.utils import get_column_letter
        change_rows(self.path, second, insert={"A": "问界", "B": "M7", "C": "新版本",
                                                get_column_letter(month): 7})
        status = self.sync.sync()
        self.assertEqual(status["status"], "synced", status)
        added = next(product for product in self.store.list_products()["items"] if product["model"] == "M7")
        plan = self.store.db.execute("SELECT sop FROM plans WHERE product_id=? AND month='2026-10'",
                                     (added["id"],)).fetchone()
        self.assertEqual(plan["sop"], 7)

        spec = self._spec()
        first = next(row["row"] for row in spec["rows"] if row.get("product_id") == "a")
        change_rows(self.path, first, delete=True)
        status = self.sync.sync()
        self.assertEqual(status["status"], "synced", status)
        product = next(item for item in self.store.list_products()["items"] if item["id"] == "a")
        self.assertFalse(product["active"])

    def test_formula_like_catalog_labels_are_literal_text(self):
        sales = json.loads(self.source.read_text(encoding="utf-8"))
        sales["products"][0].update(brand="=问界", model="+M5", version="@高配")
        self.source.write_text(json.dumps(sales, ensure_ascii=False), encoding="utf-8")
        status = self.sync.sync()
        self.assertEqual(status["status"], "synced", status)
        workbook = load_workbook(self.path, data_only=False)
        self.addCleanup(workbook.close)
        sheet = workbook["SOP计划"]
        self.assertEqual([sheet[f"{column}5"].value for column in "ABC"], ["=问界", "+M5", "@高配"])
        self.assertTrue(all(sheet[f"{column}5"].data_type == "s" for column in "ABC"))
        self.assertEqual(self.sync._parse(self.path.read_bytes())[4]["a\t2026-10"], None)


if __name__ == "__main__":
    unittest.main()
