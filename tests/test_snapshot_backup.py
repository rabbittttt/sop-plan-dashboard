import http.client
import json
import tempfile
import threading
import unittest
import zipfile
from datetime import datetime, timedelta
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path

from server import make_handler
from sop_backup import SnapshotBackup
from sop_store import SOPStore, TZ


def fixture_writer(payload):
    """Minimal temporary archive for backend atomicity tests, not a user workbook."""
    with zipfile.ZipFile(payload["output_path"], "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("xl/workbook.xml", "<workbook/>")
        archive.writestr("test_payload.json", json.dumps(payload["snapshot_data"], ensure_ascii=False))


class SnapshotBackupTest(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(prefix="sop-snapshot-backup-test-")
        self.root = Path(self.temp.name)
        self.current = datetime(2026, 10, 22, 12, tzinfo=TZ)
        sales = {"meta": {"data_through": "2026-08"}, "products": [
            {"id": "a", "brand": "A", "model": "M", "version": "V1"},
            {"id": "b", "brand": "A", "model": "M", "version": "V2"}], "sales": [], "sources": []}
        (self.root / "sales.json").write_text(json.dumps(sales), encoding="utf-8")
        self.store = SOPStore(self.root / "sop.sqlite3", self.root / "sales.json", lambda: self.current)
        self.store.save_plans({"owner": "原负责人", "reason": "原原因", "notes": "原备注", "changes": [
            {"product_id": "a", "month": "2026-10", "sop": 0, "forecast": None, "revision": 0},
            {"product_id": "b", "month": "2026-10", "sop": None, "forecast": 45, "revision": 0}]})
        self.snapshot = self.store.manual_snapshot("原始快照")
        self.snapshot_id = self.snapshot["snapshot"]["id"]
        self.current += timedelta(minutes=1)
        self.store.save_plans({"owner": "新负责人", "reason": "新原因", "notes": "新备注", "changes": [
            {"product_id": "a", "month": "2026-10", "sop": 999, "forecast": 888, "revision": 1}]})
        self.backup = SnapshotBackup(self.store, writer=fixture_writer)

    def tearDown(self):
        self.store.close()
        self.temp.cleanup()

    def test_unit_frozen_values_restart_skip_and_missing_rebuild(self):
        self.assertEqual(self.backup.status()["pending"], 1)
        result = self.backup.ensure(self.snapshot_id)
        self.assertEqual(result["status"], "ready")
        path = Path(result["path"])
        self.assertEqual(path.parent, (self.root / "backups/sop_excel/2026-10").resolve())
        self.assertIn("2026-10-22_manual_000001", path.name)
        with zipfile.ZipFile(path) as archive:
            frozen = json.loads(archive.read("test_payload.json"))
        self.assertEqual(frozen["plans"][0]["sop"], 0)
        self.assertIsNone(frozen["plans"][0]["forecast"])
        self.assertIsNone(frozen["plans"][1]["sop"])
        self.assertEqual(frozen["plans"][0]["owner"], "原负责人")
        original = path.read_bytes()

        def unexpected_write(_):
            raise AssertionError("Existing historical backups must not be rewritten")

        restarted = SnapshotBackup(self.store, writer=unexpected_write)
        self.assertEqual(restarted.ensure(self.snapshot_id)["status"], "ready")
        self.assertEqual(path.read_bytes(), original)
        path.unlink()
        restarted.writer = fixture_writer
        self.assertEqual(restarted.ensure(self.snapshot_id)["status"], "ready")
        with zipfile.ZipFile(path) as archive:
            regenerated = json.loads(archive.read("test_payload.json"))
        self.assertEqual(regenerated["plans"], frozen["plans"])

    def test_unit_failure_is_visible_retriable_and_keeps_snapshot(self):
        def failing(_):
            raise PermissionError("模拟文件夹暂不可写")

        self.backup.writer = failing
        result = self.backup.ensure(self.snapshot_id)
        self.assertEqual(result["status"], "error")
        self.assertEqual(self.backup.status()["failed"], 1)
        self.assertIn("模拟文件夹", self.backup.status()["last_error"])
        self.assertEqual(len(self.store.snapshots()["items"]), 1)
        self.assertEqual(self.backup.export_pending(), 0)  # Backoff avoids a hot retry loop.
        self.backup.writer = fixture_writer
        self.assertEqual(self.backup.ensure(self.snapshot_id)["status"], "ready")
        self.assertEqual(self.backup.status()["failed"], 0)
        self.assertFalse(list(self.backup.directory.rglob("*.pending.xlsx")))

    def test_unit_writer_does_not_hold_database_lock_or_overwrite_existing(self):
        observed = []

        def concurrent_writer(payload):
            thread = threading.Thread(target=lambda: observed.append(self.store.state()["plans"][0]["sop"]))
            thread.start()
            thread.join(timeout=1)
            self.assertFalse(thread.is_alive(), "XLSX generation must not hold store.lock")
            fixture_writer(payload)
            target = self.backup.path_for(self.snapshot["snapshot"])
            with zipfile.ZipFile(target, "w") as archive:
                archive.writestr("xl/workbook.xml", "<workbook/>")
                archive.writestr("existing.txt", "existing historical export")

        self.backup.writer = concurrent_writer
        result = self.backup.ensure(self.snapshot_id)
        self.assertEqual(result["status"], "ready")
        self.assertEqual(observed, [999])
        with zipfile.ZipFile(result["path"]) as archive:
            self.assertEqual(archive.read("existing.txt"), b"existing historical export")

    def test_unit_existing_corrupt_backup_is_not_ready_or_overwritten(self):
        path = self.backup.path_for(self.snapshot["snapshot"])
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_bytes(b"corrupt workbook")
        self.assertEqual(self.backup.ensure(self.snapshot_id)["status"], "error")
        self.assertEqual(self.backup.status()["failed"], 1)
        self.assertEqual(path.read_bytes(), b"corrupt workbook")
        path.unlink()
        self.assertEqual(self.backup.ensure(self.snapshot_id)["status"], "ready")

    def test_unit_worker_exports_daily_sunday_and_manual(self):
        self.current = datetime(2026, 10, 26, 1, tzinfo=TZ)
        self.store.ensure_daily_snapshots()
        self.assertEqual(self.backup.status()["total"], 5)
        self.assertEqual(self.backup.export_pending(limit=20), 5)
        status = self.backup.status()
        self.assertEqual((status["ready"], status["pending"], status["failed"]), (5, 0, 0))
        names = [item["filename"] for item in status["items"]]
        self.assertTrue(any("_sunday_" in name for name in names))
        self.assertTrue(any("_daily_" in name for name in names))
        self.assertTrue(any("_manual_" in name for name in names))

    def test_unit_http_status_snapshot_download_and_failure(self):
        server = ThreadingHTTPServer(("127.0.0.1", 0), BaseHTTPRequestHandler)
        server.RequestHandlerClass = make_handler(self.store, self.root, server.server_port, snapshot_backup=self.backup)
        thread = threading.Thread(target=server.serve_forever, daemon=True)
        thread.start()

        def request(method, path, body=None):
            connection = http.client.HTTPConnection("127.0.0.1", server.server_port, timeout=10)
            headers = {"Content-Type": "application/json"} if body else {}
            connection.request(method, path, json.dumps(body).encode() if body else None, headers)
            response = connection.getresponse()
            result = response.status, response.read()
            connection.close()
            return result

        try:
            code, raw = request("GET", "/api/backups/status")
            self.assertEqual((code, json.loads(raw)["pending"]), (200, 1))
            code, raw = request("GET", f"/api/snapshots/{self.snapshot_id}")
            self.assertEqual(json.loads(raw)["excel_backup"]["status"], "pending")
            code, raw = request("GET", f"/api/snapshots/{self.snapshot_id}/excel")
            self.assertEqual(code, 200)
            self.assertTrue(raw.startswith(b"PK"))
            code, raw = request("GET", "/api/snapshots")
            self.assertEqual(json.loads(raw)["items"][0]["excel_backup"]["status"], "ready")
            self.backup.path_for(self.snapshot["snapshot"]).write_bytes(b"corrupt existing file")
            code, raw = request("GET", f"/api/snapshots/{self.snapshot_id}/excel")
            self.assertEqual(code, 503)
            self.assertEqual(json.loads(raw)["excel_backup"]["status"], "error")
            code, raw = request("POST", "/api/snapshots", {"label": "HTTP手工"})
            self.assertEqual(json.loads(raw)["excel_backup"]["status"], "ready")
            failed = self.store.manual_snapshot("故障仍保留")
            self.backup.writer = lambda _: (_ for _ in ()).throw(RuntimeError("模拟写失败"))
            code, raw = request("GET", f"/api/snapshots/{failed['snapshot']['id']}/excel")
            self.assertEqual(code, 503)
            self.assertEqual(json.loads(raw)["excel_backup"]["status"], "error")
            self.assertEqual(len(self.store.snapshots()["items"]), 3)
        finally:
            server.shutdown()
            server.server_close()
            thread.join(timeout=2)

    def test_unit_manual_snapshot_imports_saved_sheet_and_blocks_conflict(self):
        store = self.store

        class SavedWorkbook:
            _input_checked = True
            mode = "pending"

            def sync(self):
                if self.mode == "pending":
                    store.save_plans({"owner": "Excel计划表", "reason": "表格同步", "changes": [
                        {"product_id": "a", "month": "2026-10", "sop": 555, "forecast": 888, "revision": 2}]})
                return {"status": self.mode, "message": "测试同步状态", "conflicts": []}

        saved_workbook = SavedWorkbook()
        server = ThreadingHTTPServer(("127.0.0.1", 0), BaseHTTPRequestHandler)
        server.RequestHandlerClass = make_handler(self.store, self.root, server.server_port, saved_workbook, self.backup)
        thread = threading.Thread(target=server.serve_forever, daemon=True)
        thread.start()

        def create():
            connection = http.client.HTTPConnection("127.0.0.1", server.server_port, timeout=10)
            connection.request("POST", "/api/snapshots", b'{"label":"test"}', {"Content-Type": "application/json"})
            response = connection.getresponse()
            result = response.status, json.loads(response.read())
            connection.close()
            return result

        try:
            code, result = create()
            self.assertEqual(code, 200)
            self.assertEqual(result["plans"][0]["sop"], 555)
            self.assertEqual(result["excel_backup"]["status"], "ready")
            saved_workbook.mode = "conflict"
            self.assertEqual(create()[0], 409)
            saved_workbook.mode = "error"
            self.assertEqual(create()[0], 400)
            self.assertEqual(len(self.store.snapshots()["items"]), 2)
        finally:
            server.shutdown()
            server.server_close()
            thread.join(timeout=2)

    def test_actual_writer_preserves_snapshot_zero_blank_and_original_time(self):
        import openpyxl
        backup = SnapshotBackup(self.store)
        result = backup.ensure(self.snapshot_id)
        self.assertEqual(result["status"], "ready", result["message"])
        workbook = openpyxl.load_workbook(result["path"], read_only=True, data_only=False)
        try:
            self.assertEqual(workbook.sheetnames, ["SOP计划", "计划明细", "备份信息"])
            sheet = workbook["SOP计划"]
            self.assertEqual(sheet["E10"].value, 0)
            self.assertIsNone(sheet["E11"].value)
            details = list(workbook["计划明细"].iter_rows(values_only=True))
            headers = list(details[4])
            row_a = next(row for row in details[5:] if row[headers.index("product_id")] == "a")
            row_b = next(row for row in details[5:] if row[headers.index("product_id")] == "b")
            self.assertEqual(row_a[headers.index("sop")], 0)
            self.assertIsNone(row_a[headers.index("forecast")])
            self.assertIsNone(row_b[headers.index("sop")])
            self.assertEqual(row_b[headers.index("forecast")], 45)
            self.assertEqual(row_a[headers.index("owner")], "原负责人")
            self.assertEqual(row_a[headers.index("updated_at")], self.snapshot["plans"][0]["updated_at"])
            all_values = [cell.value for worksheet in workbook for row in worksheet for cell in row]
            self.assertIn(self.snapshot["snapshot"]["cutoff_at"], all_values)
            self.assertNotIn(999, all_values)
            self.assertNotIn(888, all_values)
        finally:
            workbook.close()


if __name__ == "__main__":
    unittest.main()
