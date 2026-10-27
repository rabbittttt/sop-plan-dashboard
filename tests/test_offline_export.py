"""The intranet writer must preserve the real download/import contract."""
import json
import os
import tempfile
import unittest
from datetime import datetime
from pathlib import Path
from unittest.mock import patch

import openpyxl

from sop_export import PlanExport
from sop_roundtrip import PlanRoundtrip
from sop_store import SOPStore, TZ
from sop_workbook import WorkbookSync
from tests.test_roundtrip import export_body
from tests.test_workbook_sync import edit_cell


class OfflineExportTest(unittest.TestCase):
    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory()
        self.root = Path(self.temporary.name)
        source = self.root / "sales.json"
        source.write_text(json.dumps({"meta": {}, "products": [
            {"id": "a", "brand": "A", "model": "M", "version": "V1"},
            {"id": "b", "brand": "A", "model": "M", "version": "V2"}], "sales": []}), encoding="utf-8")
        self.store = SOPStore(self.root / "data.sqlite3", source,
                              lambda: datetime(2026, 10, 26, 12, tzinfo=TZ))
        self.service = PlanRoundtrip(self.store)
        self.path = self.root / "export.xlsx"

    def tearDown(self):
        self.store.close()
        self.temporary.cleanup()

    def save(self, pid, month, value):
        self.store.save_plans({"owner": "test", "changes": [{"product_id": pid,
            "month": month, "sop": value, "forecast": None, "revision": 0}]})

    def body(self, *, locks=True, versions=True):
        body = export_body(locks=locks, versions=versions)
        for row, pid in zip(body["rows"], body["row_product_ids"]):
            if pid:
                row[3] = "V1" if pid == "a" else "V2"
        return body

    def generate(self, body):
        clean, context = self.service.prepare(body)
        # Exercise the same PlanExport default path used by the HTTP endpoint.
        with patch.dict(os.environ, {"SOP_EXCEL_ENGINE": "python"}):
            with PlanExport(temp_root=self.root).generate(clean, roundtrip=context) as result:
                self.path.write_bytes(result.path.read_bytes())
        self.service.register(self.path, context)
        return context

    def test_real_export_registers_and_imports_changed_plan(self):
        self.save("a", "2026-10", 10)
        self.save("a", "2026-11", 25)  # Outside the selected months, still in annual total.
        context = self.generate(self.body())
        book = openpyxl.load_workbook(self.path, data_only=False)
        try:
            self.assertEqual(book.sheetnames, ["SOP计划", "_sop_import"])
            self.assertEqual(book["_sop_import"].sheet_state, "veryHidden")
            sheet = book["SOP计划"]
            self.assertEqual(sheet["D4"].value, "10月")
            self.assertEqual(sheet["F4"].value, "年度合计")
            self.assertIn("A5:A6", {str(item) for item in sheet.merged_cells.ranges})
            self.assertEqual(sheet["D9"].value, 999)  # Model lock, not version sum.
            self.assertEqual(sheet["C7"].value, "V1")
            self.assertEqual(sheet["D7"].value, 10)
            self.assertTrue(sheet["F7"].data_type == "f")
            self.assertFalse(sheet["D7"].protection.locked)
            self.assertTrue(sheet["D5"].protection.locked)
            self.assertEqual(sheet["G7"].value, "a")
            self.assertEqual(sheet["I7"].value, "plan:a")
        finally:
            book.close()
        cached = openpyxl.load_workbook(self.path, data_only=True, read_only=True)
        try:
            self.assertEqual(cached["SOP计划"]["F7"].value, 35)
            self.assertEqual(cached["SOP计划"]["D10"].value, 10)
        finally:
            cached.close()
        edit_cell(self.path, "D7", 41)
        result = self.service.import_content(self.path.read_bytes())
        self.assertTrue(result["ok"], result)
        self.assertEqual(result["imported"], 1)
        self.assertEqual(self.store.db.execute(
            "SELECT sop FROM plans WHERE product_id='a' AND month='2026-10'").fetchone()[0], 41)
        self.assertEqual(context["plan_values_by_product"]["a"]["2026-10"], 10)

    def test_first_web_template_exports_and_imports_without_creating_maintenance_file(self):
        maintenance = self.root / "input_file" / "SOP计划维护.xlsx"
        sync = WorkbookSync(self.store, maintenance, auto_create=False)
        self.service = PlanRoundtrip(self.store, sync)
        self.assertEqual(sync.sync()["status"], "absent")
        self.assertFalse(maintenance.exists())
        self.save("a", "2026-10", 10)
        self.generate(self.body())
        self.assertEqual(sync.status()["status"], "absent")
        self.assertFalse(maintenance.exists())
        edit_cell(self.path, "D7", 41)
        result = self.service.import_content(self.path.read_bytes())
        self.assertTrue(result["ok"], result)
        self.assertEqual(result["imported"], 1)
        self.assertFalse(maintenance.exists())
        self.assertEqual(self.store.db.execute(
            "SELECT sop FROM plans WHERE product_id='a' AND month='2026-10'").fetchone()[0], 41)

    def test_without_locks_or_versions_keeps_selected_months_and_read_only(self):
        self.save("a", "2026-10", 10)
        self.generate(self.body(locks=False, versions=False))
        book = openpyxl.load_workbook(self.path)
        try:
            sheet = book["SOP计划"]
            self.assertEqual(sheet["C4"].value, "10月")
            self.assertEqual(sheet["E4"].value, "年度合计")
            self.assertTrue(sheet.protection.sheet)
            self.assertEqual(sheet["B5"].value, "M 合计")
            self.assertEqual(sheet["C5"].value, 10)
            self.assertEqual(sheet["F5"].value, None)
            self.assertEqual(book["_sop_import"]["B3"].value, False)
        finally:
            book.close()

    def test_external_names_are_literal_text_not_excel_formulas(self):
        current = next(p for p in self.store.list_products()["items"] if p["id"] == "a")
        self.store.save_products({"changes": [{"id": "a", "revision": current["revision"],
                                  "brand": "=Brand", "version": "=1+1"}]})
        body = export_body(locks=False, selected=("a",))
        for row in body["rows"]:
            row[1] = "=Brand"
            if row[0] == "版本明细":
                row[3] = "=1+1"
        self.generate(body)
        book = openpyxl.load_workbook(self.path)
        try:
            sheet = book["SOP计划"]
            self.assertEqual(sheet["A5"].value, "=Brand")
            self.assertEqual(sheet["A5"].data_type, "s")
            self.assertEqual(sheet["C5"].value, "=1+1")
            self.assertEqual(sheet["C5"].data_type, "s")
        finally:
            book.close()

    def test_selected_months_show_each_december_annual_total(self):
        self.save("a", "2026-11", 7)
        self.save("a", "2027-12", 9)
        body = export_body(selected=("a",), locks=False)
        body["months"] = ["2026-11", "2026-12", "2027-01", "2027-12"]
        body["headers"] = ["层级", "品牌", "车型", "版本", "指标", *body["months"], "数据说明"]
        for row, pid in zip(body["rows"], body["row_product_ids"]):
            row[:] = row[:5] + [None] * len(body["months"]) + [""]
            if pid:
                row[3] = "V1"
        body["annual_plans"] = [{"2026": None, "2027": None} for _ in body["rows"]]
        self.generate(body)
        book = openpyxl.load_workbook(self.path, data_only=True, read_only=True)
        try:
            sheet = book["SOP计划"]
            self.assertEqual([sheet.cell(4, col).value for col in range(4, 10)],
                             ["11月", "12月", "年度合计", "1月", "12月", "年度合计"])
            self.assertEqual(sheet["F5"].value, 7)
            self.assertEqual(sheet["I5"].value, 9)
        finally:
            book.close()


if __name__ == "__main__":
    unittest.main()
