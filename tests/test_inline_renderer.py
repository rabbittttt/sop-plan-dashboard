"""Both real XLSX writers support editing the same primary table structure."""
import copy
import json
import shutil
import sqlite3
import subprocess
import tempfile
import unittest
from contextlib import closing
from pathlib import Path
from unittest.mock import patch

import openpyxl

from sop_workbook import _inline_layout

ROOT = Path(__file__).resolve().parents[1]
NODE = Path.home() / ".cache/codex-runtimes/codex-primary-runtime/dependencies/node/bin/node.exe"
OUT = ROOT / "tests/artifacts/inline-layout"
MONTHS = [f"{year}-{month:02}" for year in (2026, 2027) for month in range(1, 13)]
PRODUCTS = [{"id": "a", "brand": "问界", "model": "M5", "version": "Max", "active": True},
            {"id": "b", "brand": "问界", "model": "M5", "version": "Ultra", "active": False, "end_month": "2026-09"},
            {"id": "c", "brand": "智界", "model": "R7", "version": "=1+1", "active": True}]
PLANS = [{"product_id": "a", "month": "2026-08", "sop": 10, "revision": 1},
         {"product_id": "a", "month": "2026-10", "sop": 0, "revision": 1},
         {"product_id": "a", "month": "2026-12", "sop": 30, "revision": 1},
         {"product_id": "b", "month": "2026-08", "sop": 20, "revision": 1},
         {"product_id": "b", "month": "2026-12", "sop": 40, "revision": 1},
         {"product_id": "c", "month": "2026-08", "sop": 5, "revision": 1},
         {"product_id": "c", "month": "2026-10", "sop": 0, "revision": 1}]
LOOKUP = {(p["product_id"], p["month"]): p for p in PLANS}


def sum_known(values):
    values = [v for v in values if v is not None]
    return sum(values) if values else None


def input_payload():
    layout = _inline_layout(PRODUCTS, MONTHS, MONTHS, "2026-10", LOOKUP)
    return {"schema_version": 5, "layout_version": "continuous-inline-v1", "token": "input-test-token",
            "generated_at": "2026-10-26T12:00:00+08:00", "months": MONTHS, "products": PRODUCTS,
            "plans": PLANS, "editable_months": MONTHS, "current_month": "2026-10", "layout": layout,
            "baseline": [{"product_id": p["id"], "month": m, "baseline_sop": LOOKUP.get((p["id"], m), {}).get("sop"),
                          "revision": LOOKUP.get((p["id"], m), {}).get("revision", 0)} for p in PRODUCTS for m in MONTHS]}


def export_payload(months=MONTHS, locks=False, versions=True):
    groups = {}
    for p in PRODUCTS:
        groups.setdefault((p["brand"], p["model"]), []).append(p)
    rows, annual, ids, counts = [], [], [], []
    for (brand, model), group in groups.items():
        sources = [(group, False)] + ([( [p], True) for p in group] if versions else [])
        for members, detail in sources:
            for metric in (["锁单数", "SOP计划量"] if locks else ["SOP计划量"]):
                values = [sum_known(LOOKUP.get((p["id"], m), {}).get("sop") for p in members) if metric == "SOP计划量" else (70 if not detail and m == "2026-08" else None) for m in months]
                totals, known = {}, {}
                for year in {m[:4] for m in months if m.endswith("-12")}:
                    monthly = [sum_known(LOOKUP.get((p["id"], f"{year}-{m:02}"), {}).get("sop") for p in members) for m in range(1, 13)]
                    totals[year] = sum_known(monthly) if metric == "SOP计划量" else None
                    known[year] = sum(v is not None for v in monthly) if metric == "SOP计划量" else 0
                rows.append(["版本明细" if detail else "车型汇总", brand, model, *([members[0]["version"] if detail else ""] if versions else []), metric, *values, ""])
                annual.append(totals); counts.append(known); ids.append(members[0]["id"] if detail else None)
    layout = _inline_layout(PRODUCTS, months, months, "2026-10", LOOKUP)["sheets"][0]
    return {"schema_version": 1, "current_month": "2026-10", "months": list(months), "include_locks": locks,
            "include_versions": versions, "scope": "验收", "rows": rows, "annual_plans": annual,
            "headers": ["层级", "品牌", "车型", *(["版本"] if versions else []), "指标", *months, "数据说明"],
            "roundtrip": {"inline_schema": 2, "token": "0123456789abcdef0123456789abcdef", "products": PRODUCTS,
                "row_product_ids": ids, "annual_known_counts": counts, "plan_revisions_by_product": layout["plan_revisions_by_product"],
                "editable_months_by_product": layout["editable_months_by_product"], "manual_editable_months": list(months)}}


def generate(script, body, name, preview=False):
    output = OUT / (name + ".xlsx")
    request = {**body, "output_path": str(output)}
    if preview:
        request["preview_path"] = str(OUT / (name + ".png"))
    result = subprocess.run([str(NODE), str(ROOT / "scripts" / script)], input=json.dumps(request, ensure_ascii=False),
                            cwd=ROOT, text=True, encoding="utf-8", capture_output=True, timeout=120)
    (OUT / (name + ".log")).write_text(result.stderr, encoding="utf-8")
    if result.returncode:
        raise AssertionError(result.stderr)
    return output


def manifest(book):
    meta = book["_sop_import"]
    return json.loads("".join(row[1] for row in meta.iter_rows(min_row=6, values_only=True)))


@unittest.skipUnless(NODE.is_file(), "Bundled spreadsheet runtime unavailable")
class InlineRendererTest(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        OUT.mkdir(parents=True, exist_ok=True)
        cls.primary = generate("build_plan_workbook.mjs", input_payload(), "maintenance")
        cls.export = generate("build_export_workbook.mjs", export_payload(), "export")
        cls.partial = generate("build_export_workbook.mjs", export_payload(["2026-10", "2026-12", "2027-02"], locks=True), "with-locks", preview=True)
        cls.report = generate("build_export_workbook.mjs", export_payload(["2026-12"], versions=False), "summary")

    def test_input_and_export_have_same_visible_grid_and_no_second_input_area(self):
        source = openpyxl.load_workbook(self.primary)["SOP计划"]
        exported = openpyxl.load_workbook(self.export)["SOP计划"]
        visible = source.max_column - 3
        self.assertEqual(source.max_row, exported.max_row)
        self.assertEqual(source.freeze_panes, exported.freeze_panes)
        self.assertEqual(set(map(str, source.merged_cells.ranges)), set(map(str, exported.merged_cells.ranges)))
        for r in range(1, source.max_row + 1):
            self.assertEqual(source.row_dimensions[r].height, exported.row_dimensions[r].height)
            self.assertEqual(source.row_dimensions[r].hidden, exported.row_dimensions[r].hidden)
            for c in range(1, visible + 1):
                left, right = source.cell(r, c), exported.cell(r, c)
                if (r, c) != (2, 1):
                    self.assertEqual(left.value, right.value, left.coordinate)
                for field in ("font", "fill", "border", "alignment", "number_format"):
                    self.assertEqual(copy.copy(getattr(left, field)), copy.copy(getattr(right, field)), f"{left.coordinate} {field}")
        for c in range(1, visible + 1):
            letter = openpyxl.utils.get_column_letter(c)
            self.assertEqual(source.column_dimensions[letter].width, exported.column_dimensions[letter].width)
        self.assertFalse(source.protection.sheet)
        self.assertFalse(exported.protection.sheet)
        self.assertFalse(any(isinstance(c.value, str) and "新增车型/版本" in c.value for row in exported for c in row))

    def test_row_identity_and_inactive_rows_remain_attached_to_data(self):
        wb = openpyxl.load_workbook(self.partial)
        sheet, meta = wb["SOP计划"], wb["_sop_import"]
        self.assertEqual(meta.sheet_state, "veryHidden")
        self.assertEqual(meta["B1"].value, 2)
        spec = manifest(wb)
        for row in spec["rows"]:
            r = row["row"]
            self.assertEqual(sheet.cell(r, spec["key_column"]).value, row["key"])
            self.assertEqual(sheet.cell(r, spec["kind_column"]).value, row["kind"])
            if row["kind"] == "product":
                self.assertEqual(sheet.cell(r, spec["product_column"]).value, row["product_id"])
                self.assertEqual([sheet.cell(r, c).value for c in (1, 2, 3)], row["labels"])
                self.assertEqual(sheet.row_dimensions[r].hidden, row["product_id"] == "b")
        self.assertFalse(sheet.protection.sheet, "Excel must permit real row insertion/deletion")
        for c in (spec["product_column"], spec["kind_column"], spec["key_column"]):
            self.assertTrue(sheet.column_dimensions[openpyxl.utils.get_column_letter(c)].hidden)
        plan = next(r for r in spec["rows"] if r["key"] == "plan:c")
        self.assertEqual(sheet.cell(plan["row"], 3).data_type, "s")
        self.assertEqual(sheet.cell(plan["row"], 3).value, "=1+1")

    def test_summary_only_remains_readonly_without_extra_input_area(self):
        wb = openpyxl.load_workbook(self.report)
        self.assertFalse(wb["_sop_import"]["B3"].value)
        self.assertTrue(wb["SOP计划"].protection.sheet)
        self.assertTrue(all(r["kind"] != "manual_entry" for r in manifest(wb)["rows"]))

    def test_legacy_parity_reader_selects_new_contract_after_upgrade(self):
        from tests import test_export_workbook as legacy
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            (root / "input_file").mkdir(); (root / "data").mkdir()
            shutil.copy2(self.primary, root / "input_file/SOP计划维护.xlsx")
            with closing(sqlite3.connect(root / "data/sop.sqlite3")) as db, db:
                db.execute("CREATE TABLE workbook_exports (token TEXT, payload TEXT)")
                db.execute("INSERT INTO workbook_exports VALUES (?,?)", ("input-test-token", json.dumps(input_payload(), ensure_ascii=False)))
            with patch.object(legacy, "ROOT", root):
                body, source, width = legacy.actual_input_payload()
            self.assertEqual(body["roundtrip"]["inline_schema"], 2)
            self.assertEqual(body["roundtrip"]["products"], PRODUCTS)
            self.assertEqual(source.max_column, width + 3)
            self.assertEqual(len(body["rows"]), len(body["roundtrip"]["row_product_ids"]))

    def test_partial_year_formulas_keep_offscreen_amount_and_recalculate(self):
        script = r'''
import {FileBlob,SpreadsheetFile} from '@oai/artifact-tool';
const wb=await SpreadsheetFile.importXlsx(await FileBlob.load(process.argv[1]));
const sheet=wb.worksheets.getItem('SOP计划'),meta=wb.worksheets.getItem('_sop_import');
const manifest=JSON.parse(meta.getUsedRange().values.slice(5).map(r=>r[1]).join(''));
const row=manifest.rows.find(r=>r.key==='plan:a').row;
sheet.getRange('D'+row).values=[[15]];wb.recalculate();
if(sheet.getRange('F'+row).values[0][0]!==55)throw new Error('Changed monthly value did not flow into annual amount');
sheet.getRange('D'+row).values=[[null]];sheet.getRange('E'+row).values=[[null]];wb.recalculate();
if(sheet.getRange('F'+row).values[0][0]!==10)throw new Error('Offscreen August amount was lost');
console.log('Verified annual formulas retain offscreen input values');
'''
        result = subprocess.run([str(NODE), "--input-type=module", "-e", script, str(self.partial)],
                                cwd=ROOT / "scripts", text=True, encoding="utf-8", capture_output=True, timeout=120)
        self.assertEqual(result.returncode, 0, result.stderr)


if __name__ == "__main__":
    unittest.main()
