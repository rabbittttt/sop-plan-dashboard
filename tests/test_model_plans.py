"""Model totals override details per month without modifying their history."""
import unittest
from datetime import datetime

import openpyxl

from sop_plan_scope import MODEL_PLAN_LABEL, is_model_plan, plan_total
from sop_store import TZ, ValidationError
from sop_workbook import WorkbookSync
from sop_xlsx_plan import write_plan_workbook
from sop_xlsx_export import write_export_workbook
from sop_roundtrip import PlanRoundtrip
from tests import test_grid_batch as grid
from tests.test_workbook_sync import edit_cell
from tests.test_inline_workbook import change_rows


class ModelPlansTest(unittest.TestCase):
    setUp = grid.GridBatchTest.setUp
    body = grid.GridBatchTest.body
    plan = grid.GridBatchTest.plan
    product = grid.GridBatchTest.product

    def create_model(self, model="A 2026款", value=100):
        result = self.store.save_grid_batch(self.body(entries=[{
            "client_id": "model", "brand": "智界", "model": model, "plan_level": "model",
            "plans": [{"month": "2026-10", "sop": value}]}]))
        return result["entry_ids"]["model"]

    def total(self, sales=None, month="2026-10"):
        sales = sales or self.store.read_sales()
        plans = {(r["product_id"], r["month"]): r for r in self.store.view_plans(sales)}
        return plan_total(sales["products"], month, plans)

    def test_model_priority_zero_clear_and_year_separation(self):
        pid = self.create_model()
        self.assertEqual(self.total(), 120)
        self.assertEqual(self.plan("a")["sop"], 10)
        self.create_model("A 2025款", 50)
        self.assertEqual(self.total(), 170)
        for value, expected in [(0, 70), (None, 80)]:
            self.store.save_grid_batch(self.body(changes=[{"product_id": pid, "month": "2026-10", "sop": value,
                                                          "revision": self.plan(pid)["revision"]}]))
            self.assertEqual(self.total(), expected)
        self.assertEqual(self.plan("a")["forecast"], 110)
        with self.assertRaisesRegex(ValidationError, "年款"):
            self.create_model("A", 22)

    def test_summary_uses_each_member_model_priority_and_keeps_versions(self):
        self.create_model()
        self.store.update_settings({"generation_mode": "summary"})
        sales = self.store.read_sales()
        self.assertEqual(self.total(sales), 120)
        total = next(p for p in sales["products"] if is_model_plan(p))
        self.assertTrue(total["read_only"])
        self.assertEqual(next(r["sop"] for r in self.store.view_plans(sales) if r["product_id"] == total["id"]), 120)
        self.assertEqual(self.plan("a")["sop"], 10)

    def test_catalog_maintains_model_year_identity(self):
        pid = self.create_model()
        for change in ({"version": "Max"}, {"model": "A"}, {"bind_source_id": "a"}):
            with self.subTest(change=change), self.assertRaises(ValidationError):
                self.store.save_products({"changes": [{"id": pid, "revision": self.product(pid)["revision"], **change}]})
        with self.assertRaisesRegex(ValidationError, "一条车型总量"):
            self.store.save_products({"changes": [{"brand": "智界", "model": "A 2026款", "version": MODEL_PLAN_LABEL}]})

    def test_matching_uses_frozen_model_total_and_version_shipments(self):
        pid = self.create_model()
        self.store.save_grid_batch(self.body(changes=[
            {"product_id": p, "month": "2026-11", "sop": value, "revision": 0} for p, value in [("a", 10), (pid, 100)]]))
        self.store.clock = lambda: datetime(2026, 11, 3, 12, tzinfo=TZ)
        self.store.ensure_daily_snapshots()
        self.store.save_grid_batch(self.body(changes=[{"product_id": pid, "month": "2026-11", "sop": 200,
                                                       "revision": self.plan(pid, "2026-11")["revision"]}]))
        self.store.clock = lambda: datetime(2026, 12, 1, 12, tzinfo=TZ)
        self.store.save_shipments({"owner": "test", "changes": [{"product_id": "a", "month": "2026-11",
            "quantity": 80, "status": "settled", "source": "SP", "batch": "test", "revision": 0}]})
        result = self.store.matching("2026-11", model="A 2026款")
        self.assertEqual(result["actual"], 80)
        self.assertEqual(result["rows"][-1]["plan"], 100)
        self.assertEqual(result["rows"][-1]["match"], 80)

    def test_excel_new_row_without_version_creates_model_plan(self):
        service = PlanRoundtrip(self.store)
        clean, context = service.prepare(self.export_body(["2026-10", "2026-12"]))
        exported = self.root / "new-model.xlsx"
        write_export_workbook({**clean, "roundtrip": context, "output_path": str(exported)})
        service.register(exported, context)
        book = openpyxl.load_workbook(exported)
        try:
            at = book["SOP计划"].max_row
        finally:
            book.close()
        change_rows(exported, at, insert={"A": "智界", "B": "A 2026款", "D": 55})
        result = service.import_content(exported.read_bytes())
        self.assertTrue(result["ok"])
        self.assertEqual(self.total(), 75)
        self.assertEqual(self.plan("a")["sop"], 10)
        self.assertEqual(len([p for p in self.store.read_sales()["products"] if is_model_plan(p)]), 1)

    def export_body(self, months, locks=False):
        products = self.store.read_sales()["products"]
        groups = {}
        for product in products:
            groups.setdefault((product["brand"], product["model"]), []).append(product)
        rows, ids = [], []
        metrics = ["锁单数", "SOP计划量"] if locks else ["SOP计划量"]
        for (brand, model), members in groups.items():
            for metric in metrics:
                rows.append(["车型汇总", brand, model, "", metric, *([None] * len(months)), ""])
                ids.append(None)
            for p in members:
                for metric in metrics:
                    if metric == "锁单数" and is_model_plan(p):
                        continue
                    rows.append(["版本明细", brand, model, p["version"], metric, *([None] * len(months)), ""])
                    ids.append(p["id"])
        return {"schema_version": 1, "current_month": "2026-10", "months": months, "include_versions": True,
                "include_locks": locks, "scope": "all", "headers": ["层级", "品牌", "车型", "版本", "指标", *months, "数据说明"],
                "rows": rows, "row_product_ids": ids, "selected_product_ids": [p["id"] for p in products],
                "annual_plans": [{m[:4]: None for m in months if m.endswith("-12")} for _ in rows]}

    def test_export_with_locks_model_only_and_offscreen_annual(self):
        self.create_model()
        self.create_model("Solo 2027款", 0)
        self.store.save_grid_batch(self.body(changes=[
            {"product_id": "a", "month": month, "sop": value, "revision": 0}
            for month, value in [("2026-11", 5), ("2026-12", 9)]]))
        service = PlanRoundtrip(self.store)
        clean, context = service.prepare(self.export_body(["2026-10", "2026-12"], locks=True))
        path = self.root / "with-locks.xlsx"
        write_export_workbook({**clean, "roundtrip": context, "output_path": str(path)})
        service.register(path, context)
        book = openpyxl.load_workbook(path, data_only=True)
        try:
            sheet = book["SOP计划"]
            total = next(r for r in sheet.iter_rows() if r[1].value == "A 2026款 合计")
            self.assertEqual([c.value for c in total[3:6]], [100, 9, 114])
            solo = next(r for r in sheet.iter_rows() if r[1].value == "Solo 2027款 合计")
            self.assertEqual([c.value for c in solo[3:6]], [0, None, 0])
            self.assertFalse(any(r[2].value == MODEL_PLAN_LABEL + "（锁单数）" for r in sheet.iter_rows()))
        finally:
            book.close()
        self.assertTrue(service.import_content(path.read_bytes())["ok"])

    def test_maintenance_export_cached_totals_and_roundtrip_preserve_details(self):
        pid = self.create_model()
        path = self.root / "maintenance.xlsx"
        sync = WorkbookSync(self.store, path, writer=write_plan_workbook)
        self.assertEqual(sync.sync()["status"], "synced")
        book = openpyxl.load_workbook(path, data_only=True)
        try:
            sheet = book["SOP计划"]
            model_row = next(r for r in sheet.iter_rows() if r[1].value == "A 2026款 合计")
            self.assertEqual((model_row[12].value, model_row[15].value), (100, 100))  # Oct, annual.
        finally:
            book.close()
        roundtrip = PlanRoundtrip(self.store)
        clean, context = roundtrip.prepare(self.export_body(["2026-10", "2026-12"]))
        exported = self.root / "export.xlsx"
        write_export_workbook({**clean, "roundtrip": context, "output_path": str(exported)})
        roundtrip.register(exported, context)
        book = openpyxl.load_workbook(exported, data_only=True)
        try:
            sheet = book["SOP计划"]
            model = next(r for r in sheet.iter_rows() if r[1].value == "A 2026款 合计")
            self.assertEqual((model[3].value, model[5].value), (100, 100))
            row = next(r[0].row for r in sheet.iter_rows() if r[2].value == MODEL_PLAN_LABEL)
        finally:
            book.close()
        edit_cell(exported, f"D{row}", 70)
        result = roundtrip.import_content(exported.read_bytes())
        self.assertTrue(result["ok"])
        self.assertEqual((self.plan(pid)["sop"], self.plan("a")["sop"]), (70, 10))
        self.assertEqual(self.total(), 90)
        # Regenerate the Excel to get a new common baseline, then clear the override.
        clean, context = roundtrip.prepare(self.export_body(["2026-10", "2026-12"]))
        write_export_workbook({**clean, "roundtrip": context, "output_path": str(exported)})
        roundtrip.register(exported, context)
        edit_cell(exported, f"D{row}", None)
        self.assertTrue(roundtrip.import_content(exported.read_bytes())["ok"])
        self.assertEqual(self.total(), 30)
