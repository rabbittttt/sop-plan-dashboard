"""Editable export layout uses the same visible template and stable hidden IDs."""
import copy
import json
import subprocess
import unittest
from pathlib import Path

import openpyxl

from test_export_workbook import NODE, OUT, ROOT, actual_input_payload, generate, payload


def add_roundtrip(body):
    body = copy.deepcopy(body)
    if body.get("roundtrip", {}).get("inline_schema") == 2:
        return body
    ids, counts, editable = [], [], {}
    years = [m[:4] for m in body["months"] if m.endswith("-12")]
    metric = 4 if body["include_versions"] else 3
    counters = {}
    for index, row in enumerate(body["rows"]):
        pid = None
        if row[0] == "版本明细":
            key = (row[1], row[2], row[metric])
            number = counters.get(key, 0)
            counters[key] = number + 1
            pid = f"{row[1]}/{row[2]}/{number}"
            editable[pid] = list(body["months"])
        ids.append(pid)
        annual_count = {}
        for year in years:
            selected = sum(row[metric + 1 + i] is not None
                           for i, month in enumerate(body["months"]) if month[:4] == year)
            full = sum(m[:4] == year for m in body["months"]) == 12
            annual_count[year] = 0 if body["annual_plans"][index][year] is None else selected + (not full)
        counts.append(annual_count)
    body["roundtrip"] = {"token": "synthetic-renderer-token", "row_product_ids": ids,
                         "annual_known_counts": counts, "editable_months_by_product": editable,
                         "manual_editable_months": list(body["months"])}
    return body


@unittest.skipUnless(NODE.is_file(), "Bundled spreadsheet runtime unavailable")
class RoundtripRendererTest(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        OUT.mkdir(parents=True, exist_ok=True)
        cls.body = add_roundtrip(payload(True, True))
        cls.first_pid = next(pid for pid in cls.body["roundtrip"]["row_product_ids"] if pid)
        cls.body["roundtrip"]["editable_months_by_product"][cls.first_pid].remove("2027-02")
        cls.body["roundtrip"]["manual_editable_months"].remove("2027-02")
        result, cls.editable_path = generate(cls.body, "roundtrip-editable", preview=True)
        if result.returncode:
            raise AssertionError(result.stderr)
        body, cls.reference, cls.visible_last = actual_input_payload()
        cls.inline = body.get("roundtrip", {}).get("inline_schema") == 2
        result, cls.parity_path = generate(add_roundtrip(body), "roundtrip-input-parity")
        if result.returncode:
            raise AssertionError(result.stderr)
        result, cls.report_path = generate(add_roundtrip(payload(True, False)), "roundtrip-readonly")
        if result.returncode:
            raise AssertionError(result.stderr)
        edge = add_roundtrip(payload(False, True))
        edge["rows"], edge["annual_plans"] = edge["rows"][:3], edge["annual_plans"][:3]
        edge["rows"][0][5:9] = [0, None, None, None]
        edge["rows"][1][5:9] = [0, None, None, None]
        edge["rows"][2][5:9] = [None, None, None, None]
        edge["annual_plans"] = [{"2026": 0}] * 3
        edge = add_roundtrip(edge)
        edge["roundtrip"]["annual_known_counts"] = [{"2026": 2}, {"2026": 1}, {"2026": 1}]
        result, cls.zero_path = generate(edge, "roundtrip-zero-baseline")
        if result.returncode:
            raise AssertionError(result.stderr)

    def test_hidden_mapping_and_only_plan_inputs_editable(self):
        wb = openpyxl.load_workbook(self.editable_path, data_only=False)
        self.assertEqual(wb.sheetnames, ["SOP计划", "_sop_import"])
        s, meta = wb["SOP计划"], wb["_sop_import"]
        self.assertEqual(meta.sheet_state, "veryHidden")
        self.assertEqual([meta.cell(r, 2).value for r in range(1, 5)],
                         [1, "synthetic-renderer-token", True, "SOP计划"])
        self.assertEqual([c.value for c in meta[5]], ["kind", "product_id", "month", "cell", "baseline_sop"])
        rows = list(meta.iter_rows(min_row=6, values_only=True))
        plan = [r for r in rows if r[0] == "plan"]
        manual = [r for r in rows if r[0] == "manual"]
        self.assertEqual(len(plan), 11)
        self.assertEqual(len(manual), 60)
        mapped = {r[3] for r in rows}
        unlocked = {c.coordinate for row in s for c in row if not c.protection.locked}
        self.assertEqual(unlocked, mapped)
        self.assertTrue(s.protection.sheet)
        self.assertEqual(len(s.data_validations.dataValidation), len(plan) + 30)
        self.assertFalse(any(r[1] == self.first_pid and r[2] == "2027-02" for r in plan))
        self.assertFalse(any(r[2] == "2027-02" for r in manual))
        self.assertEqual({r[1] for r in plan}, set(self.body["roundtrip"]["editable_months_by_product"]))
        for _, _, _, address, baseline in plan:
            self.assertEqual(s[address].value, baseline)
        for row in s:
            if isinstance(row[2].value, str) and row[2].value.endswith("（锁单数）"):
                self.assertTrue(all(c.protection.locked for c in row))
        literal = [c for row in s for c in row if c.value == "=1+1"]
        self.assertEqual(len(literal), 1)
        self.assertEqual(literal[0].data_type, "s")

    def test_full_horizon_keeps_actual_input_visual_format(self):
        actual = openpyxl.load_workbook(self.parity_path)["SOP计划"]
        expected = self.reference
        self.assertEqual(actual.max_column, self.visible_last + (3 if self.inline else 0))
        self.assertEqual(actual.max_row, expected.max_row)
        self.assertEqual(actual.freeze_panes, expected.freeze_panes)
        # The existing business workbook may still use the earlier row-1 title merge.
        # Compare its working template while allowing the new freeze-safe title layout.
        actual_merges = {str(cell_range) for cell_range in actual.merged_cells.ranges
                         if cell_range.min_row > 1}
        expected_merges = {str(cell_range) for cell_range in expected.merged_cells.ranges
                           if cell_range.min_row > 1}
        self.assertEqual(actual_merges, expected_merges)
        guidance = {2} | {c.row for row in expected for c in row
                         if isinstance(c.value, str) and c.value.startswith("新增车型/版本：")}
        for r in range(2, actual.max_row + 1):
            self.assertEqual(actual.row_dimensions[r].height, expected.row_dimensions[r].height)
            for c in range(1, self.visible_last + 1):
                left, right = actual.cell(r, c), expected.cell(r, c)
                if not (r in guidance and c == 1):
                    self.assertEqual(left.value, right.value, left.coordinate)
                for field in ("font", "fill", "border", "alignment", "number_format"):
                    self.assertEqual(copy.copy(getattr(left, field)), copy.copy(getattr(right, field)),
                                     f"{left.coordinate} {field}")
        for c in range(1, self.visible_last + 1):
            key = openpyxl.utils.get_column_letter(c)
            self.assertEqual(actual.column_dimensions[key].width, expected.column_dimensions[key].width)

    def test_report_without_versions_is_recognized_but_readonly(self):
        wb = openpyxl.load_workbook(self.report_path)
        meta = wb["_sop_import"]
        self.assertEqual(meta["B3"].value, False)
        self.assertEqual(meta.max_row, 5)
        self.assertTrue(all(c.protection.locked for row in wb["SOP计划"] for c in row))

    def test_partial_annual_recalculates_visible_edits_and_preserves_offscreen(self):
        script = r'''
import {FileBlob,SpreadsheetFile} from '@oai/artifact-tool';
const wb=await SpreadsheetFile.importXlsx(await FileBlob.load(process.argv[1]));
const sheet=wb.worksheets.getItem('SOP计划'),meta=wb.worksheets.getItem('_sop_import');
const entries=meta.getUsedRange().values.filter(r=>r[0]==='plan');
const ids=[...new Set(entries.map(r=>r[1]))];
const first=entries.filter(r=>r[1]===ids[0]), row=Number(first[0][3].match(/\d+/)[0]);
sheet.getRange(first.find(r=>r[2]==='2026-08')[3]).values=[[25]];
wb.recalculate();
if(sheet.getRange('G'+row).values[0][0]!==315)throw new Error('Annual total did not change with visible month');
for(const e of first.filter(r=>r[2].startsWith('2026-')))sheet.getRange(e[3]).values=[[null]];
wb.recalculate();
if(sheet.getRange('G'+row).values[0][0]!==260)throw new Error('Offscreen annual amount was lost');
const totalRow=sheet.getUsedRange().values.findIndex(r=>r[0]==='总计')+1;
if(sheet.getRange('G'+totalRow).values[0][0]!==760)throw new Error('Annual grand total did not recalculate');
const zero=await SpreadsheetFile.importXlsx(await FileBlob.load(process.argv[2]));
const zs=zero.worksheets.getItem('SOP计划');
zs.getRange('D5').values=[[null]];
zero.recalculate();
if(![null,undefined,''].includes(zs.getRange('G5').values[0][0]))throw new Error('Cleared visible zero must return to blank');
if(zs.getRange('G6').values[0][0]!==0)throw new Error('An offscreen explicit zero must remain zero');
console.log('Partial annual input edits and offscreen totals verified');
'''
        result = subprocess.run([str(NODE), "--input-type=module", "-e", script, self.editable_path, self.zero_path],
                                cwd=ROOT / "scripts", capture_output=True, text=True, encoding="utf-8", timeout=120)
        self.assertEqual(result.returncode, 0, result.stderr)


if __name__ == "__main__":
    unittest.main()
