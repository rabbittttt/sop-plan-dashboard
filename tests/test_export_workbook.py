"""Exercise the real Excel download writer using disposable, synthetic records."""
import copy
import json
import sqlite3
import subprocess
import unittest
from contextlib import closing
from pathlib import Path

import openpyxl


ROOT = Path(__file__).resolve().parents[1]
NODE = Path.home() / ".cache/codex-runtimes/codex-primary-runtime/dependencies/node/bin/node.exe"
OUT = ROOT / "tests/artifacts/export-layout"


def payload(locks=True, versions=True):
    months = ["2026-08", "2026-10", "2026-12", "2027-02"]
    records = [
        ("车型汇总", "问界", "M5", "", "锁单数", [120, None, 0, None], None, "交车锁单原数，未分摊版本"),
        ("车型汇总", "问界", "M5", "", "SOP计划量", [30, 0, 70, None], 700, "部分月份缺项，仅已知合计"),
        ("版本明细", "问界", "M5", "同名版本", "锁单数", [None] * 4, None, "版本锁单缺失，未分摊"),
        ("版本明细", "问界", "M5", "同名版本", "锁单数", [None] * 4, None, "版本锁单缺失，同名版本第二项"),
        ("版本明细", "问界", "M5", "同名版本", "SOP计划量", [10, 0, 30, None], 300, "同名版本第一项"),
        ("版本明细", "问界", "M5", "同名版本", "SOP计划量", [20, None, 40, None], 400, "同名版本第二项"),
        ("车型汇总", "智界", "M5", "", "锁单数", [50, 0, None, None], None, "交车锁单原数，未分摊版本"),
        ("车型汇总", "智界", "M5", "", "SOP计划量", [5, 0, None, None], 100, "2027-02未适用"),
        ("版本明细", "智界", "M5", "=1+1", "锁单数", [None] * 4, None, "来源缺失"),
        ("版本明细", "智界", "M5", "=1+1", "SOP计划量", [5, 0, None, None], 100, "2027-02未适用"),
    ]
    records = [r for r in records if (locks or r[4] != "锁单数") and (versions or r[0] != "版本明细")]
    rows = [[r[0], r[1], r[2], *([r[3]] if versions else []), r[4], *r[5], r[7]] for r in records]
    return {"schema_version": 1, "current_month": "2026-10", "months": months,
            "include_locks": locks, "include_versions": versions, "scope": "测试范围",
            "headers": ["层级", "品牌", "车型", *(["版本"] if versions else []), "指标", *months, "数据说明"],
            "rows": rows, "annual_plans": [{"2026": r[6]} for r in records]}


def generate(body, name, preview=False):
    request = {**body, "output_path": str(OUT / f"{name}.xlsx")}
    if preview:
        request["preview_path"] = str(OUT / f"{name}.png")
    (OUT / f"{name}.json").write_text(json.dumps(request, ensure_ascii=False, indent=2), encoding="utf-8")
    result = subprocess.run([str(NODE), str(ROOT / "scripts/build_export_workbook.mjs")],
                            input=json.dumps(request, ensure_ascii=False), text=True, encoding="utf-8",
                            capture_output=True, cwd=ROOT, timeout=120)
    (OUT / f"{name}.log").write_text(result.stderr, encoding="utf-8")
    return result, request["output_path"]


def actual_input_payload():
    """Read the real maintenance template; do not manufacture a lookalike baseline."""
    path = ROOT / "input_file/SOP计划维护.xlsx"
    book = openpyxl.load_workbook(path, data_only=False)
    source = book["SOP计划"]
    cached = openpyxl.load_workbook(path, data_only=True)["SOP计划"]
    kind_column = next(c.column for c in source[4] if c.value == "row_kind")
    visible_last = kind_column - 2
    columns, annual, year = [], {}, None
    for c in range(4, visible_last + 1):
        if source.cell(3, c).value:
            year = str(source.cell(3, c).value).removesuffix("年")
        label = source.cell(4, c).value
        if label == "年度合计":
            annual[year] = c
        else:
            columns.append((f"{year}-{int(str(label).removesuffix('月')):02}", c))
    months = [m for m, _ in columns]
    current_month = next(m for m, c in columns if source.cell(4, c).fill.fgColor.rgb == "FF0871B8")
    rows, annual_plans, ids, known_counts, brand, model = [], [], [], [], "", ""
    for r in range(5, source.max_row + 1):
        kind = source.cell(r, kind_column).value
        if source.cell(r, 1).value and kind == "product":
            brand = source.cell(r, 1).value
        if kind == "product" and source.cell(r, 2).value:
            model = source.cell(r, 2).value
        if kind not in ("product", "model_total"):
            continue
        # Formula caches preserve missing values as blanks, distinct from real 0.
        quantities = [cached.cell(r, c).value for _, c in columns]
        quantities = [None if value in (None, "") else value for value in quantities]
        rows.append(["版本明细" if kind == "product" else "车型汇总", brand, model,
                     source.cell(r, 3).value or "", "SOP计划量", *quantities, ""])
        ids.append(source.cell(r, kind_column-1).value if kind == "product" else None)
        known_counts.append({y: sum(cached.cell(r, col).value not in (None, "") for m, col in columns if m.startswith(y + "-")) for y in annual})
        annual_plans.append({y: cached.cell(r, c).value or None for y, c in annual.items()})
        for y, c in annual.items():
            if cached.cell(r, c).value == 0:
                annual_plans[-1][y] = 0
    body = {"schema_version": 1, "current_month": current_month, "months": months,
            "include_locks": False, "include_versions": True, "scope": "全部品牌 · 全部车型",
            "headers": ["层级", "品牌", "车型", "版本", "指标", *months, "数据说明"],
            "rows": rows, "annual_plans": annual_plans}
    metadata = {row[0]: row[1] for row in book["_meta"].iter_rows(values_only=True) if row[0]}
    if metadata.get("schema_version") == 5:
        # Read the immutable source context that authored this exact real template.
        # A later schema must be compared with its matching writer contract.
        with closing(sqlite3.connect((ROOT / "data/sop.sqlite3").as_uri() + "?mode=ro", uri=True)) as db:
            saved = db.execute("SELECT payload FROM workbook_exports WHERE token=?", (metadata["token"],)).fetchone()
        expected = json.loads(saved[0])
        spec = expected["layout"]["sheets"][0]
        body["roundtrip"] = {"inline_schema": 2, "token": "synthetic-parity-token", "products": expected["products"],
            "row_product_ids": ids, "annual_known_counts": known_counts,
            "editable_months_by_product": spec["editable_months_by_product"],
            "manual_editable_months": spec["manual_editable_months"],
            "plan_revisions_by_product": spec["plan_revisions_by_product"]}
    return body, source, visible_last


@unittest.skipUnless(NODE.is_file(), "Bundled spreadsheet runtime unavailable")
class ExportWorkbookTest(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        OUT.mkdir(parents=True, exist_ok=True)
        cls.outputs = {}
        for locks, versions in [(True, True), (False, True), (True, False), (False, False)]:
            name = f"locks-{int(locks)}-versions-{int(versions)}"
            result, path = generate(payload(locks, versions), name, preview=True)
            if result.returncode:
                raise AssertionError(result.stderr)
            cls.outputs[(locks, versions)] = Path(path)
        body, cls.reference, cls.visible_last = actual_input_payload()
        cls.inline = body.get("roundtrip", {}).get("inline_schema") == 2
        result, path = generate(body, "actual-input-parity", preview=True)
        if result.returncode:
            raise AssertionError(result.stderr)
        cls.parity_path = Path(path)

    def test_full_plan_export_is_visibly_identical_to_actual_input(self):
        """Compare the entire real template, not a few matching blue swatches."""
        expected = self.reference
        actual_book = openpyxl.load_workbook(self.parity_path, data_only=False)
        self.assertEqual(actual_book.sheetnames, ["SOP计划", "_sop_import"] if self.inline else ["SOP计划"])
        actual = actual_book.active
        self.assertEqual(actual.max_column, self.visible_last + (3 if self.inline else 0), "Only trusted hidden identity columns may follow the visible grid")
        self.assertEqual(actual.max_row, expected.max_row)
        self.assertEqual(actual.freeze_panes, expected.freeze_panes)
        self.assertEqual(actual.sheet_view.showGridLines, expected.sheet_view.showGridLines)
        self.assertEqual(set(map(str, actual.merged_cells.ranges)), set(map(str, expected.merged_cells.ranges)))
        for c in range(1, self.visible_last + 1):
            letter = openpyxl.utils.get_column_letter(c)
            self.assertEqual(actual.column_dimensions[letter].width, expected.column_dimensions[letter].width, letter)
        # The only text differences explain that the downloaded copy does not sync.
        guidance_rows = {2} | {c.row for row in expected for c in row
                               if isinstance(c.value, str) and c.value.startswith("新增车型/版本：")}
        for r in range(1, expected.max_row + 1):
            self.assertEqual(actual.row_dimensions[r].height, expected.row_dimensions[r].height, f"row {r}")
            for c in range(1, self.visible_last + 1):
                left, right = actual.cell(r, c), expected.cell(r, c)
                if not (r in guidance_rows and c == 1):
                    self.assertEqual(left.value, right.value, f"value/formula {left.coordinate}")
                for attribute in ("font", "fill", "border", "alignment", "number_format"):
                    self.assertEqual(copy.copy(getattr(left, attribute)), copy.copy(getattr(right, attribute)),
                                     f"{attribute} {left.coordinate}")
        self.assertEqual(bool(actual.protection.sheet), not self.inline)
        if not self.inline:
            self.assertTrue(all(c.protection.locked for row in actual for c in row), "Legacy export is a read-only copy")

    def test_all_option_combinations_keep_values_style_and_formula_totals(self):
        reference = openpyxl.load_workbook(ROOT / "input_file/SOP计划维护.xlsx", data_only=False)["SOP计划"]
        for (locks, versions), path in self.outputs.items():
            with self.subTest(locks=locks, versions=versions):
                values = openpyxl.load_workbook(path, data_only=True)
                formulas = openpyxl.load_workbook(path, data_only=False)
                self.assertEqual(values.sheetnames, ["SOP计划"])
                s, f = values.active, formulas.active
                first = 3 + int(versions)
                self.assertEqual(f.freeze_panes, f"{openpyxl.utils.get_column_letter(first)}5")
                self.assertEqual(s["A3"].fill.fgColor.rgb, reference["A3"].fill.fgColor.rgb)
                self.assertEqual(s.cell(4, first).fill.fgColor.rgb, reference["D4"].fill.fgColor.rgb)
                self.assertEqual(s.cell(4, first + 3).value, "年度合计")
                self.assertEqual(s.cell(5, first + 3).fill.fgColor.rgb, reference["P5"].fill.fgColor.rgb)
                # Only chosen months are shown, with the full calendar year SOP after December.
                self.assertEqual([s.cell(4, first + i).value for i in range(5)], ["8月", "10月", "12月", "年度合计", "2月"])
                self.assertEqual(s.max_column, first + 4, "No extra metric/notes column")
                self.assertEqual(s["A1"].value, "SOP 跨年计划（2026-08—2027-02）")
                self.assertFalse(any(c.value in ("指标", "数据说明", "说明索引") for row in s for c in row))
                total = next(r for r in range(5, s.max_row + 1) if s.cell(r, 1).value == "总计")
                self.assertEqual([s.cell(total, first + i).value for i in range(5)], [35, 0, 70, 800, None])
                if locks:
                    lock_total = next(r for r in range(5, s.max_row + 1) if s.cell(r, 1).value == "锁单总计")
                    self.assertEqual([s.cell(lock_total, first + i).value for i in range(5)], [170, 0, 0, None, None])
                else:
                    self.assertFalse(any(isinstance(c.value, str) and "锁单" in c.value for row in s if row[0].row > 4 for c in row))
                model_rows = [r for r in range(5, s.max_row + 1) if s.cell(r, 2).value == "M5 合计"]
                self.assertEqual(len(model_rows), 2, "same model labels in different brands must remain separate")
                brand_rows = [r for r in range(5, s.max_row + 1) if s.cell(r, 2).value in ("问界 合计", "智界 合计")]
                for row in brand_rows:
                    self.assertIn(f"{openpyxl.utils.get_column_letter(first)}{row}", f.cell(total, first).value)
                self.assertEqual(s.cell(total, first).font.color.rgb, reference["P5"].font.color.rgb)
                self.assertEqual(s.row_dimensions[total].height, 28)
                self.assertEqual(s.row_dimensions[2].height, 29)
                self.assertFalse(s.sheet_view.showGridLines)
                self.assertTrue(s.protection.sheet)
                self.assertFalse(any(c.data_type == "e" for row in s for c in row))
                if versions:
                    self.assertEqual(s.cell(5, first).fill.fgColor.rgb, reference["D5"].fill.fgColor.rgb)
                    duplicate_rows = [r for r in range(5, s.max_row + 1) if s.cell(r, 3).value == "同名版本"]
                    self.assertEqual(len(duplicate_rows), 2)
                    if locks:
                        lock_rows = [r for r in range(5, s.max_row + 1) if s.cell(r, 3).value == "同名版本（锁单数）"]
                        self.assertEqual(len(lock_rows), 2)
                        self.assertLess(max(lock_rows), min(duplicate_rows))
                    literal = [c for row in f for c in row if c.value in ("=1+1", "=1+1（锁单数）")]
                    self.assertEqual(len(literal), 2 if locks else 1)
                    self.assertTrue(all(c.data_type == "s" for c in literal))
                else:
                    self.assertFalse(any(c.value in ("同名版本", "=1+1") for row in s for c in row))

    def test_safe_integer_contract_and_overflow(self):
        body = payload(False, False)
        body["rows"] = body["rows"][:1]
        body["annual_plans"] = body["annual_plans"][:1]
        body["rows"][0][4] = 2**53 - 1
        result, path = generate(body, "safe-integer")
        self.assertEqual(result.returncode, 0, result.stderr)
        s = openpyxl.load_workbook(path, data_only=True).active
        self.assertEqual(s["C5"].value, 2**53 - 1)
        self.assertEqual(s.column_dimensions["C"].width, 10, "Month width matches the input template")
        for index, number in enumerate([1.5, True, -1, 2**53]):
            bad = copy.deepcopy(body)
            bad["rows"][0][4] = number
            result, _ = generate(bad, f"invalid-{index}")
            self.assertNotEqual(result.returncode, 0)
        body["rows"].append(["车型汇总", "智界", "M5", "SOP计划量", 1, 0, None, None, ""])
        body["annual_plans"].append({"2026": None})
        result, _ = generate(body, "overflow")
        self.assertNotEqual(result.returncode, 0)
        self.assertIn("safe integer", result.stderr)

    def test_artifact_recalculation_and_reference_render(self):
        # Inspect the actual template and verify a changed source summary updates
        # derived brand/grand totals in memory. No production workbook is edited.
        script = r'''
import fs from 'node:fs/promises';
import {FileBlob,SpreadsheetFile} from '@oai/artifact-tool';
const root=process.argv[1], out=process.argv[2];
const template=await SpreadsheetFile.importXlsx(await FileBlob.load(root+'/input_file/SOP计划维护.xlsx'));
console.log((await template.inspect({kind:'region',sheetId:'SOP计划',range:'A3:P8',maxChars:1500,tableMaxRows:6,tableMaxCols:16})).ndjson);
const preview=await template.render({sheetName:'SOP计划',range:'A1:AC13',scale:1.2,format:'png'});
await fs.writeFile(out+'/reference.png',new Uint8Array(await preview.arrayBuffer()));
const wb=await SpreadsheetFile.importXlsx(await FileBlob.load(out+'/locks-0-versions-0.xlsx'));
const s=wb.worksheets.getItemAt(0), grid=s.getUsedRange().values;
const modelRow=grid.findIndex(row=>row[1]==='M5 合计')+1, totalRow=grid.findIndex(row=>row[0]==='总计')+1;
s.getRange(`C${modelRow}`).values=[[77]];
wb.recalculate();
if(s.getRange(`C${totalRow}`).values[0][0]!==82)throw new Error('Grand total failed to recalculate from model summary');
console.log((await wb.inspect({kind:'formula',sheetId:s.name,range:`A${totalRow}:H${totalRow}`,maxChars:1000})).ndjson);
'''
        result = subprocess.run([str(NODE), "--input-type=module", "-e", script, str(ROOT), str(OUT)],
                                cwd=ROOT / "scripts", capture_output=True, text=True, encoding="utf-8", timeout=120)
        (OUT / "reference-and-recalc.log").write_text(result.stdout + result.stderr, encoding="utf-8")
        self.assertEqual(result.returncode, 0, result.stderr)


if __name__ == "__main__":
    unittest.main()
