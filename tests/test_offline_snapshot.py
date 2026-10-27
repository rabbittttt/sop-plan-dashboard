"""The offline snapshot writer keeps frozen values and the established tabs."""

import json
import subprocess
import tempfile
import unittest
from pathlib import Path

import openpyxl

from sop_runtime import node_executable
from sop_xlsx_backup import write_snapshot_workbook


ROOT = Path(__file__).resolve().parents[1]


def frozen_payload(output):
    return {
        "output_path": str(output),
        "exported_at": "2026-10-22T13:00:00+08:00",
        "snapshot_data": {
            "snapshot": {"id": 7, "kind": "sunday", "label": "周日记录", "business_date": "2026-10-18",
                         "cutoff_at": "2026-10-18T24:00:00+08:00", "created_at": "2026-10-19T00:01:00+08:00"},
            "meta": {"months": ["2026-10", "2026-11"], "data_through": "2026-08",
                     "source_files": ["锁单来源.xlsx"]},
            "products": [{"id": "a", "brand": "问界", "model": "M5", "version": "Max"},
                         {"id": "b", "brand": "问界", "model": "M5", "version": "Ultra"}],
            "plans": [{"product_id": "a", "month": "2026-10", "sop": 0, "forecast": None,
                       "owner": "原负责人", "reason": "周日冻结", "notes": '=HYPERLINK("https://example.test")',
                       "revision": 2, "updated_at": "2026-10-18T23:59:59+08:00", "extra": {"a": 1}},
                      {"product_id": "b", "month": "2026-10", "sop": None, "forecast": 45,
                       "owner": "另一负责人", "reason": "", "notes": "", "revision": 1,
                       "updated_at": "2026-10-18T22:00:00+08:00"}],
        },
    }


class OfflineSnapshotWriterTest(unittest.TestCase):
    def test_frozen_values_exact_text_and_layout(self):
        with tempfile.TemporaryDirectory() as temporary:
            path = Path(temporary) / "snapshot.xlsx"
            write_snapshot_workbook(frozen_payload(path))
            book = openpyxl.load_workbook(path, data_only=False)
            try:
                self.assertEqual(book.sheetnames, ["SOP计划", "计划明细", "备份信息"])
                self.assertEqual(book["SOP计划"]["E10"].value, 0)
                self.assertIsNone(book["SOP计划"]["F10"].value)
                self.assertIsNone(book["SOP计划"]["E11"].value)
                self.assertEqual(book["SOP计划"]["B5"].value, "2026-10-18T24:00:00+08:00")
                self.assertEqual(book["计划明细"]["J6"].value, '=HYPERLINK("https://example.test")')
                self.assertEqual(book["计划明细"]["J6"].data_type, "s")
                self.assertEqual(book["计划明细"]["M6"].value, '{"a":1}')
                self.assertEqual(book["计划明细"]["L6"].value, "2026-10-18T23:59:59+08:00")
                self.assertEqual(book["备份信息"]["B18"].value, 2)
                self.assertEqual(book["SOP计划"].freeze_panes, "E10")
                self.assertEqual(book["计划明细"].freeze_panes, "F6")
                self.assertEqual(book["备份信息"].freeze_panes, "A6")
                self.assertTrue(all(cell.data_type != "f" for sheet in book for row in sheet for cell in row))
                self.assertNotIn(999, [cell.value for sheet in book for row in sheet for cell in row])
            finally:
                book.close()

    def test_rejects_duplicate_product_and_plan_keys(self):
        with tempfile.TemporaryDirectory() as temporary:
            path = Path(temporary) / "snapshot.xlsx"
            payload = frozen_payload(path)
            payload["snapshot_data"]["products"].append(dict(payload["snapshot_data"]["products"][0]))
            with self.assertRaisesRegex(ValueError, "unique"):
                write_snapshot_workbook(payload)
            self.assertFalse(path.exists())
            payload = frozen_payload(path)
            payload["snapshot_data"]["plans"].append(dict(payload["snapshot_data"]["plans"][0]))
            with self.assertRaisesRegex(ValueError, "Duplicate"):
                write_snapshot_workbook(payload)
            self.assertFalse(path.exists())

    def test_matches_node_writer_values_when_bundled_package_is_available(self):
        node = node_executable()
        package = ROOT / "scripts" / "node_modules" / "@oai" / "artifact-tool" / "package.json"
        if not node or not package.is_file():
            self.skipTest("Node spreadsheet package unavailable")
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            offline = root / "offline.xlsx"
            existing = root / "node.xlsx"
            payload = frozen_payload(offline)
            write_snapshot_workbook(payload)
            payload["output_path"] = str(existing)
            done = subprocess.run([node, str(ROOT / "scripts" / "build_snapshot_workbook.mjs")],
                                  input=json.dumps(payload, ensure_ascii=False), capture_output=True,
                                  text=True, encoding="utf-8", timeout=120, cwd=ROOT)
            self.assertEqual(done.returncode, 0, done.stderr[-1000:])
            a = openpyxl.load_workbook(offline, read_only=True, data_only=False)
            b = openpyxl.load_workbook(existing, read_only=True, data_only=False)
            try:
                for name in a.sheetnames:
                    self.assertEqual(list(a[name].values), list(b[name].values), name)
            finally:
                a.close()
                b.close()


if __name__ == "__main__":
    unittest.main()
