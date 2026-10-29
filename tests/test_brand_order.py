"""Brand ordering is shared by the catalog and both editable Excel writers."""
import json
import os
import tempfile
import unittest
from datetime import datetime
from pathlib import Path
from unittest.mock import patch

import openpyxl

from sop_brand import BRAND_ORDER, brand_order_key
from sop_roundtrip import PlanRoundtrip
from sop_store import SOPStore, TZ
from sop_workbook import WorkbookSync
from sop_xlsx_export import _layout, write_export_workbook
from sop_xlsx_plan import write_plan_workbook
from tests import test_inline_renderer as fixtures
from tests.test_workbook_sync import edit_cell


PRODUCTS = [{"id": str(index), "brand": brand, "model": f"M{index}", "version": "V1", "active": True}
            for index, brand in enumerate(("问界", "OtherZ", "尊界", "享界", "智界", "尚界", "OtherA"))]
PRODUCTS += [{"id": "z2", "brand": "智界", "model": "M4", "version": "V2", "active": True},
             {"id": "za", "brand": "智界", "model": "A", "version": "V1", "active": True}]
PLANS = [{"product_id": product["id"], "month": "2026-10", "sop": index + 1, "revision": 1}
         for index, product in enumerate(PRODUCTS)]
LOOKUP = {(plan["product_id"], plan["month"]): plan for plan in PLANS}
EXPECTED = [*BRAND_ORDER, "OtherA", "OtherZ"]


def payloads(products=PRODUCTS, **options):
    with patch.multiple(fixtures, PRODUCTS=products, PLANS=PLANS, LOOKUP=LOOKUP):
        return fixtures.input_payload(), fixtures.export_payload(**options)


def brands_from_sheet(sheet):
    return list(dict.fromkeys(cell.value for row in sheet.iter_rows(min_row=5, max_col=1)
                              for cell in row if cell.value in EXPECTED))


def assert_brand_blocks(test, sheet):
    """Each brand spans its models, metrics and subtotal exactly once in A."""
    areas = list(sheet.merged_cells.ranges)
    for brand in EXPECTED:
        labels = [row[0].row for row in sheet.iter_rows(min_row=5, max_col=1)
                  if row[0].value == brand]
        test.assertEqual(len(labels), 1, brand)
        first = labels[0]
        last = next(row[1].row for row in sheet.iter_rows(min_row=first, max_col=2)
                    if row[1].value == brand + " 合计")
        test.assertIn(f"A{first}:A{last}", {str(area) for area in areas}, brand)
        intersecting = [area for area in areas if area.min_col == 1
                        and area.min_row >= 5 and area.min_row <= last and area.max_row >= first]
        test.assertEqual(len(intersecting), 1, brand)


class BrandOrderTest(unittest.TestCase):
    def setUp(self):
        temporary = tempfile.TemporaryDirectory(prefix="sop-brand-order-")
        self.addCleanup(temporary.cleanup)
        self.root = Path(temporary.name)
        source = self.root / "sales.json"
        source.write_text(json.dumps({"meta": {}, "products": PRODUCTS, "sales": []}, ensure_ascii=False),
                          encoding="utf-8")
        self.store = SOPStore(self.root / "test.sqlite3", source,
                              lambda: datetime(2026, 10, 28, 12, tzinfo=TZ))
        self.addCleanup(self.store.close)

    def test_catalog_orders_all_brands_and_retains_model_version_sorting(self):
        actual = self.store.list_products()["items"]
        self.assertEqual(list(dict.fromkeys(item["brand"] for item in actual)), EXPECTED)
        expected = sorted(PRODUCTS, key=lambda item: (brand_order_key(item["brand"]),
                                                     item["model"], item["version"], item["id"]))
        self.assertEqual([item["id"] for item in actual], [item["id"] for item in expected])

    def test_new_maintenance_and_download_layouts_share_brand_order(self):
        maintenance, _ = payloads()
        spec = maintenance["layout"]["sheets"][0]
        self.assertEqual([row["labels"][1] for row in spec["rows"] if row["kind"] == "brand_total"],
                         [brand + " 合计" for brand in EXPECTED])
        expected_ids = [row["product_id"] for row in spec["rows"] if row["kind"] == "product"]
        for locks in (False, True):
            for versions in (False, True):
                with self.subTest(locks=locks, versions=versions):
                    _, body = payloads(locks=locks, versions=versions)
                    layout = _layout(body)
                    labels = [row["labels"][1] for row in layout["rows"]
                              if row["kind"] == "brand_total" and row["metric"] == "SOP计划量"]
                    self.assertEqual(labels, [brand + " 合计" for brand in EXPECTED])
                    if versions:
                        ids = [row["product_id"] for row in layout["rows"]
                               if row["kind"] == "product" and row["metric"] == "SOP计划量"]
                        self.assertEqual(ids, expected_ids)
                    path = self.root / f"layout-{locks}-{versions}.xlsx"
                    write_export_workbook({**body, "output_path": str(path)})
                    book = openpyxl.load_workbook(path)
                    try:
                        assert_brand_blocks(self, book["SOP计划"])
                    finally:
                        book.close()

    def test_python_maintenance_export_parity_keeps_values_with_product_identity(self):
        maintenance, exported = payloads()
        paths = [self.root / "maintenance.xlsx", self.root / "export.xlsx"]
        write_plan_workbook({**maintenance, "output_path": str(paths[0])})
        write_export_workbook({**exported, "output_path": str(paths[1])})
        books = [openpyxl.load_workbook(path, data_only=True) for path in paths]
        try:
            left, right = (book["SOP计划"] for book in books)
            self.assertEqual(brands_from_sheet(left), EXPECTED)
            self.assertEqual(brands_from_sheet(right), EXPECTED)
            assert_brand_blocks(self, left)
            assert_brand_blocks(self, right)
            visible = maintenance["layout"]["sheets"][0]["visible_last_column"]
            self.assertEqual(list(left.iter_rows(max_col=visible, values_only=True)),
                             list(right.iter_rows(max_col=visible, values_only=True)))
            for row in maintenance["layout"]["sheets"][0]["rows"]:
                if row["kind"] == "product":
                    self.assertEqual(left.cell(row["row"], 13).value,
                                     LOOKUP[(row["product_id"], "2026-10")]["sop"])
        finally:
            for book in books:
                book.close()

    def test_prior_order_imports_unsynced_changes_before_refresh(self):
        path = self.root / "maintenance.xlsx"
        with patch.dict(os.environ, {"SOP_EXCEL_ENGINE": "python"}):
            with patch("sop_store.brand_order_key", lambda brand: brand), \
                    patch("sop_workbook.brand_order_key", lambda brand: brand), \
                    patch("sop_workbook.PRESENTATION_VERSION", 3):
                previous = WorkbookSync(self.store, path)
                self.assertEqual(previous.sync()["status"], "synced")
                token, context, *_ = previous._parse(path.read_bytes())
                cell = next(item["cell"] for item in context["layout"]["cells"]
                            if item["product_id"] == "0" and item["month"] == "2026-10")
            edit_cell(path, cell, 41)
            current = WorkbookSync(self.store, path)
            status = current.sync()
            self.assertEqual(status["status"], "synced", status)
            self.assertNotEqual(current._parse(path.read_bytes())[0], token)
        self.assertEqual(self.store.db.execute(
            "SELECT sop FROM plans WHERE product_id='0' AND month='2026-10'").fetchone()[0], 41)
        book = openpyxl.load_workbook(path, data_only=True)
        try:
            self.assertEqual(brands_from_sheet(book["SOP计划"]), EXPECTED)
        finally:
            book.close()

    def test_filtered_download_import_preserves_excluded_wenjie(self):
        self.store.save_plans({"owner": "test", "changes": [
            {**plan, "forecast": None, "revision": 0} for plan in PLANS]})
        selected = [product for product in PRODUCTS if product["brand"] != "问界"]
        _, body = payloads(products=selected)
        body["selected_product_ids"] = [product["id"] for product in selected]
        body["row_product_ids"] = body.pop("roundtrip")["row_product_ids"]
        service = PlanRoundtrip(self.store)
        clean, context = service.prepare(body)
        path = self.root / "selected.xlsx"
        write_export_workbook({**clean, "roundtrip": context, "output_path": str(path)})
        service.register(path, context)
        book = openpyxl.load_workbook(path)
        try:
            spec = fixtures.manifest(book)
            row = next(item["row"] for item in spec["rows"] if item.get("product_id") == "4")
            self.assertNotIn("问界", brands_from_sheet(book["SOP计划"]))
        finally:
            book.close()
        edit_cell(path, f"M{row}", 55)
        result = service.import_content(path.read_bytes())
        self.assertTrue(result["ok"], result)
        self.assertEqual(result["imported"], 1)
        wenjie = next(item for item in self.store.list_products()["items"] if item["id"] == "0")
        self.assertTrue(wenjie["active"])
        self.assertEqual(self.store.db.execute(
            "SELECT sop FROM plans WHERE product_id='0' AND month='2026-10'").fetchone()[0], 1)

    @unittest.skipUnless(fixtures.NODE.is_file(), "Bundled spreadsheet runtime unavailable")
    def test_node_export_uses_same_order_and_identity_mapping(self):
        _, body = payloads(locks=True)
        with patch.object(fixtures, "OUT", self.root):
            path = fixtures.generate("build_export_workbook.mjs", body, "brand-order")
        book = openpyxl.load_workbook(path, data_only=True)
        try:
            sheet = book["SOP计划"]
            self.assertEqual(brands_from_sheet(sheet), EXPECTED)
            assert_brand_blocks(self, sheet)
            spec = fixtures.manifest(book)
            for row in spec["rows"]:
                if row["kind"] == "product" and row.get("metric") == "SOP计划量":
                    self.assertEqual(sheet.cell(row["row"], 13).value,
                                     LOOKUP[(row["product_id"], "2026-10")]["sop"])
        finally:
            book.close()


if __name__ == "__main__":
    unittest.main()
