"""Insert/delete/rename table rows against immutable download identities."""
import copy
import json
import shutil
import tempfile
import unittest
import zipfile
import xml.etree.ElementTree as ET
from datetime import datetime
from pathlib import Path
import openpyxl

from sop_roundtrip import PlanRoundtrip
from sop_export import PlanExport
from sop_store import SOPStore, TZ, ValidationError
from sop_workbook import WorkbookSync
from tests.test_roundtrip import download_fixture, export_body
from tests.test_workbook_sync import NS, col_name, edit_cell, fixture_writer


def rewrite_grid(path, grid, sheet=1):
    """Test-only saved-user-edit simulation; does not author user workbooks."""
    with zipfile.ZipFile(path) as archive:
        files = {name: archive.read(name) for name in archive.namelist()}
    root = ET.Element(f"{{{NS}}}worksheet")
    data = ET.SubElement(root, f"{{{NS}}}sheetData")
    for r, values in enumerate(grid, 1):
        line = ET.SubElement(data, f"{{{NS}}}row", {"r": str(r)})
        for c, value in enumerate(values, 1):
            cell = ET.SubElement(line, f"{{{NS}}}c", {"r": f"{col_name(c)}{r}"})
            if isinstance(value, tuple) and value[0] == "formula":
                ET.SubElement(cell, f"{{{NS}}}f").text = value[1]
            elif isinstance(value, str):
                cell.set("t", "inlineStr")
                ET.SubElement(ET.SubElement(cell, f"{{{NS}}}is"), f"{{{NS}}}t").text = value
            elif isinstance(value, bool):
                cell.set("t", "b")
                ET.SubElement(cell, f"{{{NS}}}v").text = str(int(value))
            elif value is not None:
                ET.SubElement(cell, f"{{{NS}}}v").text = str(value)
    files[f"xl/worksheets/sheet{sheet}.xml"] = ET.tostring(root, encoding="utf-8")
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as archive:
        for name, content in files.items():
            archive.writestr(name, content)


def inline_fixture(path, clean, context):
    download_fixture(path, clean, context)
    columns = [{"column": 4, "kind": "month", "month": "2026-10"},
               {"column": 5, "kind": "month", "month": "2026-12"},
               {"column": 6, "kind": "year_total", "year": "2026"}]
    headers = [["品牌", "车型", "版本", "2026年", None, None, None, None, None],
               [None, None, None, "10月", "12月", "年度合计", "product_id", "row_kind", "row_key"]]
    grid = [["SOP计划", *([None] * 8)], ["可插删明细行", *([None] * 8)], *copy.deepcopy(headers)]
    specs = []
    metric_col = 4 if clean["include_versions"] else 3
    for record, pid in zip(clean["rows"], context["row_product_ids"]):
        number = len(grid) + 1
        metric = record[metric_col]
        kind = "product" if pid is not None else "model_total"
        key = (("plan:" if metric == "SOP计划量" else "lock:") + pid) if pid else f"row:{number}"
        labels = record[1:4]
        if pid and metric == "锁单数":
            labels = [*labels[:2], labels[2] + "（锁单数）"]
        quantities = record[metric_col+1:metric_col+3]
        grid.append([*labels, *quantities, ("formula", f'SUM(D{number}:E{number})'), pid, kind, key])
        specs.append({"row": number, "key": key, "kind": kind, "product_id": pid,
                      "metric": metric, "labels": labels, "values": dict(zip(clean["months"], quantities))})
    manifest = {"columns": columns, "products": context["products"], "rows": specs,
                "product_column": 7, "kind_column": 8, "key_column": 9, "last_column": 9,
                "headers": headers, "editable_months_by_product": context["editable_months_by_product"],
                "manual_editable_months": context["manual_editable_months"],
                "plan_revisions_by_product": context["plan_revisions_by_product"]}
    meta = [["SOP_IMPORT", 2], ["token", context["token"]], ["include_versions", clean["include_versions"]],
            ["sheet", "SOP计划"], ["manifest_json", None], [1, json.dumps(manifest, ensure_ascii=False)]]
    rewrite_grid(path, grid)
    rewrite_grid(path, meta, sheet=2)
    return grid, manifest


class InlineRoundtripTest(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)
        source = self.root / "sales.json"
        source.write_text(json.dumps({"meta": {}, "products": [
            {"id": "a", "brand": "A", "model": "M", "version": "V1"},
            {"id": "b", "brand": "A", "model": "M", "version": "V2"}], "sales": []}), encoding="utf-8")
        self.store = SOPStore(self.root / "test.sqlite3", source, lambda: datetime(2026, 10, 26, 12, tzinfo=TZ))
        self.service = PlanRoundtrip(self.store)
        self.path = self.root / "export.xlsx"
        self.save("a", "2026-10", 10)
        self.save("b", "2026-12", 20)

    def tearDown(self):
        self.store.close()
        self.temp.cleanup()

    def save(self, pid, month, sop):
        row = self.store.db.execute("SELECT revision FROM plans WHERE product_id=? AND month=?", (pid, month)).fetchone()
        self.store.save_plans({"owner": "网页", "changes": [{"product_id": pid, "month": month, "sop": sop,
                              "forecast": 99, "revision": row[0] if row else 0}]})

    def export(self, selected=("a", "b")):
        body = export_body(selected=selected)
        for row, pid in zip(body["rows"], body["row_product_ids"]):
            if pid:
                row[3] = "V1" if pid == "a" else "V2"
        clean, context = self.service.prepare(body)
        self.grid, self.manifest = inline_fixture(self.path, clean, context)
        self.service.register(self.path, context)
        self.original = self.path.read_bytes()

    def row(self, pid, metric="SOP计划量"):
        key = ("plan:" if metric == "SOP计划量" else "lock:") + pid
        return next(row for row in self.grid[4:] if row[8] == key)

    def apply(self, choice=None):
        rewrite_grid(self.path, self.grid)
        return self.service.import_content(self.path.read_bytes(), choice)

    def product(self, pid):
        return next(p for p in self.store.list_products()["items"] if p["id"] == pid)

    def quantity(self, pid, month):
        return self.store.db.execute("SELECT sop FROM plans WHERE product_id=? AND month=?", (pid, month)).fetchone()[0]

    def test_insert_delete_rename_and_old_copy_do_not_undo_changes(self):
        self.export()
        self.row("a")[2:4] = ["2027款 改名版", 11]
        self.grid.remove(self.row("b"))
        self.grid.insert(5, ["新品牌", "新车型", "新版本", 0, None, None, None, None, None])
        result = self.apply()
        self.assertTrue(result["ok"], result)
        self.assertEqual(result["created"], 1)
        self.assertEqual(result["deactivated"], 1)
        self.assertEqual(self.product("a")["version"], "改名版")
        self.assertFalse(self.product("b")["active"])
        self.assertEqual(self.quantity("b", "2026-12"), 20)
        self.assertEqual(self.quantity("a", "2026-10"), 11)
        self.assertTrue(self.service.import_content(self.original)["ok"])
        self.assertFalse(self.product("b")["active"])
        self.assertEqual(self.product("a")["version"], "改名版")

    def test_copy_row_new_name_gets_new_id_and_duplicate_ambiguity_rejected(self):
        self.export()
        duplicate = copy.deepcopy(self.row("a"))
        duplicate[2] = "新增V3"
        self.grid.insert(7, duplicate)
        self.assertEqual(self.apply()["created"], 1)
        self.assertEqual(len(self.store.list_products()["items"]), 3)
        for index, row in enumerate(self.grid):
            if len(row) > 8 and row[8] == "plan:a":
                row[2] = "两行都改名" + str(index)
        with self.assertRaises(ValidationError):
            self.apply()

    def test_plain_pasted_existing_labels_keep_identity(self):
        self.export()
        self.row("a")[6:] = [None, None, None]
        self.row("b")[4] = 25
        result = self.apply()
        self.assertEqual(result["created"], 0)
        self.assertEqual(result["deactivated"], 0)
        self.assertEqual(self.quantity("b", "2026-12"), 25)

    def test_delete_conflicts_with_later_web_plan_and_preserves_history(self):
        self.export()
        self.grid.remove(self.row("b"))
        self.save("b", "2026-12", 30)
        blocked = self.apply()
        self.assertFalse(blocked["ok"])
        self.assertTrue(self.product("b")["active"])
        selected = self.apply("workbook")
        self.assertTrue(selected["ok"])
        self.assertFalse(self.product("b")["active"])
        self.assertEqual(self.quantity("b", "2026-12"), 30)

    def test_catalog_conflict_is_atomic_with_plan_edits(self):
        self.export()
        original = self.product("a")
        self.store.save_products({"changes": [{"id": "a", "revision": original["revision"], "version": "网页新名"}]})
        self.row("a")[2:4] = ["表格新名", 44]
        blocked = self.apply()
        self.assertFalse(blocked["ok"])
        self.assertEqual(self.quantity("a", "2026-10"), 10)
        self.assertTrue(self.apply("db")["ok"])
        self.assertEqual(self.product("a")["version"], "网页新名")
        self.assertEqual(self.quantity("a", "2026-10"), 44)

    def test_header_unknown_identity_formula_and_lock_edits_are_rejected(self):
        self.export()
        baseline = copy.deepcopy(self.grid)
        mutations = [lambda: self.grid[3].__setitem__(3, "11月"),
                     lambda: self.row("a").__setitem__(6, "unknown"),
                     lambda: self.row("a").__setitem__(3, ("formula", "1+1")),
                     lambda: self.row("a", "锁单数").__setitem__(3, 88),
                     lambda: next(row for row in self.grid[4:] if row[7] == "model_total").__setitem__(3, 88)]
        for mutation in mutations:
            self.grid = copy.deepcopy(baseline)
            mutation()
            with self.assertRaises(ValidationError):
                self.apply()
            self.assertEqual(self.quantity("a", "2026-10"), 10)

    def test_partial_scope_deletion_only_deactivates_exported_product(self):
        self.export(selected=("a",))
        self.grid.remove(self.row("a"))
        self.assertTrue(self.apply()["ok"])
        self.assertFalse(self.product("a")["active"])
        self.assertTrue(self.product("b")["active"])

    def test_imported_new_row_cannot_be_renamed_again_without_fresh_identity(self):
        self.export()
        inserted = ["B", "New", "V3", 0, None, None, None, None, None]
        self.grid.insert(5, inserted)
        first = self.apply()
        self.assertEqual(first["created"], 1)
        self.assertIn("重新导出", first.get("warning", ""))
        inserted[2] = "V4"
        with self.assertRaisesRegex(ValidationError, "重新导出"):
            self.apply()
        self.assertEqual(len(self.store.list_products()["items"]), 3)
        self.assertTrue(self.service.import_content(self.original)["ok"])
        self.assertEqual(len(self.store.list_products()["items"]), 3)

    def test_web_rename_of_imported_new_row_cannot_recreate_old_name(self):
        self.export()
        self.grid.insert(5, ["B", "New", "V3", 0, None, None, None, None, None])
        self.assertEqual(self.apply()["created"], 1)
        added = next(p for p in self.store.list_products()["items"] if p["id"] not in {"a", "b"})
        self.store.save_products({"changes": [{"id": added["id"], "revision": added["revision"], "version": "网页改名"}]})
        with self.assertRaisesRegex(ValidationError, "重新导出"):
            self.apply()
        self.assertEqual(len(self.store.list_products()["items"]), 3)

    @unittest.skipUnless((Path.home() / ".cache/codex-runtimes/codex-primary-runtime/dependencies/node/bin/node.exe").is_file(), "Bundled spreadsheet runtime unavailable")
    def test_real_export_without_primary_workbook_can_be_imported(self):
        sync = WorkbookSync(self.store, self.root / "primary.xlsx", writer=fixture_writer)
        self.service = PlanRoundtrip(self.store, sync)
        self.assertEqual(sync.sync()["status"], "synced")
        sync.path.unlink()

        body = export_body()
        for row, pid in zip(body["rows"], body["row_product_ids"]):
            if pid:
                row[3] = "V1" if pid == "a" else "V2"
        clean, context = self.service.prepare(body)
        with PlanExport().generate(clean, roundtrip=context) as result:
            self.service.register(result.path, context)
            shutil.copy2(result.path, self.path)
        self.assertFalse(sync.path.exists())
        book = openpyxl.load_workbook(self.path, read_only=True, data_only=True)
        try:
            row = next(index for index, values in enumerate(book["SOP计划"].iter_rows(values_only=True), 1)
                       if "plan:a" in values)
            self.assertEqual(book["SOP计划"][f"D{row}"].value, 10)
        finally:
            book.close()
        edit_cell(self.path, f"D{row}", 41)
        imported = self.service.import_content(self.path.read_bytes())
        self.assertTrue(imported["ok"], imported)
        self.assertEqual(self.quantity("a", "2026-10"), 41)

    @unittest.skipUnless((Path.home() / ".cache/codex-runtimes/codex-primary-runtime/dependencies/node/bin/node.exe").is_file(), "Bundled spreadsheet runtime unavailable")
    def test_real_renderer_merged_labels_import_plan_edit(self):
        body = export_body()
        for row, pid in zip(body["rows"], body["row_product_ids"]):
            if pid:
                row[3] = "V1" if pid == "a" else "V2"
                if row[4] == "锁单数":
                    row[5] = 40 if pid == "a" else 60
        clean, context = self.service.prepare(body)
        with PlanExport().generate(clean, roundtrip=context) as result:
            self.service.register(result.path, context)
            shutil.copy2(result.path, self.path)
        with zipfile.ZipFile(self.path) as archive:
            root = ET.fromstring(archive.read("xl/worksheets/sheet1.xml"))
        merges = {area.attrib["ref"] for area in root.findall(f".//{{{NS}}}mergeCell")}
        book = openpyxl.load_workbook(self.path, read_only=True, data_only=False)
        rows = list(book["SOP计划"].iter_rows(values_only=True))
        first = next(index for index, row in enumerate(rows, 1) if "plan:a" in row)
        second = next(index for index, row in enumerate(rows, 1) if "plan:b" in row)
        model_lock = next(row for row in rows if "model_total" in row and row[1] == "M 锁单 合计")
        version_locks = [row for row in rows if any(isinstance(value, str) and value.startswith("lock:") for value in row)]
        self.assertEqual(model_lock[3], 999)
        self.assertEqual([row[3] for row in version_locks], [40, 60])
        brand_first = next(index for index, row in enumerate(rows, 1) if row[0] == "A")
        brand_last = next(index for index, row in enumerate(rows, 1) if row[1] == "A 合计")
        self.assertIn(f"A{brand_first}:A{brand_last}", merges)
        self.assertIn(f"B{first}:B{second}", merges)
        self.assertIsNone(rows[second - 1][0])
        self.assertIsNone(rows[second - 1][1])
        book.close()
        edit_cell(self.path, f"D{second}", 37)
        result = self.service.import_content(self.path.read_bytes())
        self.assertTrue(result["ok"], result)
        self.assertEqual(self.quantity("b", "2026-10"), 37)

    @unittest.skipUnless((Path.home() / ".cache/codex-runtimes/codex-primary-runtime/dependencies/node/bin/node.exe").is_file(), "Bundled spreadsheet runtime unavailable")
    def test_real_renderer_register_and_structural_roundtrip(self):
        body = export_body()
        for row, pid in zip(body["rows"], body["row_product_ids"]):
            if pid:
                row[3] = "V1" if pid == "a" else "V2"
        clean, context = self.service.prepare(body)
        with PlanExport().generate(clean, roundtrip=context) as result:
            self.service.register(result.path, context)
            shutil.copy2(result.path, self.path)
        book = openpyxl.load_workbook(self.path, data_only=False)
        self.grid = [[("formula", cell.value[1:]) if cell.data_type == "f" else cell.value for cell in row]
                     for row in book["SOP计划"]]
        self.assertEqual(book["_sop_import"]["B1"].value, 2)
        self.assertFalse(book["SOP计划"].protection.sheet)
        book.close()
        self.row("a")[2:4] = ["直接改名", 31]
        self.grid.remove(self.row("b"))
        self.grid.insert(5, ["B", "New", "V3", 0, None, None, None, None, None])
        result = self.apply()
        self.assertTrue(result["ok"], result)
        self.assertEqual(result["created"], 1)
        self.assertEqual(result["deactivated"], 1)
        self.assertEqual(self.product("a")["version"], "直接改名")
        self.assertEqual(self.quantity("a", "2026-10"), 31)
        self.assertEqual(self.quantity("b", "2026-12"), 20)


if __name__ == "__main__":
    unittest.main()
