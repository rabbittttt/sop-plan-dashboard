"""Readonly generation projections share real plans across Excel view changes."""
import json
import tempfile
import unittest
from datetime import datetime
from pathlib import Path
from unittest.mock import patch

import openpyxl

from sop_store import SOPStore, TZ, ValidationError
from sop_workbook import WorkbookSync
from sop_roundtrip import PlanRoundtrip
from sop_xlsx_plan import write_plan_workbook
from sop_xlsx_export import write_export_workbook
from tests import test_inline_renderer as fixtures
from tests.test_workbook_sync import edit_cell


class GenerationExcelTest(unittest.TestCase):
    def setUp(self):
        temp = tempfile.TemporaryDirectory(prefix="sop-generation-excel-")
        self.addCleanup(temp.cleanup)
        self.root = Path(temp.name)
        self.config = self.root / "shared_config.py"
        self.config.write_text("MERGE_GROUPS = {'智界 Combo 2026款汇总': ['智界 A 2025款', '智界 B 2026款']}\n", encoding="utf-8")
        rules = patch("sop_source_rules.DEFAULT_CONFIG_PATH", self.config)
        rules.start(); self.addCleanup(rules.stop)
        products = [{"id": pid, "brand": "智界", "model": model, "version": "Max", "generation": generation}
                    for pid, model, generation in [("a", "A", "智界 A 2025款"), ("b", "B", "智界 B 2026款"),
                                                  ("c", "C", "智界 C 2026款"), ("raw", "Combo", "智界 Combo 2026款汇总")]]
        source = self.root / "sales.json"
        source.write_text(json.dumps({"meta": {}, "products": products, "sales": [], "model_sales": []}), encoding="utf-8")
        self.store = SOPStore(self.root / "test.sqlite3", source, lambda: datetime(2026, 10, 28, 12, tzinfo=TZ))
        self.addCleanup(self.store.close)
        self.store.save_plans({"owner": "test", "changes": [
            {"product_id": pid, "month": "2026-10", "sop": value, "forecast": None, "revision": 0}
            for pid, value in [("a", 10), ("b", 20), ("c", 3)]]})
        self.path = self.root / "maintenance.xlsx"
        self.sync = WorkbookSync(self.store, self.path, writer=write_plan_workbook)

    def summary(self):
        self.store.update_settings({"generation_mode": "summary"})
        sales = self.store.read_sales()
        self.aggregate = next(p["id"] for p in sales["products"] if p.get("read_only"))
        return sales

    def spec(self):
        return self.sync._parse(self.path.read_bytes())[1]["layout"]["sheets"][0]

    def download(self, name="export", writer=write_export_workbook):
        sales = self.store.read_sales()
        plans = self.store.view_plans(sales)
        lookup = {(p["product_id"], p["month"]): p for p in plans}
        with patch.multiple(fixtures, PRODUCTS=sales["products"], PLANS=plans, LOOKUP=lookup):
            body = fixtures.export_payload(months=["2026-10"], locks=True)
        body["selected_product_ids"] = [p["id"] for p in sales["products"]]
        body["row_product_ids"] = body.pop("roundtrip")["row_product_ids"]
        body["generation_rule_signature"] = sales["generation_view"]["rule_signature"]
        body["generation_mode"] = sales["generation_view"]["mode"]
        service = PlanRoundtrip(self.store)
        clean, context = service.prepare(body)
        path = self.root / (name + ".xlsx")
        writer({**clean, "roundtrip": context, "output_path": str(path)})
        service.register(path, context)
        return service, path, body

    def test_maintenance_switches_views_preserving_pending_detail_edits(self):
        self.assertEqual(self.sync.sync()["status"], "synced")
        first = self.spec()
        row_a = next(r["row"] for r in first["rows"] if r.get("product_id") == "a")
        edit_cell(self.path, f"M{row_a}", 11)
        self.summary()
        self.assertEqual(self.sync.sync()["status"], "synced")
        spec = self.spec()
        row = next(r["row"] for r in spec["rows"] if r.get("product_id") == self.aggregate)
        self.assertEqual(spec["editable_months_by_product"][self.aggregate], [])
        book = openpyxl.load_workbook(self.path, data_only=True)
        self.assertEqual(book["SOP计划"][f"M{row}"].value, 31)
        self.assertTrue(book["SOP计划"][f"M{row}"].protection.locked)
        self.assertTrue(book["SOP计划"][f"C{row}"].protection.locked)
        book.close()
        self.store.update_settings({"generation_mode": "detail"})
        self.assertEqual(self.sync.sync()["status"], "synced")
        self.assertEqual({p["id"] for p in self.spec()["products"]}, {"a", "b", "c"})
        self.assertEqual(self.store.db.execute("SELECT sop FROM plans WHERE product_id='a'").fetchone()[0], 11)
        self.assertEqual(self.store.db.execute("SELECT COUNT(*) FROM plans WHERE product_id LIKE 'group_%'").fetchone()[0], 0)

    def test_export_derives_values_and_rejects_projection_changes(self):
        self.summary()
        service, path, _ = self.download()
        original = path.read_bytes()
        book = openpyxl.load_workbook(path)
        spec = fixtures.manifest(book)
        row = next(r["row"] for r in spec["rows"] if r.get("product_id") == self.aggregate and r.get("metric") == "SOP计划量")
        plain = next(r["row"] for r in spec["rows"] if r.get("product_id") == "c" and r.get("metric") == "SOP计划量")
        self.assertEqual(book["SOP计划"][f"D{row}"].value, 30)
        self.assertTrue(book["SOP计划"][f"D{row}"].protection.locked)
        self.assertTrue(book["SOP计划"][f"C{row}"].protection.locked)
        book.close()
        for field, value in [("D", 999), ("C", "Changed")]:
            with self.subTest(field=field):
                path.write_bytes(original)
                edit_cell(path, f"{field}{row}", value)
                with self.assertRaises(ValidationError):
                    service.import_content(path.read_bytes())
        path.write_bytes(original)
        book = openpyxl.load_workbook(path); book["SOP计划"].delete_rows(row); book.save(path); book.close()
        with self.assertRaisesRegex(ValidationError, "不能删除"):
            service.import_content(path.read_bytes())
        path.write_bytes(original)
        edit_cell(path, f"D{plain}", 8)
        self.store.update_settings({"generation_mode": "detail"})
        self.assertEqual(service.import_content(path.read_bytes())["imported"], 1)
        self.assertEqual(self.store.db.execute("SELECT sop FROM plans WHERE product_id='c'").fetchone()[0], 8)
        self.assertEqual(self.store.db.execute("SELECT COUNT(*) FROM plans WHERE product_id LIKE 'group_%'").fetchone()[0], 0)

    def test_detail_download_imports_after_switch_and_updates_derived_total(self):
        service, path, _ = self.download()
        book = openpyxl.load_workbook(path)
        row = next(r["row"] for r in fixtures.manifest(book)["rows"] if r.get("product_id") == "a" and r.get("metric") == "SOP计划量")
        book.close()
        self.summary()
        edit_cell(path, f"D{row}", 25)
        self.assertEqual(service.import_content(path.read_bytes())["imported"], 1)
        derived = next(r for r in self.store.view_plans() if r["product_id"] == self.aggregate)
        self.assertEqual(derived["sop"], 45)

    def test_changed_rule_rejects_stale_export_and_refreshes_old_readonly_maintenance(self):
        self.summary()
        _, _, body = self.download()
        self.assertEqual(self.sync.sync()["status"], "synced")
        self.config.write_text("MERGE_GROUPS = {'智界 Combo 2026款汇总': ['智界 A 2025款']}\n", encoding="utf-8")
        with self.assertRaisesRegex(ValidationError, "规则或展示模式已变化"):
            PlanRoundtrip(self.store).prepare(body)
        self.assertEqual(self.sync.sync()["status"], "synced")
        self.assertEqual({p["model"] for p in self.spec()["products"]}, {"Combo", "B", "C"})

    @unittest.skipUnless(fixtures.NODE.is_file(), "Bundled spreadsheet runtime unavailable")
    def test_node_writer_keeps_aggregate_cells_readonly_and_roundtrippable(self):
        self.summary()
        payload, _ = self.sync._context()
        with patch.object(fixtures, "OUT", self.root):
            maintenance = fixtures.generate("build_plan_workbook.mjs",
                {**payload, "token": "node-maintenance", "generated_at": "2026-10-28T12:00:00+08:00"}, "node-maintenance")
        row = next(r["row"] for r in payload["layout"]["sheets"][0]["rows"] if r.get("product_id") == self.aggregate)
        book = openpyxl.load_workbook(maintenance)
        self.assertEqual(book["SOP计划"][f"M{row}"].value, 30)
        self.assertTrue(book["SOP计划"][f"M{row}"].protection.locked)
        self.assertTrue(book["SOP计划"][f"C{row}"].protection.locked)
        book.close()
        def writer(payload):
            with patch.object(fixtures, "OUT", self.root):
                fixtures.generate("build_export_workbook.mjs", payload, "node-export")
        service, path, _ = self.download("node-export", writer)
        book = openpyxl.load_workbook(path)
        row = next(r["row"] for r in fixtures.manifest(book)["rows"] if r.get("product_id") == self.aggregate and r.get("metric") == "SOP计划量")
        self.assertEqual(book["SOP计划"][f"D{row}"].value, 30)
        self.assertTrue(book["SOP计划"][f"C{row}"].protection.locked)
        book.close()
        self.assertEqual(service.import_content(path.read_bytes())["imported"], 0)


if __name__ == "__main__":
    unittest.main()
