"""Real Excel writer -> user edits -> shared store and maintenance workbook."""
import json
import shutil
import tempfile
import unittest
from datetime import datetime
from pathlib import Path

import openpyxl

from sop_export import PlanExport
from sop_roundtrip import PlanRoundtrip
from sop_store import SOPStore, TZ
from sop_workbook import WorkbookSync
from tests.test_workbook_sync import fixture_writer, edit_cell
from tests.test_inline_workbook import change_rows


NODE = Path.home() / ".cache/codex-runtimes/codex-primary-runtime/dependencies/node/bin/node.exe"


@unittest.skipUnless(NODE.is_file(), "Bundled spreadsheet runtime unavailable")
class RoundtripE2ETest(unittest.TestCase):
    def test_real_download_edits_merge_back_without_losing_web_edits(self):
        temporary = tempfile.TemporaryDirectory()
        self.addCleanup(temporary.cleanup)
        root = Path(temporary.name)
        source = root / "sales.json"
        source.write_text(json.dumps({"meta": {}, "products": [
            {"id": "a", "brand": "问界", "model": "M5", "version": "Max"},
            {"id": "b", "brand": "问界", "model": "M5", "version": "Ultra"}], "sales": []}), encoding="utf-8")
        store = SOPStore(root / "test.sqlite3", source, lambda: datetime(2026, 10, 25, 12, tzinfo=TZ))
        self.addCleanup(store.close)
        def save(pid, month, value, forecast=330):
            row = store.db.execute("SELECT revision FROM plans WHERE product_id=? AND month=?", (pid, month)).fetchone()
            store.save_plans({"owner": "验收用户", "changes": [{"product_id": pid, "month": month,
                "sop": value, "forecast": forecast, "revision": row[0] if row else 0}]})
        def value(pid, month):
            row = store.db.execute("SELECT sop,forecast FROM plans WHERE product_id=? AND month=?", (pid, month)).fetchone()
            return tuple(row) if row else (None, None)
        save("a", "2026-01", 100)
        save("a", "2026-10", 10)
        save("a", "2026-12", 20)
        save("b", "2026-10", 5)
        sync = WorkbookSync(store, root / "maintenance.xlsx", writer=fixture_writer)
        self.assertEqual(sync.sync()["status"], "synced")
        importer = PlanRoundtrip(store, sync)
        months = ["2026-10", "2026-12", "2027-12"]
        rows = [["车型汇总", "问界", "M5", "", "锁单数", 70, None, None, ""],
                ["车型汇总", "问界", "M5", "", "SOP计划量", None, None, None, ""]]
        ids = [None, None]
        for metric in ("锁单数", "SOP计划量"):
            for pid, version in (("a", "Max"), ("b", "Ultra")):
                rows.append(["版本明细", "问界", "M5", version, metric, None, None, None, ""])
                ids.append(pid)
        body = {"schema_version": 1, "current_month": "2026-10", "months": months,
                "include_locks": True, "include_versions": True, "scope": "验收",
                "headers": ["层级", "品牌", "车型", "版本", "指标", *months, "数据说明"],
                "rows": rows, "annual_plans": [{"2026": None, "2027": None} for _ in rows],
                "row_product_ids": ids, "selected_product_ids": ["a", "b"]}
        clean, context = importer.prepare(body)
        export = PlanExport()
        original, edited = root / "original.xlsx", root / "edited.xlsx"
        with export.generate(clean, roundtrip=context) as result:
            importer.register(result.path, context)
            shutil.copy2(result.path, original)
        shutil.copy2(original, edited)
        book = openpyxl.load_workbook(original, data_only=False)
        metadata = list(book["_sop_import"].values)
        self.assertEqual(book["_sop_import"].sheet_state, "veryHidden")
        self.assertEqual(metadata[0][1], 2)
        manifest = json.loads("".join(row[1] for row in metadata[5:] if row[0] is not None))
        month_columns = {col["month"]: openpyxl.utils.get_column_letter(col["column"]) for col in manifest["columns"] if col["kind"] == "month"}
        cells = {(row["product_id"], month): f"{column}{row['row']}" for row in manifest["rows"]
                 if row["kind"] == "product" and row["metric"] == "SOP计划量" for month, column in month_columns.items()}
        appended_row = book["SOP计划"].max_row + 1
        book.close()
        edit_cell(edited, cells[("a", "2026-10")], 12)
        edit_cell(edited, cells[("b", "2026-12")], 0)
        change_rows(edited, appended_row, insert={"A": "问界", "B": "新车型", "C": "2027款 新版本", month_columns["2027-12"]: 21})
        # Web edits after export: one untouched cell inside its scope, one outside.
        save("b", "2026-10", 9)
        save("a", "2026-11", 7)
        result = importer.import_content(edited.read_bytes())
        self.assertTrue(result["ok"], result)
        self.assertEqual(result["created"], 1)
        self.assertEqual(value("a", "2026-10"), (12, 330))
        self.assertEqual(value("b", "2026-10"), (9, 330))
        self.assertEqual(value("b", "2026-12"), (0, None))
        self.assertEqual(value("a", "2026-11"), (7, 330))
        self.assertEqual(value("a", "2026-01"), (100, 330))
        # Reimporting the original old copy cannot undo either import or web edits.
        unchanged = importer.import_content(original.read_bytes())
        self.assertTrue(unchanged["ok"], unchanged)
        self.assertEqual(unchanged["imported"], 0)
        self.assertEqual(value("a", "2026-10"), (12, 330))
        save("a", "2026-10", 50)
        edit_cell(edited, cells[("a", "2026-10")], 14)
        conflict = importer.import_content(edited.read_bytes())
        self.assertFalse(conflict["ok"])
        self.assertEqual(conflict["status"], "conflict")
        self.assertEqual(value("a", "2026-10"), (50, 330))
        kept = importer.import_content(edited.read_bytes(), choice="db")
        self.assertTrue(kept["ok"], kept)
        self.assertEqual(value("a", "2026-10"), (50, 330))
        # Deleting an original nonblank plan clears only that month.
        cleared = root / "cleared.xlsx"
        shutil.copy2(original, cleared)
        edit_cell(cleared, cells[("a", "2026-12")], None)
        self.assertTrue(importer.import_content(cleared.read_bytes())["ok"])
        self.assertEqual(value("a", "2026-12"), (None, 330))
        self.assertEqual(value("a", "2026-10"), (50, 330))
        self.assertEqual(sync.status()["status"], "synced")
        self.assertEqual(len(store.list_products()["items"]), 3)



if __name__ == "__main__":
    unittest.main()
