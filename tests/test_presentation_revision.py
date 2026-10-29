"""A workbook created by an older renderer is refreshed after its edits import."""

import json
import os
import tempfile
import unittest
from datetime import datetime
from pathlib import Path
from unittest.mock import patch

from openpyxl import load_workbook

from sop_store import SOPStore, TZ
from sop_workbook import WorkbookSync, _signature
from tests.test_workbook_sync import edit_cell


class PresentationRevisionTest(unittest.TestCase):
    def test_prior_presentation_imports_saved_plan_before_layout_refresh(self):
        with tempfile.TemporaryDirectory() as temporary, patch.dict(os.environ, {"SOP_EXCEL_ENGINE": "python"}):
            root = Path(temporary)
            source = root / "sales.json"
            source.write_text(json.dumps({"meta": {}, "products": [
                {"id": "a", "brand": "问界", "model": "M5", "version": "Max"}],
                "sales": []}, ensure_ascii=False), encoding="utf-8")
            store = SOPStore(root / "test.sqlite3", source,
                             lambda: datetime(2026, 10, 27, 12, tzinfo=TZ))
            try:
                path = root / "SOP计划维护.xlsx"
                with patch("sop_workbook.PRESENTATION_VERSION", 0):
                    previous = WorkbookSync(store, path)
                    self.assertEqual(previous.sync()["status"], "synced")
                    old_token, expected, *_ = previous._parse(path.read_bytes())
                old_without_revision = dict(expected)
                old_without_revision.pop("presentation_version")
                self.assertEqual(_signature(old_without_revision), _signature(expected))
                address = next(item["cell"] for item in expected["layout"]["cells"]
                               if item["product_id"] == "a" and item["month"] == "2026-10")

                # Simulate the previous full-width title while retaining its
                # trusted token, hidden mapping and a user's unsynced plan edit.
                workbook = load_workbook(path)
                try:
                    sheet = workbook["SOP计划"]
                    for merged in list(sheet.merged_cells.ranges):
                        if merged.min_row == merged.max_row == 1:
                            sheet.unmerge_cells(str(merged))
                    sheet["D1"] = None
                    sheet.merge_cells(start_row=1, start_column=1, end_row=1,
                                      end_column=expected["layout"]["sheets"][0]["visible_last_column"])
                    sheet["A1"] = "SOP 跨年计划（2026-01—2027-12）"
                    workbook.save(path)
                finally:
                    workbook.close()
                edit_cell(path, address, 41)

                restarted = WorkbookSync(store, path)
                status = restarted.sync()
                self.assertEqual(status["status"], "synced", status)
                saved = store.db.execute(
                    "SELECT sop FROM plans WHERE product_id='a' AND month='2026-10'").fetchone()
                self.assertEqual(saved[0], 41)
                new_token, *_ = restarted._parse(path.read_bytes())
                self.assertNotEqual(new_token, old_token)
                refreshed = load_workbook(path, read_only=False)
                try:
                    self.assertEqual(refreshed["SOP计划"]["A1"].value, "SOP计划表")
                    self.assertEqual(refreshed["SOP计划"]["D1"].value, "2026-01—2027-12")
                    self.assertIn("A1:C1", {str(area) for area in refreshed["SOP计划"].merged_cells.ranges})
                finally:
                    refreshed.close()
            finally:
                store.close()


if __name__ == "__main__":
    unittest.main()
