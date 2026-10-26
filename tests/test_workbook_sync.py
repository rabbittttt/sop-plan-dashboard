"""Workbook sync tests use synthetic OOXML fixtures; production authoring uses artifact-tool."""
import io
import json
import tempfile
import unittest
import zipfile
import xml.etree.ElementTree as ET
from datetime import datetime
from pathlib import Path
from unittest.mock import patch
from uuid import uuid4

from sop_store import SOPStore, TZ
from sop_workbook import WorkbookSync, _annual_layout, _continuous_layout

NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"


def col_name(index):
    result = ""
    while index:
        index, remainder = divmod(index - 1, 26)
        result = chr(65 + remainder) + result
    return result


def fixture_writer(payload):
    """Minimal test-only OOXML; avoids invoking an external writer in every merge unit test."""
    plans = {(r["product_id"], r["month"]): r.get("sop") for r in payload["plans"]}
    rows = [["product_id", "品牌", "车型", "版本", *payload["months"]]]
    rows += [[p["id"], p["brand"], p["model"], p["version"], *[plans.get((p["id"], m)) for m in payload["months"]]] for p in payload["products"]]
    baselines = [["product_id", "month", "baseline_sop", "revision"]]
    baselines += [[r["product_id"], r["month"], r["baseline_sop"], r["revision"]] for r in payload["baseline"]]
    meta = [["key", "value"], ["token", payload["token"]], ["generated_at", payload["generated_at"]], ["schema_version", 1]]
    sheets = [("SOP计划", rows), ("_sync", baselines), ("_meta", meta)]
    formulas = {}
    merged_ranges = {}
    if payload.get("schema_version") in (2, 3, 4, 5):
        sheets = []
        for spec in payload["layout"]["sheets"]:
            continuous = payload.get("schema_version") in (3, 4, 5)
            inline = payload.get("schema_version") == 5
            width = spec["last_column"] if continuous else 18
            grid = [[None] * width for _ in range(spec["last_row"])]
            grid[0][0] = "SOP 计划"
            grid[1][0] = "灰格不可填"
            if continuous:
                grid[2:4] = [list(r) for r in spec["headers"]]
                if spec.get("manual_section"):
                    section = spec["manual_section"]
                    grid[section["title_row"]-1][0] = section["title"]
                    grid[section["header_row"]-1:section["header_row"]+1] = [list(r) for r in spec["headers"]]
            else:
                grid[2] = ["品牌", "车型", "版本", *[f"{m}月" for m in range(1,13)], "年度合计", "product_id", "row_kind"]
            for item in spec["rows"]:
                row = grid[item["row"]-1]
                row[:3] = [item.get(f"label_{letter}") or None for letter in ("a","b","c")]
                row[(spec["kind_column"] if continuous else 18)-1] = item["kind"]
                if inline:
                    row[spec["key_column"]-1] = item["key"]
                if item["kind"] == "product":
                    row[(spec["product_column"] if continuous else 17)-1] = item["product_id"]
                    if continuous:
                        for col in spec["columns"]:
                            if col["kind"] == "month":
                                value = plans.get((item["product_id"],col["month"]))
                                cell = next(c for c in payload["layout"]["cells"] if c["product_id"]==item["product_id"] and c["month"]==col["month"])
                                row[col["column"]-1] = "未适用" if not inline and value is None and not cell["applicable"] else value
                    else:
                        for month in spec["shown_months"]:
                            row[int(month[-2:])+2] = plans.get((item["product_id"],month))
            formulas[spec["name"]] = spec["formulas"]
            merged_ranges[spec["name"]] = spec["merges"]
            sheets.append((spec["name"],grid))
        meta[3][1] = payload["schema_version"]
        meta.append(["layout_version",payload["layout_version"]])
        mapping = [["product_id","month","sheet","cell","editable"]]
        mapping += [[c["product_id"],c["month"],c["sheet"],c["cell"],int(c["editable"])] for c in payload["layout"]["cells"]]
        if payload["schema_version"]in (3, 4, 5):
            mapping[0].append("applicable")
            for row,cell in zip(mapping[1:],payload["layout"]["cells"]):
                row.append(int(cell["applicable"]))
        sheets += [("_sync",baselines),("_meta",meta),("_layout",mapping)]
    content_types = '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types"><Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/><Default Extension="xml" ContentType="application/xml"/><Override PartName="/xl/workbook.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml"/>'
    content_types += ''.join(f'<Override PartName="/xl/worksheets/sheet{i}.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml"/>' for i in range(1, len(sheets)+1)) + '</Types>'
    relationships = '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
    with zipfile.ZipFile(payload["output_path"], "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("[Content_Types].xml", content_types)
        archive.writestr("_rels/.rels", relationships + '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="xl/workbook.xml"/></Relationships>')
        book = ET.Element(f"{{{NS}}}workbook")
        sheet_elements = ET.SubElement(book,f"{{{NS}}}sheets")
        for i,(name,_) in enumerate(sheets,1):
            ET.SubElement(sheet_elements,f"{{{NS}}}sheet",{"name":name,"sheetId":str(i),"state":"hidden" if name.startswith("_") else "visible","{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id":f"rId{i}"})
        archive.writestr("xl/workbook.xml", ET.tostring(book,encoding="utf-8"))
        archive.writestr("xl/_rels/workbook.xml.rels", relationships + ''.join(f'<Relationship Id="rId{i}" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/worksheet" Target="worksheets/sheet{i}.xml"/>' for i in range(1, len(sheets)+1)) + '</Relationships>')
        for index, (name,sheet_rows) in enumerate(sheets, 1):
            root = ET.Element(f"{{{NS}}}worksheet")
            data = ET.SubElement(root, f"{{{NS}}}sheetData")
            for row_index, row in enumerate(sheet_rows, 1):
                element = ET.SubElement(data, f"{{{NS}}}row", {"r": str(row_index)})
                for column, value in enumerate(row, 1):
                    cell = ET.SubElement(element, f"{{{NS}}}c", {"r": f"{col_name(column)}{row_index}"})
                    address = f"{col_name(column)}{row_index}"
                    if address in formulas.get(name,{}):
                        ET.SubElement(cell,f"{{{NS}}}f").text = formulas[name][address][1:]
                    elif isinstance(value, str):
                        cell.set("t", "inlineStr")
                        ET.SubElement(ET.SubElement(cell, f"{{{NS}}}is"), f"{{{NS}}}t").text = value
                    elif value is not None:
                        ET.SubElement(cell, f"{{{NS}}}v").text = str(value)
            merges = merged_ranges.get(name, ())
            if merges:
                merged = ET.SubElement(root, f"{{{NS}}}mergeCells", {"count": str(len(merges))})
                for reference in merges:
                    ET.SubElement(merged, f"{{{NS}}}mergeCell", {"ref": reference})
            archive.writestr(f"xl/worksheets/sheet{index}.xml", ET.tostring(root, encoding="utf-8"))


def edit_cell(path, reference, value, sheet=1, formula=False):
    """Simulate a user's saved Excel cell edit inside a disposable test fixture."""
    with zipfile.ZipFile(path) as archive:
        contents = {name: archive.read(name) for name in archive.namelist()}
    filename = f"xl/worksheets/sheet{sheet}.xml"
    root = ET.fromstring(contents[filename])
    cell = root.find(f".//{{{NS}}}c[@r='{reference}']")
    if cell is None:
        raise AssertionError("Fixture cell does not exist: " + reference)
    for child in list(cell):
        cell.remove(child)
    cell.attrib.pop("t", None)
    if formula:
        ET.SubElement(cell, f"{{{NS}}}f").text = str(value)
    elif isinstance(value, str):
        cell.set("t", "inlineStr")
        ET.SubElement(ET.SubElement(cell, f"{{{NS}}}is"), f"{{{NS}}}t").text = value
    elif value is not None:
        ET.SubElement(cell, f"{{{NS}}}v").text = str(value)
    contents[filename] = ET.tostring(root, encoding="utf-8")
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as archive:
        for name, content in contents.items():
            archive.writestr(name, content)


class WorkbookSyncTest(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)
        sales = {"meta": {}, "products": [{"id": "a", "brand": "A", "model": "M", "version": "V1"},
                                             {"id": "b", "brand": "A", "model": "M", "version": "V2"}],
                 "sales": [{"product_id": "a", "month": "2026-08", "metric": "locked_orders", "value": 5}], "sources": []}
        (self.root / "sales.json").write_text(json.dumps(sales), encoding="utf-8")
        self.store = SOPStore(self.root / "test.sqlite3", self.root / "sales.json", lambda: datetime(2026, 10, 22, 12, tzinfo=TZ))
        self.path = self.root / "SOP计划维护.xlsx"
        self.sync = WorkbookSync(self.store, self.path, writer=fixture_writer)

    def tearDown(self):
        self.store.close()
        self.temp.cleanup()

    def save(self, pid="a", sop=100, forecast=10, revision=0):
        return self.store.save_plans({"owner": "网页用户", "changes": [{"product_id": pid, "month": "2026-10", "sop": sop, "forecast": forecast, "revision": revision}]})

    def cell(self, pid="a", month="2026-10"):
        context, _ = self.sync._context()
        return next(c["cell"] for c in context["layout"]["cells"] if c["product_id"]==pid and c["month"]==month)

    def current(self, pid="a"):
        return next(r for r in self.store.state()["plans"] if r["product_id"] == pid and r["month"] == "2026-10")

    def test_blank_zero_clear_and_audit(self):
        self.assertEqual(self.sync.sync()["status"], "synced")
        self.assertEqual(self.store.state()["plans"], [])
        edit_cell(self.path, self.cell(), 0)
        self.assertEqual(self.sync.sync()["status"], "synced")
        self.assertEqual(self.current()["sop"], 0)
        self.assertIsNone(self.current()["forecast"])
        self.assertEqual(self.current()["owner"], "Excel计划表")
        self.assertEqual(self.current()["reason"], "表格同步")
        edit_cell(self.path, self.cell(), None)
        self.sync.sync()
        self.assertIsNone(self.current()["sop"])
        self.assertEqual(self.current()["revision"], 2)

    def test_unchanged_poll_keeps_last_sync_time(self):
        initial = self.sync.sync()
        self.assertEqual(initial["status"], "synced")
        self.store.clock = lambda: datetime(2026, 10, 22, 13, tzinfo=TZ)
        unchanged = self.sync.sync()
        self.assertEqual(unchanged["last_sync"], initial["last_sync"])
        self.assertEqual(self.sync.sync()["last_sync"], initial["last_sync"])
        edit_cell(self.path, self.cell(), 25)
        changed = self.sync.sync()
        self.assertNotEqual(changed["last_sync"], initial["last_sync"])
        self.assertEqual(self.current()["sop"], 25)

    def test_field_merge_preserves_latest_forecast_and_unchanged_blanks(self):
        self.save()
        self.sync.sync()
        edit_cell(self.path, self.cell(), 120)
        self.save(sop=100, forecast=200, revision=1)
        self.save(pid="b", sop=75, forecast=30)
        self.assertEqual(self.sync.sync()["status"], "synced")
        self.assertEqual((self.current()["sop"], self.current()["forecast"]), (120, 200))
        self.assertEqual(self.current("b")["sop"], 75)

    def test_conflict_is_atomic_and_choices_are_explicit(self):
        self.save()
        self.sync.sync()
        edit_cell(self.path, self.cell(), 150)
        edit_cell(self.path, self.cell("b"), 25)
        before = self.path.read_bytes()
        self.save(sop=200, forecast=40, revision=1)
        status = self.sync.sync()
        self.assertEqual(status["status"], "conflict")
        self.assertEqual(self.path.read_bytes(), before)
        self.assertEqual(self.current()["sop"], 200)
        self.assertEqual(len(self.store.state()["plans"]), 1)
        self.assertEqual(self.sync.sync(choice="workbook")["status"], "synced")
        self.assertEqual((self.current()["sop"], self.current()["forecast"]), (150, 40))
        self.assertEqual(self.current("b")["sop"], 25)
        edit_cell(self.path, self.cell(), 300)
        self.save(sop=400, forecast=50, revision=3)
        self.assertEqual(self.sync.sync()["status"], "conflict")
        self.assertEqual(self.sync.sync(choice="db")["status"], "synced")
        self.assertEqual(self.current()["sop"], 400)

    def test_invalid_cells_and_baseline_fail_without_partial_import(self):
        self.sync.sync()
        pristine = self.path.read_bytes()
        for bad, formula in [(-1, False), (2.5, False), ("5", False), ("1+2", True)]:
            with self.subTest(bad=bad):
                self.path.write_bytes(pristine)
                edit_cell(self.path, self.cell(), 20)
                edit_cell(self.path, self.cell("b"), bad, formula=formula)
                self.assertEqual(self.sync.sync()["status"], "error")
                self.assertEqual(self.store.state()["plans"], [])
        for bad_id in ("a", "unknown"):
            self.path.write_bytes(pristine)
            context,_=self.sync._context()
            edit_cell(self.path, f"{col_name(context['layout']['sheets'][0]['product_column'])}6", bad_id)
            self.assertEqual(self.sync.sync()["status"], "error")
            self.assertEqual(self.store.state()["plans"], [])
        self.path.write_bytes(pristine)
        edit_cell(self.path, "C2", 999, sheet=2)
        self.assertEqual(self.sync.sync()["status"], "error")
        self.assertEqual(self.store.state()["plans"], [])

    def test_locked_writeback_accepts_sequential_excel_saves_and_recovers(self):
        self.save()
        self.sync.sync()
        edit_cell(self.path, self.cell(), 150)
        with patch("sop_workbook.os.replace", side_effect=PermissionError("Excel open")):
            self.assertEqual(self.sync.sync()["status"], "pending")
            self.assertEqual(self.current()["sop"], 150)
            edit_cell(self.path, self.cell(), 180)
            self.assertEqual(self.sync.sync()["status"], "pending")
            self.assertEqual(self.current()["sop"], 180)
        self.assertEqual(self.sync.sync()["status"], "synced")
        self.assertEqual(self.current()["sop"], 180)

    def test_restart_and_write_during_export_do_not_drop_newer_edits(self):
        self.save()
        self.sync.sync()
        original = self.path.read_bytes()
        restarted = WorkbookSync(self.store, self.path, writer=fixture_writer)
        self.assertEqual(restarted.sync()["status"], "synced")
        self.assertEqual(self.path.read_bytes(), original)  # Unchanged restart does not rewrite.
        self.save(sop=110, revision=1)

        def slow_writer(payload):
            fixture_writer(payload)
            edit_cell(self.path, self.cell(), 150)

        restarted.writer = slow_writer
        self.assertEqual(restarted.sync()["status"], "pending")
        restarted.writer = fixture_writer
        self.assertEqual(restarted.sync()["status"], "conflict")
        self.assertEqual(self.current()["sop"], 110)

    def test_actual_artifact_writer_contract(self):
        import openpyxl
        self.save(sop=0)
        self.save(pid="b",sop=125)
        actual = WorkbookSync(self.store, self.path)
        self.assertEqual(actual.sync()["status"], "synced", actual.status()["message"])
        token, expected, baseline, accepted, values = actual._parse(self.path.read_bytes())
        self.assertTrue(token)
        self.assertIn("2026-08", expected["months"])
        self.assertEqual(expected["months"][-1], "2027-12")
        self.assertEqual(values["a\t2026-10"],0)
        self.assertEqual(values["b\t2026-10"],125)
        self.assertTrue(all(value is None for k,value in values.items() if not k.endswith("2026-10")))
        self.assertEqual(len(values), 2 * len(expected["months"]))
        with zipfile.ZipFile(self.path) as archive:
            root=ET.fromstring(archive.read("xl/worksheets/sheet1.xml"))
            self.assertIsNone(root.find(f"{{{NS}}}sheetProtection"), "Inline rows must support Excel insertion/deletion")
            self.assertEqual(root.find(f".//{{{NS}}}pane").attrib["topLeftCell"],"D5")
            self.assertTrue(all(c.attrib.get("hidden")=="1" for c in root.findall(f"{{{NS}}}cols/{{{NS}}}col") if int(c.attrib['min'])>=30))
        workbook = openpyxl.load_workbook(self.path,read_only=True,data_only=True)
        try:
            self.assertEqual(workbook.sheetnames,["SOP计划","_sync","_meta","_layout"])
            self.assertEqual(workbook["SOP计划"]["P5"].value,0)
            self.assertEqual(workbook["SOP计划"]["P6"].value,125)
            self.assertEqual(workbook["SOP计划"]["M7"].value,125)
            self.assertEqual(workbook["SOP计划"]["P10"].value,125)
            self.assertIn(workbook["SOP计划"]["AC10"].value,(None,""))
            self.assertFalse(workbook["SOP计划"]["M5"].protection.locked)
            self.assertFalse(workbook["SOP计划"]["D5"].protection.locked)
            self.assertTrue(workbook["SOP计划"]["P5"].protection.locked)
            self.assertEqual(workbook["_layout"].sheet_state,"hidden")
        finally:
            workbook.close()
        # A real exported XLSX also round-trips back to SQLite and keeps the
        # independent forecast when the user edits the SOP input.
        edit_cell(self.path,self.cell(),15)
        self.assertEqual(actual.sync()["status"],"synced",actual.status()["message"])
        self.assertEqual((self.current()["sop"],self.current()["forecast"]),(15,10))

    def test_inline_month_headers_and_identity_mapping_are_protected(self):
        self.sync.sync()
        pristine=self.path.read_bytes()
        # Version labels are intentionally editable in schema 5. Column headers
        # and the hidden trusted mapping remain immutable.
        for address,value,sheet,formula in [("A3","Other",1,False),("D3","2030年",1,False),
              ("M4","Other",1,False),("D2","N5",4,False)]:
            with self.subTest(address=address,sheet=sheet):
                self.path.write_bytes(pristine)
                edit_cell(self.path,self.cell(),20)
                edit_cell(self.path,address,value,sheet=sheet,formula=formula)
                status=self.sync.sync()
                self.assertEqual(status["status"],"error",status)
                self.assertEqual(self.store.state()["plans"],[])

    def install_legacy_workbook(self):
        context,_=self.sync._context()
        legacy={k:context[k] for k in ("months","products","baseline","plans")}
        token=str(uuid4())
        self.store.db.execute("INSERT INTO workbook_exports VALUES (?,?,?)",(token,"2026-10-22T12:00:00+08:00",json.dumps(legacy)))
        fixture_writer({**legacy,"token":token,"generated_at":"2026-10-22T12:00:00+08:00","output_path":str(self.path)})
        return f"{col_name(legacy['months'].index('2026-10')+5)}2"

    def test_legacy_unsynced_values_import_before_annual_migration(self):
        self.save(sop=50)
        old_cell=self.install_legacy_workbook()
        edit_cell(self.path,old_cell,0)
        self.assertEqual(self.sync.sync()["status"],"synced")
        self.assertEqual((self.current()["sop"],self.current()["forecast"]),(0,10))
        token,expected,_,_,values=self.sync._parse(self.path.read_bytes())
        self.assertEqual(expected["schema_version"],5)
        self.assertEqual(values["a\t2026-10"],0)
        self.assertTrue(all(v is None for k,v in values.items() if k!="a\t2026-10"))
        original=self.path.read_bytes()
        restarted=WorkbookSync(self.store,self.path,writer=fixture_writer)
        self.assertEqual(restarted.sync()["status"],"synced")
        self.assertEqual(original,self.path.read_bytes())

    def test_legacy_conflict_blocks_upgrade_and_preserves_saved_edits(self):
        self.save(sop=50)
        old_cell=self.install_legacy_workbook()
        edit_cell(self.path,old_cell,75)
        before=self.path.read_bytes()
        self.save(sop=100,forecast=40,revision=1)
        self.assertEqual(self.sync.sync()["status"],"conflict")
        self.assertEqual(self.path.read_bytes(),before)
        self.assertEqual(self.sync.sync(choice="workbook")["status"],"synced")
        self.assertEqual((self.current()["sop"],self.current()["forecast"]),(75,40))
        self.assertEqual(self.sync._parse(self.path.read_bytes())[1]["schema_version"],5)

    def install_annual_workbook(self):
        context,_=self.sync._context()
        months=context["business_months"]
        legacy={**context,"schema_version":2,"layout_version":"annual-blue-v1","months":months,
                "editable_months":months,"baseline":[b for b in context["baseline"] if b["month"] in months]}
        # The previous exported labels, including old model-year prefixes,
        # remain authoritative while importing an old workbook's saved edits.
        legacy["products"]=[{**p,"version":"2026款 · "+p["version"]} for p in context["products"]]
        legacy["layout"]=_annual_layout(legacy["products"],months,months)
        token=str(uuid4())
        self.store.db.execute("INSERT INTO workbook_exports VALUES (?,?,?)",(token,"2026-10-22T12:00:00+08:00",json.dumps(legacy)))
        fixture_writer({**legacy,"token":token,"generated_at":"2026-10-22T12:00:00+08:00","output_path":str(self.path)})

    def test_schema2_saved_edits_are_imported_before_single_sheet_upgrade(self):
        self.save(sop=50)
        self.install_annual_workbook()
        edit_cell(self.path,"M4",0)
        self.assertEqual(self.sync.sync()["status"],"synced",self.sync.status())
        self.assertEqual((self.current()["sop"],self.current()["forecast"]),(0,10))
        _,context,_,_,values=self.sync._parse(self.path.read_bytes())
        self.assertEqual(context["schema_version"],5)
        self.assertEqual([s["name"] for s in context["layout"]["sheets"]],["SOP计划"])
        self.assertEqual(context["months"][0],"2026-01")
        self.assertEqual(context["months"][-1],"2027-12")
        self.assertEqual(len(context["months"]),24)
        self.assertEqual(len(values),48)
        self.assertEqual(values["a\t2026-10"],0)
        self.assertFalse(any("2026款" in r.get("label_c","") for r in context["layout"]["sheets"][0]["rows"] if r.get("label_c")))
        # January and February are real historical plan inputs, even when the
        # first sales record was in August; annual totals must be maintainable.
        edit_cell(self.path,self.cell(month="2026-01"),12)
        self.assertEqual(self.sync.sync()["status"],"synced")
        self.assertEqual(self.store.db.execute("SELECT sop FROM plans WHERE product_id='a' AND month='2026-01'").fetchone()["sop"],12)

    def test_schema2_conflict_prevents_layout_migration(self):
        self.save(sop=50)
        self.install_annual_workbook()
        edit_cell(self.path,"M4",75)
        before=self.path.read_bytes()
        self.save(sop=100,forecast=40,revision=1)
        self.assertEqual(self.sync.sync()["status"],"conflict")
        self.assertEqual(self.path.read_bytes(),before)
        self.assertEqual(self.sync.sync(choice="workbook")["status"],"synced")
        self.assertEqual((self.current()["sop"],self.current()["forecast"]),(75,40))

    def test_catalog_rename_add_missing_source_and_inactive_preserve_plans(self):
        self.save(sop=30)
        self.save(pid="b",sop=40)
        self.sync.sync()
        edit_cell(self.path,self.cell(),35)
        catalog={p["id"]:p for p in self.store.read_sales()["products"]}
        result=self.store.save_products({"changes":[{"id":"a","revision":catalog["a"]["revision"],"version":"2026款 · 后驱 Max"},
                   {"brand":"B","model":"New model","version":"New version"}]})
        new_id=next(p["id"] for p in result["items"] if p["id"]!="a")
        source=json.loads((self.root/"sales.json").read_text(encoding="utf-8"))
        source["products"]=[p for p in source["products"] if p["id"]!="b"]
        (self.root/"sales.json").write_text(json.dumps(source),encoding="utf-8")
        self.assertEqual(self.sync.sync()["status"],"synced",self.sync.status())
        self.assertEqual(self.current()["sop"],35)

        _,context,_,_,values=self.sync._parse(self.path.read_bytes())
        self.assertEqual({p["id"] for p in context["products"]},{"a","b",new_id})
        self.assertEqual(next(p["version"] for p in context["products"] if p["id"]=="a"),"后驱 Max")
        self.assertEqual(values["b\t2026-10"],40)
        self.assertIsNone(values[new_id+"\t2026-10"])
        catalog={p["id"]:p for p in self.store.read_sales()["products"]}
        self.store.save_products({"changes":[{"id":"a","revision":catalog["a"]["revision"],"active":False}]})
        self.assertEqual(self.sync.sync()["status"],"synced",self.sync.status())
        _,context,_,_,values=self.sync._parse(self.path.read_bytes())
        self.assertNotIn("a\t2026-10",values)
        self.assertEqual(self.current()["sop"],35)
        self.assertIn("a\t2026-01",values)
        row=next(r for r in context["layout"]["sheets"][0]["rows"] if r.get("product_id")=="a")
        self.assertEqual(row["label_c"],"后驱 Max")
        self.assertTrue(row["hidden"], "Inactive records remain in history and totals, but are hidden from the editable main table")
        edit_cell(self.path,self.cell(month="2026-01"),7)
        self.assertEqual(self.sync.sync()["status"],"synced",self.sync.status())
        edit_cell(self.path,self.cell(),36)
        self.assertEqual(self.sync.sync()["status"],"error")
        self.assertEqual(self.current()["sop"],35)

    def bind_empty_source(self):
        target=self.store.save_products({"changes":[{"brand":"B","model":"Preplanned","version":"Future version"}]})["items"][0]
        return target

    def commit_binding(self, target):
        self.store.save_products({"changes":[{"id":target["id"],"revision":target["revision"],"bind_source_id":"b"}]})

    def test_bound_empty_source_row_recovers_after_restart_for_each_schema(self):
        # Recreate isolated stores per subtest because source binding is immutable.
        for schema in (1,2,3,4,5):
            with self.subTest(schema=schema):
                if schema != 1:
                    self.tearDown()
                    self.setUp()
                self.save(sop=20)
                target=self.bind_empty_source()
                self.store.save_plans({"owner":"测试","changes":[{"product_id":target["id"],"month":"2026-10","sop":80,"forecast":90,"revision":0}]})
                if schema==1:
                    self.install_legacy_workbook()
                elif schema==2:
                    self.install_annual_workbook()
                elif schema in (3,4):
                    context,_=self.sync._context()
                    context.update(schema_version=schema, layout_version="continuous-blue-v1" if schema==3 else "continuous-blue-v2")
                    plans={(r["product_id"],r["month"]):r for r in context["plans"]}
                    context["layout"]=_continuous_layout(context["products"],context["months"],context["editable_months"],context["current_month"],plans,include_manual=schema==4)
                    token=str(uuid4())
                    self.store.db.execute("INSERT INTO workbook_exports VALUES (?,?,?)",(token,"2026-10-22T12:00:00+08:00",json.dumps(context)))
                    fixture_writer({**context,"token":token,"generated_at":"2026-10-22T12:00:00+08:00","output_path":str(self.path)})
                else:
                    self.sync.sync()
                self.commit_binding(target)
                restarted=WorkbookSync(self.store,self.path,writer=fixture_writer)
                self.assertEqual(restarted.sync()["status"],"synced",restarted.status())
                _,context,_,_,values=restarted._parse(self.path.read_bytes())
                self.assertEqual(context["schema_version"],5)
                self.assertEqual({p["id"] for p in context["products"]},{"a",target["id"]})
                self.assertEqual(values[target["id"]+"\t2026-10"],80)
                self.assertEqual(values["a\t2026-10"],20)
                self.assertFalse(any(k.startswith("b\t") for k in values))

    def test_bound_source_saved_zero_is_not_dropped_when_writeback_was_locked(self):
        target=self.bind_empty_source()
        self.sync.sync()
        old_source_cell=self.cell("b")
        self.commit_binding(target)
        restarted=WorkbookSync(self.store,self.path,writer=fixture_writer)
        with patch("sop_workbook.os.replace",side_effect=PermissionError("Excel open")):
            self.assertEqual(restarted.sync()["status"],"pending",restarted.status())
            # The user saves an old source row while Excel still holds the old
            # workbook; even explicit zero is meaningful and cannot disappear.
            edit_cell(self.path,old_source_cell,0)
            before=self.path.read_bytes()
            self.assertEqual(restarted.sync()["status"],"error",restarted.status())
            self.assertIn("未同步修改",restarted.status()["message"])
            self.assertEqual(self.path.read_bytes(),before)
            self.assertEqual(self.store.state()["plans"],[])
        restarted_again=WorkbookSync(self.store,self.path,writer=fixture_writer)
        self.assertEqual(restarted_again.sync()["status"],"error")
        self.assertEqual(self.path.read_bytes(),before)
        # Once the saved source row is empty again, upgrade can recover.
        edit_cell(self.path,old_source_cell,None)
        self.assertEqual(restarted_again.sync()["status"],"synced",restarted_again.status())

    def test_lifecycle_not_applicable_preserves_existing_zero_and_historical_correction(self):
        target=self.store.save_products({"changes":[{"brand":"B","model":"Next model","version":"New","start_month":"2027-01"}]})["items"][0]
        self.assertEqual(self.sync.sync()["status"],"synced")
        _,context,_,_,values=self.sync._parse(self.path.read_bytes())
        self.assertNotIn(target["id"]+"\t2026-10",values)
        self.assertIn(target["id"]+"\t2027-01",values)
        self.assertTrue(next(c for c in context["layout"]["cells"] if c["product_id"]=="a" and c["month"]=="2026-01")["editable"])
        edit_cell(self.path,self.cell(target["id"]),0)
        self.assertEqual(self.sync.sync()["status"],"error")
        self.assertIn("不可维护",self.sync.status()["message"])
        edit_cell(self.path,self.cell(target["id"]),None)
        self.assertEqual(self.sync.sync()["status"],"synced")
        # A prior explicit zero remains data after a lifespan correction. A
        # historic plan can be corrected; an existing future value stays locked.
        self.store.save_plans({"owner":"测试","changes":[{"product_id":"a","month":"2026-01","sop":0,"forecast":None,"revision":0},
                                                           {"product_id":"a","month":"2027-12","sop":50,"forecast":60,"revision":0}]})
        product=next(p for p in self.store.read_sales()["products"] if p["id"]=="a")
        self.store.save_products({"changes":[{"id":"a","revision":product["revision"],"start_month":"2026-03","end_month":"2027-06"}]})
        self.assertEqual(self.sync.sync()["status"],"synced",self.sync.status())
        _,context,_,_,values=self.sync._parse(self.path.read_bytes())
        self.assertEqual(values["a\t2026-01"],0)
        self.assertNotIn("a\t2026-02",values)
        self.assertNotIn("a\t2027-12",values)
        edit_cell(self.path,self.cell(month="2026-01"),1)
        self.assertEqual(self.sync.sync()["status"],"synced",self.sync.status())
        edit_cell(self.path,self.cell(month="2027-12"),51)
        self.assertEqual(self.sync.sync()["status"],"error")
        self.assertEqual(self.store.db.execute("SELECT sop FROM plans WHERE product_id='a' AND month='2027-12'").fetchone()["sop"],50)

    def test_actual_lifecycle_export_keeps_inapplicable_cells_and_totals_blank(self):
        import openpyxl
        target=self.store.save_products({"changes":[{"brand":"B","model":"Future","version":"Future version","start_month":"2027-01"}]})["items"][0]
        actual=WorkbookSync(self.store,self.path)
        self.assertEqual(actual.sync()["status"],"synced",actual.status())
        _,context,_,_,values=actual._parse(self.path.read_bytes())
        mapping=next(c for c in context["layout"]["cells"] if c["product_id"]==target["id"] and c["month"]=="2026-10")
        self.assertFalse(mapping["editable"])
        row=next(r["row"] for r in context["layout"]["sheets"][0]["rows"] if r.get("product_id")==target["id"])
        workbook=openpyxl.load_workbook(self.path,read_only=True,data_only=True)
        try:
            sheet=workbook["SOP计划"]
            self.assertIsNone(sheet[mapping["cell"]].value)
            self.assertTrue(sheet[mapping["cell"]].protection.locked)
            self.assertIn(sheet[f"P{row}"].value,(None,""))
            self.assertIsNone(sheet[f"Q{row}"].value)
            self.assertFalse(sheet[f"Q{row}"].protection.locked)
        finally:
            workbook.close()

    def test_actual_large_unicode_catalog_keeps_every_coordinate(self):
        source=json.loads((self.root/"sales.json").read_text(encoding="utf-8"))
        source["products"] += [{"id":f"unicode-{i:02d}","brand":"中文品牌","model":f"车型{i//3}",
                                "version":f"2026款 · 四驱高配版本{i:02d}"} for i in range(33)]
        (self.root/"sales.json").write_text(json.dumps(source,ensure_ascii=False),encoding="utf-8")
        actual=WorkbookSync(self.store,self.path)
        context,_=actual._context()
        # This payload exceeds an OS pipe buffer and splits Chinese text across
        # multiple chunks. Every trusted sheet/cell mapping must survive UTF-8.
        self.assertGreater(len(json.dumps(context,ensure_ascii=False).encode("utf-8")),65536)
        self.assertEqual(actual.sync()["status"],"synced",actual.status())
        _,expected,_,_,values=actual._parse(self.path.read_bytes())
        self.assertEqual(len(expected["products"]),35)
        self.assertEqual(len(values),35*24)
        self.assertTrue(all(v is None for v in values.values()))

    def test_horizon_shrink_preserves_values_but_blocks_future_edits(self):
        self.store.clock=lambda:datetime(2027,8,20,12,tzinfo=TZ)
        self.store.save_plans({"owner":"测试","changes":[{"product_id":"a","month":"2028-02","sop":80,"forecast":100,"revision":0}]})
        self.assertEqual(self.sync.sync()["status"],"synced")
        self.store.update_settings({"bridge_m6":False})
        self.assertEqual(self.sync.sync()["status"],"synced")
        _,expected,_,_,values=self.sync._parse(self.path.read_bytes())
        self.assertNotIn("a\t2028-02",values)
        self.assertEqual(expected["plan_end"],"2027-12")
        mapped=next(c for c in expected["layout"]["cells"] if c["product_id"]=="a" and c["month"]=="2028-02")
        self.assertFalse(mapped["editable"])
        sheet=next(i for i,s in enumerate(expected["layout"]["sheets"],1) if s["name"]==mapped["sheet"])
        edit_cell(self.path,mapped["cell"],81,sheet=sheet)
        self.assertEqual(self.sync.sync()["status"],"error")
        self.assertEqual(self.store.db.execute("SELECT sop FROM plans WHERE month='2028-02'").fetchone()["sop"],80)

    def test_newly_closed_month_edit_is_not_lost_during_upgrade(self):
        self.store.clock=lambda:datetime(2027,8,20,12,tzinfo=TZ)
        self.sync.sync()
        context,_=self.sync._context()
        mapped=next(c for c in context["layout"]["cells"] if c["product_id"]=="a" and c["month"]=="2028-02")
        edit_cell(self.path,mapped["cell"],81,sheet=1)
        before=self.path.read_bytes()
        self.store.update_settings({"bridge_m6":False})
        self.assertEqual(self.sync.sync()["status"],"error")
        self.assertEqual(self.path.read_bytes(),before)
        self.assertEqual(self.store.state()["plans"],[])


if __name__ == "__main__":
    unittest.main()
