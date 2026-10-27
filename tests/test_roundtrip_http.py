"""Upload/download transport checks; business merge behavior is tested separately."""
import base64
import http.client
import io
import json
import tempfile
import threading
import unittest
from datetime import datetime
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path

import openpyxl

from server import make_handler
from sop_export import PlanExport
from sop_roundtrip import PlanRoundtrip
from sop_store import ConflictError, SOPStore, TZ
from sop_workbook import WorkbookSync
from tests.test_plan_export import NoBusinessAccess, fixture_writer, payload
from tests.test_roundtrip import export_body
from tests.test_workbook_sync import edit_cell


class RoundtripStub:
    def __init__(self):
        self.uploads = []
        self.registered = []
        self.error = None
        self.result = {"ok": True, "imported": 2, "created": 0, "status": "imported"}

    def prepare(self, body):
        clean = dict(body)
        ids = clean.pop("row_product_ids")
        return clean, {"token": "trusted-token", "row_product_ids": ids}

    def register(self, path, context):
        self.registered.append((path.exists(), context))

    def import_content(self, content, choice=None):
        self.uploads.append((content, choice))
        if self.error:
            raise self.error
        return self.result


class RoundtripHTTPTest(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.root = Path(self.tmp.name)
        self.service = RoundtripStub()
        self.http = ThreadingHTTPServer(("127.0.0.1", 0), BaseHTTPRequestHandler)
        self.http.RequestHandlerClass = make_handler(
            NoBusinessAccess(), self.root, self.http.server_port,
            workbook_sync=NoBusinessAccess(),
            plan_export=PlanExport(writer=fixture_writer, temp_root=self.root),
            plan_roundtrip=self.service)
        self.thread = threading.Thread(target=lambda: self.http.serve_forever(poll_interval=.02), daemon=True)
        self.thread.start()

    def tearDown(self):
        self.http.shutdown()
        self.thread.join()
        self.http.server_close()
        self.tmp.cleanup()

    def request(self, body, path="/api/plans/import-xlsx", headers=None):
        connection = http.client.HTTPConnection("127.0.0.1", self.http.server_port, timeout=5)
        connection.request("POST", path, json.dumps(body), headers or {"Content-Type": "application/json"})
        response = connection.getresponse()
        result = response.status, response.read()
        connection.close()
        return result

    def upload(self, **extra):
        return {"filename": "SOP计划.xlsx", "content_base64": base64.b64encode(b"xlsx-content").decode(), **extra}

    def test_download_registered_before_delivery_and_cleaned(self):
        body = {**payload(), "row_product_ids": [None, None, "p1"]}
        code, data = self.request(body, "/api/plans/export-xlsx")
        self.assertEqual(code, 200)
        self.assertTrue(data.startswith(b"PK"))
        self.assertEqual(self.service.registered, [(True, {"token": "trusted-token", "row_product_ids": [None, None, "p1"]})])
        self.assertEqual(list(self.root.iterdir()), [])

    def test_upload_and_conflict_resolution_transport(self):
        code, data = self.request(self.upload())
        self.assertEqual(code, 200)
        self.assertEqual(json.loads(data)["imported"], 2)
        self.assertEqual(self.service.uploads, [(b"xlsx-content", None)])
        conflict = ConflictError([{"month": "2026-10", "baseline": 1, "database": 2, "workbook": 3}])
        self.service.error = conflict
        code, data = self.request(self.upload())
        self.assertEqual(code, 409)
        self.assertEqual(json.loads(data)["conflicts"][0]["database"], 2)
        conflict.source = "maintenance"
        code, data = self.request(self.upload(choice="workbook"))
        self.assertEqual(json.loads(data)["source"], "maintenance")
        self.assertEqual(self.service.uploads[-1], (b"xlsx-content", "workbook"))

    def test_invalid_uploads_and_cross_origin_never_reach_service(self):
        for body in [{}, self.upload(filename="bad.xlsm"), self.upload(content_base64="%%%"),
                     self.upload(content_base64=""), self.upload(choice="anything"),
                     self.upload(path="C:/unrelated.xlsx"), self.upload(filename=42)]:
            with self.subTest(body=body):
                code, data = self.request(body)
                self.assertEqual(code, 400)
                self.assertIn("error", json.loads(data))
        code, _ = self.request(self.upload(), headers={"Content-Type": "application/json", "Origin": "https://external.example"})
        self.assertEqual(code, 403)
        self.assertFalse(self.service.uploads)

    def test_returned_conflicts_and_committed_warning_keep_their_status(self):
        self.service.result = {"ok": False, "status": "conflict", "imported": 0,
                               "conflicts": [{"month": "2026-10", "database": 6, "workbook": 8}]}
        code, data = self.request(self.upload())
        self.assertEqual(code, 409)
        self.assertEqual(json.loads(data), self.service.result)
        self.service.result = {"ok": True, "status": "imported", "imported": 2, "created": 0,
                               "warning": "计划已导入；维护表待同步", "workbook_sync": {"status": "pending"}}
        code, data = self.request(self.upload(choice="workbook"))
        self.assertEqual(code, 200)
        self.assertEqual(json.loads(data), self.service.result)


@unittest.skipUnless((Path.home() / ".cache/codex-runtimes/codex-primary-runtime/dependencies/node/bin/node.exe").is_file(),
                     "Bundled spreadsheet runtime unavailable")
class MissingPrimaryWorkbookHTTPTest(unittest.TestCase):
    def test_download_and_reimport_when_maintenance_workbook_cannot_be_created(self):
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            source = root / "sales.json"
            source.write_text(json.dumps({"meta": {}, "products": [
                {"id": "a", "brand": "A", "model": "M", "version": "同名版本", "generation": "2026"},
                {"id": "b", "brand": "A", "model": "M", "version": "同名版本", "generation": "2027"}], "sales": []}), encoding="utf-8")
            store = SOPStore(root / "test.sqlite3", source, lambda: datetime(2026, 10, 25, 12, tzinfo=TZ))
            workbook_path = root / "SOP计划维护.xlsx"

            def unavailable(_payload):
                raise RuntimeError("maintenance writer unavailable")

            sync = WorkbookSync(store, workbook_path, writer=unavailable)
            service = PlanRoundtrip(store, sync)
            server = ThreadingHTTPServer(("127.0.0.1", 0), BaseHTTPRequestHandler)
            server.RequestHandlerClass = make_handler(store, root, server.server_port,
                workbook_sync=sync, plan_export=PlanExport(temp_root=root),
                source_refresh=NoBusinessAccess(), plan_roundtrip=service)
            thread = threading.Thread(target=lambda: server.serve_forever(poll_interval=.02), daemon=True)
            thread.start()
            try:
                store.save_plans({"owner": "网页", "changes": [{"product_id": "a", "month": "2026-10",
                    "forecast": None, "sop": 10, "revision": 0}]})

                def request(path, body):
                    connection = http.client.HTTPConnection("127.0.0.1", server.server_port, timeout=180)
                    try:
                        connection.request("POST", path, json.dumps(body), {"Content-Type": "application/json"})
                        response = connection.getresponse()
                        return response.status, response.read()
                    finally:
                        connection.close()

                status, content = request("/api/plans/export-xlsx", export_body())
                self.assertEqual(status, 200, content[:500])
                self.assertFalse(workbook_path.exists())
                exported = root / "download.xlsx"
                exported.write_bytes(content)
                book = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
                try:
                    self.assertIn("_sop_import", book.sheetnames)
                    row = next(index for index, values in enumerate(book["SOP计划"].iter_rows(values_only=True), 1)
                               if "plan:a" in values)
                    self.assertEqual(book["SOP计划"][f"D{row}"].value, 10)
                finally:
                    book.close()

                edit_cell(exported, f"D{row}", 41)
                status, content = request("/api/plans/import-xlsx", {
                    "filename": "download.xlsx", "content_base64": base64.b64encode(exported.read_bytes()).decode()})
                self.assertEqual(status, 200, content[:500])
                self.assertEqual(json.loads(content)["imported"], 1)
                self.assertEqual(store.db.execute("SELECT sop FROM plans WHERE product_id='a' AND month='2026-10'").fetchone()[0], 41)
            finally:
                server.shutdown()
                thread.join()
                server.server_close()
                store.close()


if __name__ == "__main__":
    unittest.main()
