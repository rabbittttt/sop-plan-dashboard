"""Offline XLSX writer for the managed SOP maintenance workbook.

The service supplies a trusted schema-5 layout and an immutable export token.
Only the writer changes here; the same parser and SQLite ledger remain authoritative.
"""
from __future__ import annotations

import re
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Protection
from openpyxl.utils import get_column_letter, range_boundaries
from openpyxl.worksheet.datavalidation import DataValidation

from sop_xlsx_style import setup_plan_sheet, style_plan_row, write_formula_caches


_SUM_FORMULA = re.compile(r'^=IF\(COUNT\((.+)\)=0,"",SUM\(\1\)\)$')


def _quantity(value):
    if value is not None and (type(value) is not int or not 0 <= value <= 10**12):
        raise ValueError("SOP 计划量须为非负整数或空白。")
    return value


def _literal(cell, value):
    """Preserve labels and identifiers as Excel text, including leading '='."""
    cell.value = value
    if isinstance(value, str):
        cell.data_type = "s"


def _formula_results(sheet, formulas):
    cache, visiting = {}, set()

    def result(address):
        if address in cache:
            return cache[address]
        if address in visiting:
            raise ValueError("计划表合计公式不能循环引用。")
        cell = sheet[address]
        if address not in formulas:
            value = cell.value
            return value if type(value) is int else None
        visiting.add(address)
        formula = formulas[address]
        if formula == '=""':
            value = None
        else:
            matched = _SUM_FORMULA.fullmatch(formula)
            if not matched:
                raise ValueError(f"不支持的计划表合计公式：{address}")
            values = []
            for reference in matched.group(1).split(","):
                min_col, min_row, max_col, max_row = range_boundaries(reference)
                for row in range(min_row, max_row + 1):
                    for column in range(min_col, max_col + 1):
                        child = f"{get_column_letter(column)}{row}"
                        amount = result(child)
                        if amount is not None:
                            values.append(amount)
            value = sum(values) if values else None
        visiting.remove(address)
        cache[address] = value
        return value

    for address in formulas:
        result(address)
    return cache


def write_plan_workbook(payload: dict) -> None:
    """Write ``payload['output_path']`` using only the declared Python dependency."""
    if (payload.get("schema_version") != 5
            or payload.get("layout_version") not in {"continuous-inline-v1", "continuous-inline-v2"}
            or not payload.get("token") or not payload.get("output_path")
            or not payload.get("layout", {}).get("sheets")):
        raise ValueError("缺少可信的 SOP 维护表布局或同步标识。")

    plans = {(row["product_id"], row["month"]): _quantity(row.get("sop"))
             for row in payload["plans"]}
    mappings_by_sheet = {}
    for mapping in payload["layout"]["cells"]:
        mappings_by_sheet.setdefault(mapping["sheet"], {})[mapping["cell"]] = mapping

    workbook = Workbook()
    workbook.remove(workbook.active)
    caches = {}
    for spec in payload["layout"]["sheets"]:
        sheet = workbook.create_sheet(spec["name"])
        columns = spec["columns"]
        month_columns = [item for item in columns if item["kind"] == "month"]
        setup_plan_sheet(sheet, last_row=spec["last_row"], last_column=spec["last_column"],
                         visible_last_column=spec["visible_last_column"], columns=columns,
                         header_row=spec["header_row"], first_data_row=spec["first_data_row"])
        _literal(sheet.cell(1, 1), f"SOP 跨年计划（{payload['months'][0]}—{payload['months'][-1]}）")
        _literal(sheet.cell(2, 1), "单位：辆。直接插入、删除明细行或修改名称；删除表示停用，历史保留。保存后重算合计。")
        for offset, header in enumerate(spec["headers"]):
            for column, value in enumerate(header, 1):
                if value is not None:
                    _literal(sheet.cell(spec["header_row"] + offset, column), value)

        mappings = mappings_by_sheet.get(spec["name"], {})
        validation = DataValidation(type="whole", operator="between", formula1=0,
                                    formula2=10**12, allow_blank=True)
        validation.showInputMessage = True
        validation.showErrorMessage = True
        validation.errorStyle = "stop"
        validation.promptTitle = "SOP计划量（辆）"
        validation.prompt = "只填非负整数。空白表示未填写，0表示明确不安排。"
        validation.errorTitle = "请输入非负整数"
        validation.error = "SOP计划量须为非负整数；未填写时请留空。"
        sheet.add_data_validation(validation)

        for item in spec["rows"]:
            row = item["row"]
            for column, label in enumerate("abc", 1):
                text = item.get("label_" + label)
                if text is not None:
                    _literal(sheet.cell(row, column), text)
            _literal(sheet.cell(row, spec["kind_column"]), item["kind"])
            _literal(sheet.cell(row, spec["key_column"]), item["key"])
            states = {}
            if item["kind"] == "product":
                _literal(sheet.cell(row, spec["product_column"]), item["product_id"])
                for column in range(1, 4):
                    sheet.cell(row, column).protection = Protection(locked=False)
                for month in month_columns:
                    address = f"{get_column_letter(month['column'])}{row}"
                    mapping = mappings.get(address)
                    if mapping is None or mapping["product_id"] != item["product_id"]:
                        raise ValueError(f"缺少可信的计划输入映射：{spec['name']}!{address}")
                    value = plans.get((item["product_id"], month["month"]))
                    if value is not None:
                        sheet[address] = value
                    states[address] = {"editable": mapping["editable"],
                                       "applicable": mapping["applicable"]}
                    if mapping["editable"]:
                        sheet[address].protection = Protection(locked=False)
                        validation.add(sheet[address])
            style_plan_row(sheet, row=row, kind=item["kind"],
                           visible_last_column=spec["visible_last_column"], columns=columns,
                           version=item.get("version", ""), month_states=states)
            if item.get("hidden"):
                sheet.row_dimensions[row].hidden = True

        for address, formula in spec["formulas"].items():
            sheet[address] = formula
        for reference in spec["merges"]:
            sheet.merge_cells(reference)
        sheet.freeze_panes = f"D{spec['first_data_row']}"
        caches[sheet.title] = _formula_results(sheet, spec["formulas"])

    sync = workbook.create_sheet("_sync")
    sync.append(["product_id", "month", "baseline_sop", "revision"])
    for item in payload["baseline"]:
        sync.append([item["product_id"], item["month"], _quantity(item["baseline_sop"]),
                     item["revision"]])
    for row in sync:
        for cell in row[:2]:
            cell.number_format = "@"
            if isinstance(cell.value, str):
                cell.data_type = "s"
    sync.sheet_state = "hidden"

    meta = workbook.create_sheet("_meta")
    for row in (["key", "value"], ["token", payload["token"]],
                ["generated_at", payload["generated_at"]],
                ["schema_version", payload["schema_version"]],
                ["layout_version", payload["layout_version"]]):
        meta.append(row)
    meta.sheet_state = "hidden"

    layout = workbook.create_sheet("_layout")
    layout.append(["product_id", "month", "sheet", "cell", "editable", "applicable"])
    for item in payload["layout"]["cells"]:
        layout.append([item["product_id"], item["month"], item["sheet"], item["cell"],
                       int(item["editable"]), int(item["applicable"])])
    for row in layout:
        for cell in row[:4]:
            cell.number_format = "@"
            if isinstance(cell.value, str):
                cell.data_type = "s"
    layout.sheet_state = "hidden"

    output = Path(payload["output_path"])
    output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output)
    write_formula_caches(output, workbook.sheetnames, caches)
