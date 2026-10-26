"""Read-only extraction of configuration-version monthly source references.

Only the version/quantity blocks in monthly option workbooks feed plan products.
Other option dimensions and aggregate rows are never appended. Consolidated
daily generation totals are isolated references, never assigned to a version.
Unknown data stays unknown, and source metric ambiguities remain visible.
"""
from __future__ import annotations

import argparse
import calendar
from collections import Counter, defaultdict
from datetime import date, datetime, timedelta, timezone
import hashlib
import json
import math
from pathlib import Path
import re
import sys

import openpyxl

from sop_paths import sales_source_files


ROOT = Path(__file__).resolve().parent
DEFAULT_SOURCE_DIR = ROOT / "input_file"
SHANGHAI = timezone(timedelta(hours=8))
SOURCE_NAME = "鸿蒙智行销量数据汇总.xlsx"
DAILY_SHEET = "小订首销平销by天"
METRICS = [
    {"id": "locked_orders", "label": "交车锁单", "source_column": "交车锁单", "unit": "单"},
    {"id": "gross_orders", "label": "大定", "source_column": "大定", "unit": "单"},
    {"id": "net_orders", "label": "留存大定", "source_column": "留存大定", "unit": "单"},
]


def _text(value):
    return str(value).strip() if value is not None else ""


def _number(value):
    if isinstance(value, bool) or not isinstance(value, (int, float)):
        return None
    if not math.isfinite(value):
        return None
    return int(value) if float(value).is_integer() else float(value)


def _date(value):
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if value:
        try:
            return date.fromisoformat(str(value)[:10])
        except ValueError:
            pass
    return None


def _table(workbook, sheet_name, required):
    if sheet_name not in workbook.sheetnames:
        raise ValueError(f"缺少工作表：{sheet_name}")
    sheet = workbook[sheet_name]
    iterator = sheet.iter_rows(values_only=True)
    header = [_text(v) for v in next(iterator)]
    missing = set(required) - set(header)
    if missing:
        raise ValueError(f"{sheet_name} 缺少字段：{'、'.join(sorted(missing))}")
    for row_number, values in enumerate(iterator, 2):
        if any(v is not None for v in values):
            yield row_number, dict(zip(header, values))


def _product(generation, secondary, brand=""):
    generation, secondary = _text(generation), _text(secondary)
    key = generation + "\x1f" + secondary
    brand = _text(brand) or generation.split(" ", 1)[0]
    model_and_year = generation[len(brand):].strip()
    match = re.fullmatch(r"(.+?)\s+(20\d{2}款)(.*)", model_and_year)
    if match:
        model, version = match.group(1).strip(), (match.group(2) + match.group(3)).strip()
    else:
        model, version = model_and_year or generation, "未标年款"
    if secondary:
        secondary_suffix = secondary[len(generation):].strip() if secondary.startswith(generation) else secondary
        version += " · " + secondary_suffix
    return {
        "id": "p_" + hashlib.sha256(key.encode("utf-8")).hexdigest()[:16],
        "brand": brand, "model": model, "version": version,
        "generation": generation, "secondary_generation": secondary or None,
        "name": secondary or generation,
    }


def _source_file(path):
    stat = path.stat()
    return {
        "name": path.name,
        "modified_at": datetime.fromtimestamp(stat.st_mtime, SHANGHAI).isoformat(timespec="seconds"),
        "size_bytes": stat.st_size,
        "sha256": hashlib.sha256(path.read_bytes()).hexdigest(),
    }


def _option_consistency(source_dir):
    """Review monthly controls without importing conflicting option dimensions."""
    path = source_dir / "鸿蒙智行大定选配比例分析.xlsx"
    if not path.exists():
        return []
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    controls = {}
    try:
        for sheet in workbook.worksheets:
            if not sheet.title.endswith("by月"):
                continue
            rows = list(sheet.iter_rows(max_row=3, values_only=True))
            if len(rows) < 3 or not rows[0] or rows[0][0] != "指标":
                continue
            header = rows[0]
            for values in rows[1:]:
                if values[0] == "净大定":
                    controls[sheet.title[:-3]] = {
                        _text(header[c]): _number(values[c])
                        for c in range(3, min(len(values), len(header)))
                        if re.fullmatch(r"\d{2}-\d{2}", _text(header[c]))
                    }
        diffs = []
        for name, control in controls.items():
            if " " not in name:
                children = [data for child, data in controls.items() if child.startswith(name + " ")]
                if not children:
                    continue
                for month, value in control.items():
                    values = [data.get(month) for data in children]
                    if value is not None and all(v is not None for v in values):
                        detail = sum(values)
                        if value != detail:
                            diffs.append({"brand": name, "month": "20" + month, "brand_value": value, "model_sum": detail})
        return diffs
    finally:
        workbook.close()


def _extract_generation_reference(source_dir=None, *, as_of=None):
    """Extract generation-level references; never modify any file.

    `as_of` may be a date or ISO date. Records for today/future and records not
    explicitly marked as finished days are excluded. The return is the JSON dict.
    Optional paths are resolved against the caller's current directory; defaults
    are anchored beside this module so refresh() also works from a launcher.
    """
    source_dir = Path(source_dir) if source_dir is not None else DEFAULT_SOURCE_DIR
    path = source_dir / SOURCE_NAME
    if not path.is_file():
        raise FileNotFoundError(f"未找到销量汇总文件：{path}")
    now = datetime.now(SHANGHAI)
    observation_date = _date(as_of) if as_of else now.date()
    if observation_date is None:
        raise ValueError("as_of 必须是 YYYY-MM-DD 日期")
    source_files = [_source_file(p) for p in sales_source_files(source_dir)]
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    products, info, days = {}, {}, {}
    notices, exclusions, duplicates = [], Counter(), 0
    source_notes, source_declared_date = [], None
    try:
        for row_number, row in _table(workbook, "说明与来源", ["类型", "项目", "内容"]):
            source_notes.append({"row": row_number, "type": row["类型"], "item": row["项目"], "content": row["内容"]})
            if row.get("项目") == "数据判定日期":
                source_declared_date = _date(row.get("内容"))
        for row_number, row in _table(workbook, "车型基本信息", ["代际名", "二级代际名", "品牌"]):
            product = _product(row["代际名"], row["二级代际名"], row["品牌"])
            product.update({
                "launch_date": _date(row.get("首销开始")).isoformat() if _date(row.get("首销开始")) else None,
                "first_sale_end": _date(row.get("首销结束")).isoformat() if _date(row.get("首销结束")) else None,
                "energy": row.get("能源类型"), "segment": row.get("产品档位"),
            })
            if product["id"] in products:
                raise ValueError(f"车型基本信息存在重复产品：{product['name']}")
            products[product["id"]] = product
            info[product["id"]] = (row_number, row)
        required = ["代际名", "二级代际名", "日期", "日期状态"] + [m["source_column"] for m in METRICS]
        for row_number, row in _table(workbook, DAILY_SHEET, required):
            dt = _date(row["日期"])
            if dt is None:
                exclusions["日期无效"] += 1
                continue
            if dt >= observation_date:
                exclusions["当日或未来日期"] += 1
                continue
            if _text(row["日期状态"]) != "已结束日":
                exclusions["不是已结束日"] += 1
                continue
            product = _product(row["代际名"], row["二级代际名"])
            pid = product["id"]
            products.setdefault(pid, product)
            key = (pid, dt)
            values = {m["id"]: _number(row[m["source_column"]]) for m in METRICS}
            entry = {"values": values, "row": row_number, "stage": row.get("订单阶段"), "source": row.get("字段来源")}
            if key in days:
                if days[key]["values"] != values:
                    raise ValueError(f"同产品同日期存在冲突记录：{product['name']} {dt}")
                duplicates += 1
                continue
            days[key] = entry
            if row.get("首销分时核对") and "不一致" in str(row["首销分时核对"]):
                notices.append({"product_id": pid, "date": dt.isoformat(), "message": row["首销分时核对"], "row": row_number})
    finally:
        workbook.close()

    buckets = defaultdict(list)
    for (pid, dt), entry in days.items():
        for metric, value in entry["values"].items():
            if value is not None:
                buckets[(pid, dt.strftime("%Y-%m"), metric)].append((dt, value, entry["row"]))
    sales = []
    for (pid, month, metric), values in sorted(buckets.items()):
        year, mon = map(int, month.split("-"))
        expected_days = calendar.monthrange(year, mon)[1]
        total = sum(v for _, v, _ in values)
        sales.append({
            "product_id": pid, "month": month, "metric": metric, "value": total,
            "known_days": len(values), "expected_days": expected_days,
            "complete": len(values) == expected_days,
            "from_date": min(dt for dt, _, _ in values).isoformat(),
            "as_of": max(dt for dt, _, _ in values).isoformat(),
            "source_id": "consolidated_daily",
        })
    metric_coverage = {}
    for metric in METRICS:
        mid = metric["id"]
        relevant = [row for row in sales if row["metric"] == mid]
        raw_values = [entry["values"][mid] for entry in days.values() if entry["values"][mid] is not None]
        extracted_total = sum(row["value"] for row in relevant)
        if extracted_total != sum(raw_values):
            raise AssertionError(f"{mid} 月汇总无法对回逐日原始量")
        metric_coverage[mid] = {
            "data_from": min((r["from_date"] for r in relevant), default=None),
            "data_through": max((r["as_of"] for r in relevant), default=None),
            "source_daily_total": sum(raw_values), "monthly_total": extracted_total,
            "known_daily_values": len(raw_values), "monthly_records": len(relevant),
            "complete_month_records": sum(r["complete"] for r in relevant),
            "product_count": len({r["product_id"] for r in relevant}), "reconciled": True,
        }

    reconciliation = []
    for pid, (row_number, row) in info.items():
        start, end = _date(row.get("首销开始")), _date(row.get("首销结束"))
        if not start or not end or end >= observation_date:
            continue
        for metric, source_total in (("gross_orders", "总大定"), ("net_orders", "首销期留存大定"), ("locked_orders", "首销期锁单")):
            target = _number(row.get(source_total))
            if target is None:
                continue
            period = [entry["values"][metric] for (p, dt), entry in days.items() if p == pid and start <= dt <= end and entry["values"][metric] is not None]
            expected = (end - start).days + 1
            complete = len(period) == expected
            reconciliation.append({
                "product_id": pid, "metric": metric, "start": start.isoformat(), "end": end.isoformat(),
                "known_days": len(period), "expected_days": expected,
                "daily_total": sum(period) if period else None, "source_total": target,
                "status": "matched" if complete and sum(period) == target else "mismatch" if complete else "partial",
                "source_sheet": "车型基本信息", "source_row": row_number,
            })
    option_conflicts = _option_consistency(source_dir)
    partial_count = sum(not r["complete"] for r in sales)
    quality = [
        {"level": "warning", "code": "partial_months", "message": f"{partial_count} 条产品月指标只有部分已知日量，显示已知合计；未覆盖日期不补0。", "count": partial_count},
        {"level": "info", "code": "daily_month_reconciled", "message": "全部月合计已逐项对回汇总文件的已知日量。"},
    ]
    if option_conflicts:
        quality.append({"level": "warning", "code": "option_control_conflicts", "message": f"选配月表有 {len(option_conflicts)} 项品牌小计与车型合计不一致，未混入销量汇总。", "count": len(option_conflicts), "examples": option_conflicts[:5]})
    if notices:
        quality.append({"level": "warning", "code": "source_hour_day_conflicts", "message": f"源汇总标记 {len(notices)} 条日量与分时不一致，保留源表优先级和日量口径。", "count": len(notices), "examples": notices[:5]})
    mismatches = [r for r in reconciliation if r["status"] == "mismatch"]
    if mismatches:
        quality.append({"level": "warning", "code": "first_sale_control_conflicts", "message": f"{len(mismatches)} 项首销日量合计与车型基本信息终值不一致；未用终值覆盖日量。", "count": len(mismatches), "examples": mismatches[:5]})
    if any(exclusions.values()):
        quality.append({"level": "info", "code": "excluded_rows", "message": "已排除当日、未来或未结束日。", "counts": dict(exclusions)})
    data_through = max((r["as_of"] for r in sales), default=None)
    result = {
        "meta": {
            "generated_at": now.isoformat(timespec="seconds"), "timezone": "Asia/Shanghai",
            "data_through": data_through,
            "source_declared_as_of": source_declared_date.isoformat() if source_declared_date else None,
            "observation_date": observation_date.isoformat(), "source_files": source_files,
            "default_metric": "locked_orders", "metrics": METRICS, "unit": "单",
            "metric_coverage": metric_coverage,
            "notes": [
                "数据仅来自销量汇总的小订首销平销by天，按品牌、车型、年款及二级代际聚合为月。",
                "交车锁单、留存大定和大定是不同订单指标；未将任何订单指标标为已交付。",
                "月值为已知日量合计；完整月要求日历月每一天都有数值，空白和无记录均不等于0。",
                "同名车型的不同年款、二级代际分别保留；选配、地区、门店等维度不重复相加。",
                "源文件修改时间不代表数据截止时间；各指标最后已知日可能不同。",
                "每周日24点为周一00:00（北京时间）；历史封版仅可由保存的导入版本和快照产生。",
            ],
            "quality": quality, "duplicate_identical_days_ignored": duplicates,
            "reconciliation": reconciliation,
        },
        "products": sorted(products.values(), key=lambda p: (p["brand"], p["model"], p["version"])),
        "sales": sales,
        "sources": [{
            "id": "consolidated_daily", "file": SOURCE_NAME, "sheet": DAILY_SHEET,
            "range": "A1:O" + str(max((r["row"] for r in days.values()), default=1)),
            "description": "按已结束日逐字段取源汇总数值，按自然月聚合；未补未知日。",
        }, {
            "id": "product_registry", "file": SOURCE_NAME, "sheet": "车型基本信息",
            "description": "品牌、车型、年款/二级代际与首销总量核对。",
        }],
    }
    return result


VERSION_METRICS = [
    {"id": "net_orders", "label": "净大定（版本参考）", "unit": "单",
     "source_file": "鸿蒙智行大定选配比例分析.xlsx", "control_label": "净大定"},
    {"id": "delivery_reference", "label": "锁单来源版本数量", "unit": "单",
     "source_file": "鸿蒙智行锁单选配比例分析.xlsx", "control_label": "交车锁单"},
]


def _month(value):
    text = _text(value)
    match = re.fullmatch(r"(\d{2}|\d{4})[-/](\d{1,2})", text)
    if match:
        year, month = int(match[1]), int(match[2])
        year = year + 2000 if year < 100 else year
        if 1 <= month <= 12:
            return f"{year:04d}-{month:02d}"
    if isinstance(value, (date, datetime)):
        return value.strftime("%Y-%m")
    return None


def _extract_version_reference(source_dir=None, *, as_of=None):
    """Read true configuration-version monthly quantities, plus isolated references.

    Missing cells stay absent. A source month's end day is not fabricated from
    the workbook mtime, another tab, or the consolidated daily reference. IDs
    encode generation and configuration, so year-model changes remain separate.
    """
    source_dir = Path(source_dir) if source_dir is not None else DEFAULT_SOURCE_DIR
    now = datetime.now(SHANGHAI)
    observation = _date(as_of) if as_of else now.date()
    if observation is None:
        raise ValueError("as_of 必须是 YYYY-MM-DD 日期")
    products, sales, sources, checks, missing, duplicate_keys = {}, [], [], [], [], set()
    for metric in VERSION_METRICS:
        path = source_dir / metric["source_file"]
        if not path.exists():
            raise FileNotFoundError(f"未找到版本数量来源：{path}")
        workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
        try:
            for sheet in workbook.worksheets:
                if not sheet.title.endswith("by月") or " " not in sheet.title:
                    continue
                rows = list(sheet.iter_rows(values_only=True))
                if not rows or tuple(rows[0][:3]) != ("指标", "统计类型", "分类"):
                    continue
                generation = sheet.title[:-3]
                base = _product(generation, "")
                month_columns = [(col, _month(value)) for col, value in enumerate(rows[0]) if col >= 3 and _month(value)]
                version_rows, controls = [], {}
                section, statistic = None, None
                for row_num, values in enumerate(rows[1:], 2):
                    if len(values) < 3:
                        continue
                    if values[0] is not None:
                        section = values[0]
                    if values[1] is not None:
                        statistic = values[1]
                    if values[0] in ("大定", "净大定", "交车锁单", "已交付") and values[1] == "数量":
                        controls[values[0]] = (row_num, values)
                    if section == "版本" and statistic == "数量" and _text(values[2]):
                        version_rows.append((row_num, values))
                if not version_rows:
                    continue
                source_id = "src_" + hashlib.sha256((path.name + "|" + sheet.title).encode()).hexdigest()[:12]
                sources.append({"id": source_id, "file": path.name, "sheet": sheet.title,
                                "range": f"A{min(r for r, _ in version_rows)}:{openpyxl.utils.get_column_letter(len(rows[0]))}{max(r for r, _ in version_rows)}",
                                "metric": metric["id"], "description": "仅取版本/数量块和年月列；不取占比、总计或其他选配维度。"})
                for row_num, values in version_rows:
                    config = _text(values[2])
                    pid = "v_" + hashlib.sha256((generation + "\x1f" + config).encode("utf-8")).hexdigest()[:16]
                    products.setdefault(pid, {"id": pid, "brand": base["brand"], "model": base["model"],
                                             "version": config, "source_version": f"{base['version']} · {config}", "model_year": base["version"],
                                             "configuration": config, "generation": generation,
                                             "name": f"{generation} {config}", "source_backed": True})
                    for col, month in month_columns:
                        if month > observation.strftime("%Y-%m"):
                            continue
                        value = _number(values[col]) if col < len(values) else None
                        if value is None:
                            missing.append({"product_id": pid, "month": month, "metric": metric["id"]})
                            continue
                        key = (pid, month, metric["id"])
                        if key in duplicate_keys:
                            raise ValueError(f"版本月指标重复：{generation} {config} {month} {metric['id']}")
                        duplicate_keys.add(key)
                        sales.append({"product_id": pid, "month": month, "metric": metric["id"], "value": value,
                                      "complete": False, "coverage_status": "monthly_source_cutoff_unknown",
                                      "known_days": None, "expected_days": calendar.monthrange(*map(int, month.split("-")))[1],
                                      "as_of": None, "source_id": source_id, "source_row": row_num,
                                      "source_cell": f"{openpyxl.utils.get_column_letter(col + 1)}{row_num}"})
                control_row, control_values = controls.get(metric["control_label"], (None, ()))
                for col, month in month_columns:
                    if month > observation.strftime("%Y-%m"):
                        continue
                    values = [_number(v[col]) if col < len(v) else None for _, v in version_rows]
                    expected = _number(control_values[col]) if col < len(control_values) else None
                    actual = sum(values) if all(v is not None for v in values) else None
                    delta = actual - expected if actual is not None and expected is not None else None
                    checks.append({"generation": generation, "metric": metric["id"], "month": month,
                                   "version_sum": actual, "source_total": expected, "difference": delta,
                                   "status": "matched" if delta == 0 else "difference" if delta is not None else "unavailable",
                                   "control_label": metric["control_label"], "source_id": source_id,
                                   "source_cell": f"{openpyxl.utils.get_column_letter(col + 1)}{control_row}" if control_row else None})
        finally:
            workbook.close()
    if not products:
        raise ValueError("未在选配月表找到真实版本数量块，未替换已有数据")
    # Reference-only generation data never enters the configuration product set.
    reference = _extract_generation_reference(source_dir, as_of=observation) if (source_dir / SOURCE_NAME).exists() else None
    months = sorted({r["month"] for r in sales})
    latest_month = months[-1] if months else None
    discrepancies = [r for r in checks if r["difference"] not in (None, 0)]
    option_conflicts = _option_consistency(source_dir)
    metric_coverage = {}
    for metric in VERSION_METRICS:
        records = [r for r in sales if r["metric"] == metric["id"]]
        mchecks = [r for r in checks if r["metric"] == metric["id"]]
        metric_coverage[metric["id"]] = {
            "data_from": min((r["month"] for r in records), default=None),
            "data_through": max((r["month"] for r in records), default=None),
            "exact_cutoff_date": None, "cutoff_status": "月表未提供精确截止日",
            "monthly_records": len(records), "product_count": len({r["product_id"] for r in records}),
            "monthly_total": sum(r["value"] for r in records),
            "source_controls_total": sum(r["source_total"] for r in mchecks if r["source_total"] is not None),
            "matched_controls": sum(r["status"] == "matched" for r in mchecks),
            "controls_with_difference": sum(r["status"] == "difference" for r in mchecks),
        }
    quality = [
        {"level": "warning", "code": "stale_version_data", "message": f"版本月表最新为 {latest_month}，当前为 {observation:%Y-%m}；后续月份缺失，不能作实时完成量。"},
        {"level": "warning", "code": "unknown_month_cutoff", "message": "版本月表未注明精确截止日，且与日表不能直接对齐；所有月值均作为来源参考，完整月状态未确认。"},
        {"level": "warning", "code": "version_control_difference", "message": f"{len(discrepancies)} 项车型月版本合计与同页小计存在差异，保留原数，不假定差异原因。", "count": len(discrepancies), "examples": discrepancies[:5]},
        {"level": "info", "code": "version_lock_source_rows", "message": "锁单选配表的版本数量取“版本/数量”明细，车型锁单取“交车锁单/数量”总行；两者原值独立保留。"},
        {"level": "warning", "code": "brand_control_conflicts", "message": f"源表有 {len(option_conflicts)} 项品牌小计与车型合计不同；看板只汇总真实版本行，不叠加品牌或车型总计。", "count": len(option_conflicts), "examples": option_conflicts[:3]},
    ]
    if missing:
        quality.append({"level": "warning", "code": "missing_version_cells", "message": f"{len(missing)} 个版本月单元格缺失，未补0。", "count": len(missing)})
    result = {
        "meta": {"generated_at": now.isoformat(timespec="seconds"), "timezone": "Asia/Shanghai",
                 "observation_date": observation.isoformat(), "data_through": latest_month, "latest_month": latest_month,
                 "data_through_label": f"{latest_month}（月表，截止日未提供）", "exact_cutoff_date": None,
                 "source_files": [_source_file(p) for p in sales_source_files(source_dir)],
                 "default_metric": "net_orders", "metrics": VERSION_METRICS, "metric_coverage": metric_coverage,
                 "unit": "单", "product_grain": "年款内具体配置版本", "brands": ["问界", "智界", "享界", "尊界", "尚界"],
                 "notes": [
                     "真实版本：仅从车型by月表中“版本/数量”块提取，不把年款或二级代际冒充配置版本。",
                     "不同年款独立；品牌、车型、版本逐级汇总，但不叠加原品牌总计、车型总计及其他选配维度。",
                     "大定选配表版本数量按净大定参考；锁单选配表版本数量与交车锁单总行分别读取，不要求合计相等。",
                     "月表最后一列“总计”部分行实际等于最后一个月，未作为跨月总计使用。",
                     "所有原始数量保持不变。版本合计与小计的±1差异不擅自修正或归因为四舍五入。",
                     "汇总文件代际日量保存在单独 reference 区域，未拆分为配置销量，未计入版本合计。",
                     "缺失月、空白和未知截止日均不补0；文件更新时间不作为数据日期。",
                 ], "quality": quality, "reconciliation": checks},
        "products": sorted(products.values(), key=lambda p: (p["brand"], p["model"], p["version"])),
        "sales": sorted(sales, key=lambda r: (r["month"], r["product_id"], r["metric"])),
        "sources": sources,
        "reference": {"description": "代际级逐日汇总参考，不属于配置版本月度实际量。", **reference} if reference else None,
    }
    return result


LOCK_BASIS_OPTIONS = {"unconfirmed", "source_version_quantity", "model_lock_row"}


def _resolve_lock_basis(explicit=None, settings_path=None):
    """A user-confirmed source interpretation persists across ordinary refreshes."""
    basis = explicit
    if basis is None:
        settings_path = Path(settings_path) if settings_path else ROOT / "data" / "source_settings.json"
        if settings_path.is_file():
            settings = json.loads(settings_path.read_text(encoding="utf-8-sig"))
            basis = settings.get("version_lock_basis", "unconfirmed")
        else:
            basis = "unconfirmed"
    if basis not in LOCK_BASIS_OPTIONS:
        raise ValueError("version_lock_basis 必须为 unconfirmed、source_version_quantity 或 model_lock_row")
    return basis


def _model_monthly_locks(source_dir, *, as_of=None):
    """Keep model lock controls at their actual grain; never allocate to versions."""
    path = Path(source_dir) / "鸿蒙智行锁单选配比例分析.xlsx"
    observation = _date(as_of) if as_of else datetime.now(SHANGHAI).date()
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    result = []
    try:
        for sheet in workbook:
            if not sheet.title.endswith("by月") or " " not in sheet.title:
                continue
            rows = list(sheet.iter_rows(values_only=True))
            if not rows or tuple(rows[0][:3]) != ("指标", "统计类型", "分类"):
                continue
            generation = sheet.title[:-3]
            product = _product(generation, "")
            source_id = "src_" + hashlib.sha256((path.name + "|" + sheet.title).encode()).hexdigest()[:12]
            lock_rows = [(index, row) for index, row in enumerate(rows, 1)
                         if len(row) >= 3 and row[0] == "交车锁单" and row[1] == "数量"]
            if len(lock_rows) > 1:
                raise ValueError(f"车型月交车锁单总量行重复：{sheet.title}")
            for row_number, row in lock_rows:
                for col, header in enumerate(rows[0]):
                    month = _month(header)
                    if col < 3 or not month or month > observation.strftime("%Y-%m"):
                        continue
                    value = _number(row[col]) if col < len(row) else None
                    result.append({"brand": product["brand"], "model": product["model"],
                                   "generation": generation, "model_year": product["version"],
                                   "month": month, "metric": "locked_orders", "value": value,
                                   "grain": "model_generation", "source_id": source_id,
                                   "source_cell": f"{openpyxl.utils.get_column_letter(col + 1)}{row_number}",
                                   "complete": False, "coverage_status": "monthly_source_cutoff_unknown",
                                   "as_of": None, "is_allocated": False})
    finally:
        workbook.close()
    return result


def _build_lock_data(version_data, model_sales, basis):
    """Keep version detail and model lock totals at their separate source grains."""
    if basis not in LOCK_BASIS_OPTIONS:
        raise ValueError("无效的版本锁单口径")
    confirmed = basis == "source_version_quantity"
    model_only = basis == "model_lock_row"
    version_rows = [r for r in version_data["sales"] if r["metric"] == "delivery_reference"]
    locks = []
    for row in version_rows:
        lock = {**row, "metric": "locked_orders", "value": row["value"] if confirmed else None,
                "value_status": "source_quantity_user_confirmed" if confirmed else "version_lock_basis_unconfirmed",
                "source_version_value": row["value"], "version_lock_basis": basis,
                "is_allocated": False}
        if not confirmed:
            lock["coverage_status"] = "not_available_at_version" if model_only else "version_lock_basis_unconfirmed"
        if model_only:
            lock["value_status"] = "not_available_at_version"
            lock.pop("source_version_value", None)
            lock["source_cell"] = None
            lock["source_row"] = None
        locks.append(lock)
    # A product can still exist when one of its source months is empty. Preserve
    # that missingness instead of deriving a quantity from another metric.
    known = [r for r in locks if r["value"] is not None]
    model_by_key = {(r["generation"], r["month"]): r for r in model_sales}
    product_by_id = {p["id"]: p for p in version_data["products"]}
    totals = defaultdict(list)
    for row in version_rows:
        product = product_by_id[row["product_id"]]
        totals[(product["generation"], row["month"])].append(row["value"])
    checks = []
    for key, model_row in model_by_key.items():
        values = totals.get(key, [])
        actual = sum(values) if values and all(v is not None for v in values) else None
        target = model_row["value"]
        delta = actual - target if actual is not None and target is not None else None
        checks.append({"generation": key[0], "month": key[1], "metric": "locked_orders",
                       "source_version_sum": actual, "source_total": target,
                       "difference": delta, "status": "matched" if delta == 0 else "difference" if delta is not None else "unavailable",
                       "source_id": model_row["source_id"], "source_cell": model_row["source_cell"],
                       "basis_confirmed": confirmed})
    differences = [c for c in checks if c["status"] == "difference"]
    meta = dict(version_data["meta"])
    metric = {"id": "locked_orders", "label": "锁单", "unit": "单",
              "source_file": "鸿蒙智行锁单选配比例分析.xlsx", "control_label": "交车锁单",
              "version_lock_basis": basis}
    meta.update({"default_metric": "locked_orders", "metrics": [metric],
                 "source_directory": "input_file", "version_lock_basis": basis,
                 "model_lock_months": len(model_sales), "reconciliation": checks,
                 "metric_coverage": {"locked_orders": {
                     "data_from": min((r["month"] for r in locks), default=None),
                     "data_through": max((r["month"] for r in locks), default=None),
                     "known_data_through": max((r["month"] for r in known), default=None),
                     "exact_cutoff_date": None, "cutoff_status": "月表未提供精确截止日",
                     "monthly_records": len(locks), "known_monthly_records": len(known),
                     "product_count": len({r["product_id"] for r in locks}),
                     "monthly_total": sum(r["value"] for r in known) if known else None,
                     "model_monthly_total": sum(r["value"] for r in model_sales if r["value"] is not None),
                     "model_monthly_records": len(model_sales),
                 }}})
    meta["quality"] = [q for q in meta["quality"] if q["code"] in ("stale_version_data", "unknown_month_cutoff")]
    if not model_only:
        meta["quality"].append({"level": "info" if confirmed else "warning", "code": "version_lock_control_difference",
                                "message": f"{len(differences)} 项车型月的版本数量合计与交车锁单总量不同；两个层级各取来源原值，不分摊、不强制相等。",
                                "count": len(differences), "examples": differences[:5]})
    if model_only:
        basis_note = "已按确认仅采用交车锁单数量行，版本计划汇总后与同月车型锁单对照。"
        meta["quality"].append({"level": "info", "code": "lock_basis_model_row", "message": basis_note})
    elif confirmed:
        basis_note = "版本锁单取锁单选配表“版本/数量”明细，车型锁单独立取“交车锁单/数量”总行；两者原值分别展示，不要求合计相等。"
        meta["quality"].append({"level": "info", "code": "version_lock_basis_user_confirmed", "message": basis_note})
    else:
        basis_note = "版本数量与同页“已交付”小计接近、与交车锁单不一致；确认前版本锁单留空，车型交车锁单单独展示，不能当作版本实绩。"
        meta["quality"].insert(0, {"level": "warning", "code": "version_locks_unconfirmed", "message": basis_note,
                                   "count": len(locks)})
    meta["notes"] = [
        "默认指标为锁单，月份和具体配置版本使用与 SOP 相同的产品键；原版本ID保持稳定。",
        basis_note,
        "车型月锁单来自“交车锁单/数量”行，版本月锁单来自“版本/数量”明细行；品牌和总计只累加车型数，不与版本数叠加。" if confirmed else
        "车型月锁单来自同页“交车锁单/数量”行，只在车型完整选中时独立显示，不与配置版本数叠加。",
        "只取月份列；源表“总计”列不参与跨月相加。缺失保持空白，明确的0保留为0。",
        "原始版本数量已排除在锁单口径之外；版本锁单无数据，不按占比或其他指标分摊。" if model_only else
        "版本数量按来源原值展示，不按车型总量分摊；净大定与代际日量另作参考。" if confirmed else
        "净大定和原锁单选配版本数仅在 version_reference 中留存，代际日量仅在 reference 中留存。",
        "月表未提供精确截止日，不用日表日期或文件修改时间冒充月表截止日。",
    ]
    sources = [dict(source) for source in version_data["sources"]]
    for source in sources:
        if source.get("metric") == "delivery_reference":
            source.update({"metric": "locked_orders", "version_lock_basis": basis,
                           "description": "按确认仅读取交车锁单/数量行，版本原数量不用于锁单；版本计划汇总后与车型锁单对照。" if model_only else
                           "版本锁单读取版本/数量明细行，车型锁单独立读取交车锁单/数量总行；差异保留，不分摊。" if confirmed else
                           "车型锁单读取交车锁单/数量行；配置数量仅在用户确认口径后作为版本锁单，未进行分摊。"})
            if model_only:
                cells = [r["source_cell"] for r in model_sales if r["source_id"] == source["id"]]
                if cells:
                    ordered = sorted(cells, key=lambda cell: openpyxl.utils.cell.coordinate_to_tuple(cell)[1])
                    source["range"] = ordered[0] + ":" + ordered[-1]
    return {"meta": meta, "products": version_data["products"], "sales": locks,
            "model_sales": model_sales, "sources": sources,
            "reference": version_data.get("reference"),
            "version_reference": {"description": "版本/数量原值已作为版本锁单；此处另保留来源参考，不与车型锁单相加。" if confirmed else
                                              "原始版本数量口径参考，不参与锁单与SOP主表相加。",
                                  "metrics": version_data["meta"]["metrics"], "sales": version_data["sales"],
                                  "reconciliation": version_data["meta"]["reconciliation"]}}


def extract(source_dir=None, *, as_of=None, version_lock_basis=None):
    """Read inputs without writing any file, including during source validation."""
    source_dir = Path(source_dir) if source_dir is not None else DEFAULT_SOURCE_DIR
    basis = _resolve_lock_basis(version_lock_basis)
    version_data = _extract_version_reference(source_dir, as_of=as_of)
    models = _model_monthly_locks(source_dir, as_of=as_of)
    return _build_lock_data(version_data, models, basis)


def refresh(source_dir=None, output_path=None, *, as_of=None, version_lock_basis=None):
    """Refresh the local JSON atomically from input_file; never modify sources."""
    result = extract(source_dir, as_of=as_of, version_lock_basis=version_lock_basis)
    output_path = Path(output_path) if output_path is not None else ROOT / "data" / "sales.json"
    output_path.parent.mkdir(parents=True, exist_ok=True)
    temporary = output_path.with_suffix(output_path.suffix + ".tmp")
    temporary.write_text(json.dumps(result, ensure_ascii=False, indent=2, allow_nan=False), encoding="utf-8")
    temporary.replace(output_path)
    return result


def main():
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
    parser = argparse.ArgumentParser(description="只读导入销量 Excel，为 SOP 看板刷新月度实际量")
    parser.add_argument("--source-dir", type=Path, default=DEFAULT_SOURCE_DIR)
    parser.add_argument("--output", type=Path, default=ROOT / "data" / "sales.json")
    parser.add_argument("--as-of", help="排除该日及未来数据，YYYY-MM-DD；默认北京时间今天")
    parser.add_argument("--version-lock-basis", choices=sorted(LOCK_BASIS_OPTIONS), help="model_lock_row仅采用交车锁单行；source_version_quantity仅适用于已明确确认版本数量为锁单的来源")
    args = parser.parse_args()
    result = refresh(args.source_dir, args.output, as_of=args.as_of, version_lock_basis=args.version_lock_basis)
    print(json.dumps({"products": len(result["products"]), "monthly_records": len(result["sales"]), "data_through": result["meta"]["data_through"], "metrics": result["meta"]["metric_coverage"], "quality": result["meta"]["quality"]}, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
