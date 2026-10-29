"""Offline XLSX writer for dashboard exports.

This mirrors the visible grid and trusted inline manifest produced by
``scripts/build_export_workbook.mjs``.  It only renders a disposable download;
the database and primary maintenance workbook are never changed here.
"""
from __future__ import annotations

import json
from collections import OrderedDict
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Protection
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from sop_brand import brand_order_key
from sop_plan_scope import MODEL_PLAN_LABEL
from sop_xlsx_style import PLAN_TITLE, plan_guidance, setup_plan_sheet, style_plan_row, write_formula_caches


MAX_SAFE_INTEGER = 2**53 - 1


def _column(number):
    return get_column_letter(number)


def _quantity(value):
    return value is None or (type(value) is int and 0 <= value <= MAX_SAFE_INTEGER)


def _sum_known(values):
    known = [value for value in values if value is not None]
    result = sum(known) if known else None
    if not _quantity(result):
        raise ValueError("导出汇总量超过安全整数范围。")
    return result


def _text(cell, value):
    """Keep user supplied labels literal, including names beginning with '='."""
    cell.value = value
    if isinstance(value, str):
        cell.data_type = "s"


def _title():
    return PLAN_TITLE


def _period(months):
    return months[0] if months[0] == months[-1] else f"{months[0]}—{months[-1]}"


def _subtitle(data, inline):
    if inline:
        if data["include_versions"]:
            return plan_guidance(data["include_locks"], any(p.get("read_only") for p in data.get("roundtrip", {}).get("products", [])))
        return "单位：辆。当前为车型汇总，回填计划请导出版本明细。年度合计为全年SOP计划量。"
    if data.get("roundtrip"):
        return "单位：辆。空白表示未填写，0表示明确不安排。年度合计为全年SOP；锁单行只读。"
    return "单位：辆。空白表示未填写，0表示明确不安排。年度合计为全年SOP；带“锁单”字样的行为交车锁单数。"


def _layout(data):
    months = data["months"]
    if not months or months != sorted(set(months)):
        raise ValueError("导出月份必须按时间排列且不重复。")
    identity = ["品牌", "车型"] + (["版本"] if data["include_versions"] else [])
    first_value = len(identity) + 1
    columns, years = [], []
    for year in dict.fromkeys(month[:4] for month in months):
        first = first_value + len(columns)
        year_months = [month for month in months if month.startswith(year + "-")]
        for month in year_months:
            columns.append({"column": first_value + len(columns), "kind": "month",
                            "month": month, "sourceIndex": months.index(month),
                            "label": f"{int(month[-2:])}月", "historical": month < data["current_month"]})
            if month.endswith("-12"):
                columns.append({"column": first_value + len(columns), "kind": "year_total",
                                "year": year, "label": "年度合计"})
        years.append({"year": year, "first_column": first,
                      "last_column": first_value + len(columns) - 1,
                      "full": len(year_months) == 12})
    width = len(identity) + len(columns)
    headers = [[None] * width, [None] * width]
    for index, label in enumerate(identity):
        headers[0][index] = label
    for year in years:
        headers[0][year["first_column"] - 1] = year["year"] + "年"
    for item in columns:
        headers[1][item["column"] - 1] = item["label"]
    merges = [f"{_column(i)}3:{_column(i)}4" for i in range(1, len(identity) + 1)]
    merges += [f"{_column(y['first_column'])}3:{_column(y['last_column'])}3"
               for y in years if y["last_column"] > y["first_column"]]

    metric_index = 4 if data["include_versions"] else 3
    source_rows = []
    ids = data.get("roundtrip", {}).get("row_product_ids")
    counts = data.get("roundtrip", {}).get("annual_known_counts")
    if ids is not None and len(ids) != len(data["rows"]):
        raise ValueError("版本行标识数量不一致。")
    for index, raw in enumerate(data["rows"]):
        source_rows.append({
            "index": index, "level": raw[0], "brand": raw[1], "model": raw[2],
            "version": raw[3] if data["include_versions"] else "",
            "metric": raw[metric_index], "values": raw[metric_index + 1:-1],
            "annual": data["annual_plans"][index],
            "product_id": ids[index] if ids is not None else None,
            "annual_known": counts[index] if counts is not None else None,
        })
    groups = OrderedDict()
    for source in source_rows:
        groups.setdefault(source["brand"], OrderedDict()).setdefault(source["model"], []).append(source)
    fields = ["锁单数", "SOP计划量"] if data["include_locks"] else ["SOP计划量"]
    rows, next_row = [], 5

    def add_source(source, kind):
        nonlocal next_row
        row = {"row": next_row, "kind": kind, "metric": source["metric"],
               "product_id": source["product_id"], "labels": [None] * len(identity),
               "values": {}, "formulas": {}, "version": source["version"], "annual_known": source["annual_known"]}
        next_row += 1
        for item in columns:
            column = item["column"]
            value = source["values"][item["sourceIndex"]] if item["kind"] == "month" else source["annual"][item["year"]]
            if not _quantity(value):
                raise ValueError("导出包含无效数量。")
            row["values"][column] = value
            if kind != "product" or source["metric"] != "SOP计划量" or item["kind"] != "year_total":
                continue
            year = next(y for y in years if y["year"] == item["year"])
            if not data.get("roundtrip") and not year["full"]:
                continue
            ref = f"{_column(year['first_column'])}{row['row']}:{_column(column - 1)}{row['row']}"
            if year["full"]:
                row["formulas"][column] = f'=IF(COUNT({ref})=0,"",SUM({ref}))'
                continue
            selected = [source["values"][c["sourceIndex"]] for c in columns
                        if c["kind"] == "month" and c["month"].startswith(item["year"] + "-")]
            selected_known = [v for v in selected if v is not None]
            offscreen_count = source["annual_known"][item["year"]] - len(selected_known)
            offscreen_sum = (value or 0) - sum(selected_known)
            if not 0 <= offscreen_count <= 12 - len(selected) or not _quantity(offscreen_sum) or (offscreen_count == 0 and offscreen_sum != 0) or (value is None and source["annual_known"][item["year"]] != 0):
                raise ValueError("年度计划量与所选月份不一致。")
            row["formulas"][column] = f'=IF(COUNT({ref})+{offscreen_count}=0,"",SUM({ref})+{offscreen_sum})'
        rows.append(row)
        return row

    def sum_into(row, children):
        model_inputs = [child for child in children if child.get("version") == MODEL_PLAN_LABEL]
        if row["kind"] == "model_total" and row["metric"] == "SOP计划量" and model_inputs:
            versions = [child for child in children if child not in model_inputs]
            for item in columns:
                column, number = item["column"], row["row"]
                letter = _column(column)
                if item["kind"] == "month":
                    known = [c for c in model_inputs if c["values"][column] is not None]
                    expected = _sum_known(c["values"][column] for c in (known or versions))
                    if row["values"][column] != expected:
                        raise ValueError("车型总量与有效计划口径不一致。")
                    direct = ",".join(f"{letter}{c['row']}" for c in model_inputs)
                    detail = ",".join(f"{letter}{c['row']}" for c in versions)
                    fallback = f'IF(COUNT({detail})=0,"",SUM({detail}))' if detail else '""'
                    row["formulas"][column] = f'=IF(COUNT({direct})>0,SUM({direct}),{fallback})'
                else:
                    year = next(y for y in years if y["year"] == item["year"])
                    if not year["full"] and not data.get("roundtrip"):
                        continue
                    selected = [row["values"][c["column"]] for c in columns if c["kind"] == "month" and c["month"].startswith(item["year"] + "-")]
                    known = [value for value in selected if value is not None]
                    off_count = 0 if year["full"] else row["annual_known"][item["year"]] - len(known)
                    off_sum = 0 if year["full"] else (row["values"][column] or 0) - sum(known)
                    ref = f"{_column(year['first_column'])}{number}:{_column(column-1)}{number}"
                    row["formulas"][column] = f'=IF(COUNT({ref})+{off_count}=0,"",SUM({ref})+{off_sum})'
            return
        for item in columns:
            column = item["column"]
            if item["kind"] == "year_total" and row["metric"] == "锁单数":
                row["values"][column] = None
                continue
            expected = _sum_known(child["values"][column] for child in children)
            if column in row["values"] and row["values"][column] != expected:
                raise ValueError("车型 SOP 与版本明细不一致。")
            row["values"][column] = expected
            refs = ",".join(f"{_column(column)}{child['row']}" for child in children)
            row["formulas"][column] = f'=IF(COUNT({refs})=0,"",SUM({refs}))' if refs else '=""'

    def total_row(metric, children, kind, label):
        nonlocal next_row
        row = {"row": next_row, "kind": kind, "metric": metric,
               "product_id": None, "labels": [None] * len(identity),
               "values": {}, "formulas": {}, "version": ""}
        next_row += 1
        row["labels"][0 if kind == "grand_total" else 1] = label
        sum_into(row, children)
        rows.append(row)
        return row

    def merge_total(row, grand=False):
        start = 1 if grand else 2
        if len(identity) > start:
            merges.append(f"{_column(start)}{row['row']}:{_column(len(identity))}{row['row']}")

    inline = data.get("roundtrip", {}).get("inline_schema") == 2
    brand_totals = {field: [] for field in fields}
    for brand, models in sorted(groups.items(), key=lambda item: brand_order_key(item[0])):
        brand_first = next_row
        model_totals = {field: [] for field in fields}
        for model, group in models.items():
            model_first = next_row
            summaries = {field: [r for r in group if r["level"] == "车型汇总" and r["metric"] == field] for field in fields}
            details = {field: [r for r in group if r["level"] == "版本明细" and r["metric"] == field] for field in fields}
            if any(len(value) != 1 for value in summaries.values()):
                raise ValueError(f"车型汇总行不完整：{brand}/{model}")
            if data["include_versions"] and not details["SOP计划量"]:
                raise ValueError(f"计划明细行不完整：{brand}/{model}")
            detail_rows = {field: [] for field in fields}
            if data["include_versions"]:
                version_names = [[r["version"] for r in details[field] if r["version"] != MODEL_PLAN_LABEL] for field in fields]
                if any(names != version_names[0] for names in version_names):
                    raise ValueError(f"版本顺序不一致：{brand}/{model}")
                for field in fields:
                    for source in details[field]:
                        row = add_source(source, "product")
                        row["labels"][2] = source["version"] + ("（锁单数）" if field == "锁单数" else "")
                        row["version"] = row["labels"][2]
                        detail_rows[field].append(row)
                if inline:
                    for field in fields:
                        block = detail_rows[field]
                        if not block:
                            continue
                        block[0]["labels"][1] = model
                        if len(block) > 1:
                            merges.append(f"B{block[0]['row']}:B{block[-1]['row']}")
            if next_row > model_first and not inline:
                next(row for row in rows if row["row"] == model_first)["labels"][1] = model
                if next_row - model_first > 1:
                    merges.append(f"B{model_first}:B{next_row - 1}")
            for field in fields:
                row = add_source(summaries[field][0], "model_total")
                row["labels"][1] = model + (" 锁单" if field == "锁单数" else "") + " 合计"
                if data["include_versions"] and field == "SOP计划量":
                    sum_into(row, detail_rows[field])
                merge_total(row)
                model_totals[field].append(row)
        for field in fields:
            row = total_row(field, model_totals[field], "brand_total", brand + (" 锁单" if field == "锁单数" else "") + " 合计")
            merge_total(row)
            brand_totals[field].append(row)
        first = next(r for r in rows if r["row"] == brand_first)
        first["labels"][0] = brand
        if next_row - brand_first > 1:
            merges.append(f"A{brand_first}:A{next_row - 1}")
        next_row += 1
    for field in fields:
        row = total_row(field, brand_totals[field], "grand_total", "锁单总计" if field == "锁单数" else "总计")
        merge_total(row, grand=True)
    last_row = next_row - 1
    product_column = width + 1 if inline else None
    kind_column = width + 2 if inline else None
    key_column = width + 3 if inline else None
    if inline:
        headers[0].extend([None, None, None])
        headers[1].extend(["product_id", "row_kind", "row_key"])
        catalog = {product["id"]: product for product in data["roundtrip"]["products"]}
        for row in rows:
            row["key"] = ("lock:" if row["metric"] == "锁单数" else "plan:") + row["product_id"] if row["kind"] == "product" else "row:" + str(row["row"])
            row["hidden"] = row["kind"] == "product" and catalog[row["product_id"]].get("active") is False
    return {"identity": identity, "first_value": first_value, "columns": columns, "headers": headers,
            "rows": rows, "merges": merges, "last_row": last_row, "width": width,
            "last_column": key_column or width, "product_column": product_column,
            "kind_column": kind_column, "key_column": key_column, "inline": inline}


def _manifest(data, layout):
    context = data["roundtrip"]
    return {"columns": layout["columns"], "merges": layout["merges"],
            "products": context["products"],
            "rows": [{"row": row["row"], "key": row["key"], "kind": row["kind"],
                      "product_id": row["product_id"], "metric": row["metric"],
                      "labels": row["labels"], "hidden": row["hidden"],
                      "values": {item["month"]: row["values"][item["column"]]
                                 for item in layout["columns"] if item["kind"] == "month"}}
                     for row in layout["rows"]],
            "product_column": layout["product_column"], "kind_column": layout["kind_column"],
            "key_column": layout["key_column"], "last_column": layout["last_column"],
            "headers": layout["headers"],
            "editable_months_by_product": context["editable_months_by_product"],
            "manual_editable_months": context["manual_editable_months"],
            "plan_revisions_by_product": context.get("plan_revisions_by_product", {})}


def write_export_workbook(payload: dict) -> None:
    """Write the dashboard download to ``payload['output_path']``."""
    if not isinstance(payload, dict) or payload.get("schema_version") != 1 or not payload.get("output_path"):
        raise ValueError("无效的导出参数。")
    data = payload
    layout = _layout(data)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "SOP计划"
    setup_plan_sheet(sheet, last_row=layout["last_row"], last_column=layout["last_column"],
                     visible_last_column=layout["width"], columns=layout["columns"],
                     identity_columns=len(layout["identity"]))
    _text(sheet.cell(1, 1), _title())
    _text(sheet.cell(1, layout["first_value"]), _period(data["months"]))
    _text(sheet.cell(2, 1), _subtitle(data, layout["inline"]))
    if layout["inline"] and layout["width"] <= 6:
        sheet.row_dimensions[2].height = 58 if layout["width"] <= 4 else 44
        sheet.cell(2, 1).alignment = Alignment(vertical="center", wrap_text=True)
    for row_index, header in enumerate(layout["headers"], 3):
        for column, value in enumerate(header, 1):
            if value is not None:
                _text(sheet.cell(row_index, column), value)
    validation = DataValidation(type="whole", operator="between", formula1=0,
                                formula2=1000000000000, allow_blank=True)
    validation.showErrorMessage = True
    validation.errorStyle = "stop"
    validation.errorTitle = "请输入非负整数"
    validation.error = "SOP计划量须为非负整数；未填写时请留空。"
    validation.showInputMessage = True
    validation.promptTitle = "SOP计划量（辆）"
    validation.prompt = "只填非负整数。空白表示未填写，0表示明确不安排。"
    caches = {}
    context = data.get("roundtrip", {})
    for row in layout["rows"]:
        number = row["row"]
        for index, label in enumerate(row["labels"], 1):
            if label is not None:
                _text(sheet.cell(number, index), label)
        month_states = {}
        if context and data["include_versions"] and row["kind"] == "product" and row["metric"] == "SOP计划量":
            editable = context["editable_months_by_product"][row["product_id"]]
            product = next((p for p in context.get("products", []) if p["id"] == row["product_id"]), None)
            for item in layout["columns"]:
                if item["kind"] != "month":
                    continue
                is_editable = item["month"] in editable
                state = {"editable": is_editable}
                if product:
                    state["applicable"] = is_editable or ((not product.get("start_month") or item["month"] >= product["start_month"]) and (not product.get("end_month") or item["month"] <= product["end_month"]))
                month_states[f"{_column(item['column'])}{number}"] = state
        style_plan_row(sheet, row=number, kind=row["kind"], identity_columns=len(layout["identity"]),
                       visible_last_column=layout["width"], columns=layout["columns"],
                       version=row["version"], month_states=month_states)
        for item in layout["columns"]:
            column = item["column"]
            cell = sheet.cell(number, column)
            if column in row["formulas"]:
                cell.value = row["formulas"][column]
                caches[cell.coordinate] = row["values"][column]
            else:
                cell.value = row["values"][column]
            if (context and data["include_versions"] and row["kind"] == "product"
                    and row["metric"] == "SOP计划量" and item["kind"] == "month"
                    and item["month"] in context["editable_months_by_product"][row["product_id"]]):
                cell.protection = Protection(locked=False)
                validation.add(cell)
        if layout["inline"]:
            _text(sheet.cell(number, layout["product_column"]), row["product_id"])
            _text(sheet.cell(number, layout["kind_column"]), row["kind"])
            _text(sheet.cell(number, layout["key_column"]), row["key"])
            if row["hidden"]:
                sheet.row_dimensions[number].hidden = True
            if data["include_versions"] and row["kind"] == "product" and row["metric"] == "SOP计划量":
                for column in range(1, 4):
                    sheet.cell(number, column).protection = Protection(locked=bool(product and product.get("read_only")))
    # An empty dataValidations container is invalid SpreadsheetML. Summary
    # downloads and versions with no editable months must omit it entirely.
    if validation.sqref:
        sheet.add_data_validation(validation)
    for reference in layout["merges"]:
        sheet.merge_cells(reference)
    sheet.freeze_panes = f"{_column(layout['first_value'])}5"
    if not (layout["inline"] and data["include_versions"]):
        sheet.protection.sheet = True
        sheet.protection.selectLockedCells = False
        sheet.protection.selectUnlockedCells = False

    if context:
        meta = workbook.create_sheet("_sop_import")
        meta.sheet_state = "veryHidden"
        if layout["inline"]:
            encoded = json.dumps(_manifest(data, layout), ensure_ascii=False, separators=(",", ":"))
            chunks = [encoded[index:index + 14000] for index in range(0, len(encoded), 14000)]
            meta_rows = [["SOP_IMPORT", 2], ["token", context["token"]],
                         ["include_versions", data["include_versions"]], ["sheet", sheet.title],
                         ["manifest_json", None]]
            meta_rows.extend([index, chunk] for index, chunk in enumerate(chunks, 1))
        else:
            raise ValueError("离线导出需要 inline_schema=2 的回导上下文。")
        for number, values in enumerate(meta_rows, 1):
            for column, value in enumerate(values, 1):
                if isinstance(value, str):
                    _text(meta.cell(number, column), value)
                else:
                    meta.cell(number, column, value)
    output = Path(data["output_path"])
    output.parent.mkdir(parents=True, exist_ok=True)
    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    workbook.save(output)
    hidden = {sheet.title: (layout["width"] + 1, layout["last_column"])} if layout["inline"] else {}
    write_formula_caches(output, workbook.sheetnames, {sheet.title: caches}, hidden)
